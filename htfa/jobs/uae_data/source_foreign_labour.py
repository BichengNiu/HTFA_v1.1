"""UAE 外籍劳工流入代理指标：入库与写表。

逻辑移植自 ``data/foreignEmployer/``（download_data.py 下载 +
process_data.py 解析/质检/写表）。原始文件写入 ``data/UAE/raw/foreign_labour/``；
入库表 ``foreign_labour_monthly`` 以 (period, series) 长表保存全部序列（含辅助
字段）；工作簿 ``月度_外籍劳动力`` 只写 8 列（日期+三国原始官方指标+合计）。
"""

from __future__ import annotations

import calendar
import csv
import html
import json
import os
import re
import shutil
import sys
import time
from collections import defaultdict
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path
from statistics import mean
from typing import Any, Iterator
from urllib.parse import urljoin

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

from .paths import DATA_DIR, SCRIPTS_DIR


from . import db  # noqa: E402

# ---------------------------------------------------------------------------
# 常量
# ---------------------------------------------------------------------------

RAW_DIR = DATA_DIR / "raw" / "foreign_labour"
MANIFEST_PATH = RAW_DIR / "manifest.json"
BACKUP_DIR = RAW_DIR / "backups"

SHEET_NAME = "月度_外籍劳动力"
BASELINE_START = "2025-05"
BASELINE_END = "2025-11"

DMW_CMS_URL = "https://wcms.dmw.gov.ph/api/compendium-releases"
DMW_ARCHIVE_ROOT = (
    "https://dmw.gov.ph/archives/v1/resources/dsms/DMW/Externals/2025/"
    "Statistics"
)
DOFE_ROOT = "https://dofe.gov.np"
DOFE_MONTHLY_URL = f"{DOFE_ROOT}/category/monthly/"
BMET_OEP_URL = "https://www.oep.gov.bd/reports/country-clearance"
UAE_OEP_COUNTRY_ID = 188

DEVANAGARI_DIGITS = str.maketrans("०१२३४५६७८९", "0123456789")
NEPAL_MONTHS = {
    "वैशाख": 1,
    "बैशाख": 1,
    "जेष्ठ": 2,
    "जेठ": 2,
    "अषाढ": 3,
    "असार": 3,
    "श्रावण": 4,
    "साउन": 4,
    "भाद्र": 5,
    "भदौ": 5,
    "असोज": 6,
    "आश्विन": 6,
    "कार्तिक": 7,
    "कात्तिक": 7,
    "मंसिर": 8,
    "मङ्सिर": 8,
    "पौष": 9,
    "पुष": 9,
    "माघ": 10,
    "फागुन": 11,
    "चैत": 12,
    "चैत्र": 12,
}

WORKBOOK_HEADERS = [
    "日期",
    "菲律宾_DMW部署_总计",
    "菲律宾_DMW部署_新雇",
    "菲律宾_DMW部署_再雇",
    "尼泊尔_DoFE批准_含再入境",
    "尼泊尔_DoFE批准_不含再入境",
    "孟加拉国_BMET出境许可",
    "重点三国合计_可比月",
]
SERIES_KEYS = [
    "philippines_total",
    "philippines_new_hires",
    "philippines_rehires",
    "nepal_with_reentry",
    "nepal_without_reentry",
    "bangladesh_clearance",
    "proxy_sum",
    "proxy_index",
    "source_count",
]


# ---------------------------------------------------------------------------
# 下载器（平移自 download_data.py）
# ---------------------------------------------------------------------------


def build_session() -> requests.Session:
    session = requests.Session()
    retries = Retry(
        total=4,
        connect=4,
        read=4,
        backoff_factor=1.0,
        status_forcelist=(429, 500, 502, 503, 504),
        allowed_methods=("GET",),
    )
    session.mount("https://", HTTPAdapter(max_retries=retries))
    session.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 Chrome/131 Safari/537.36"
            ),
            "Accept": "*/*",
        }
    )
    return session


def get_bytes(session: requests.Session, url: str) -> tuple[bytes, str]:
    last_error: Exception | None = None
    for attempt in range(3):
        try:
            response = session.get(url, timeout=(30, 240))
            response.raise_for_status()
            return response.content, response.headers.get("Content-Type", "")
        except (requests.ConnectionError, requests.Timeout) as error:
            last_error = error
            if attempt < 2:
                time.sleep(2**attempt)
    assert last_error is not None
    raise last_error


def save_download(
    session: requests.Session,
    url: str,
    destination: Path,
    expected_magic: bytes,
    refresh: bool,
) -> None:
    if destination.exists() and not refresh:
        return
    content, content_type = get_bytes(session, url)
    if not content.startswith(expected_magic):
        preview = content[:100].decode("utf-8", errors="replace")
        raise RuntimeError(
            f"Unexpected response for {url}: {content_type}; {preview!r}"
        )
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix(destination.suffix + ".tmp")
    temporary.write_bytes(content)
    temporary.replace(destination)


def dmw_static_2025_urls() -> list[dict[str, Any]]:
    months = [
        (1, "JAN", "2024 Jan 2025"),
        (2, "FEB", "2024 Feb 2025"),
        (3, "MAR", "2024 Mar 2025"),
        (4, "APR", "2024 Apr 2025"),
        (5, "MAY", "2024 May 2025"),
        (6, "JUNE", "2024 June 2025"),
        (7, "JULY", "2024 July 2025"),
        (8, "AUGUST", "2024 August 2025"),
        (9, "SEPTEMBER", "2024 September 2025"),
        (10, "OCTOBER", "2024 October 2025"),
    ]
    records: list[dict[str, Any]] = []
    for month, folder, label in months:
        filename = f"TAB 11 - Landbased Deployment by Destination ({label}).xlsx"
        records.append(
            {
                "year": 2025,
                "month": month,
                "url": f"{DMW_ARCHIVE_ROOT}/{folder}/{filename}",
                "source": "DMW archive",
            }
        )
    filename = (
        "TAB 11 - Landbased Deployment by Destination (2024-November 2025).xlsx"
    )
    records.append(
        {
            "year": 2025,
            "month": 11,
            "url": (
                "https://dmw.gov.ph/archives/v1/resources/dsms/DMW/"
                f"Externals/2025/Statistics/NOVEMBER/{filename}"
            ),
            "source": "DMW archive",
        }
    )
    return records


def dmw_cms_urls(session: requests.Session) -> list[dict[str, Any]]:
    params = {
        "filters[release_type][$eq]": "monthly",
        "populate[reports][populate][files][populate][0]": "file",
        "sort[0]": "year:asc",
        "sort[1]": "month:asc",
        "pagination[pageSize]": 100,
    }
    response = session.get(DMW_CMS_URL, params=params, timeout=90)
    response.raise_for_status()
    records: list[dict[str, Any]] = []
    for release in response.json().get("data", []):
        report = next(
            (
                item
                for item in release.get("reports", [])
                if str(item.get("tab_number")) == "11"
            ),
            None,
        )
        if report is None:
            continue
        file_record = next(
            (
                item
                for item in report.get("files", [])
                if item.get("file_type") == "excel"
            ),
            None,
        )
        if not file_record:
            continue
        file_url = (file_record.get("file") or {}).get("url")
        if not file_url:
            continue
        records.append(
            {
                "year": int(release["year"]),
                "month": int(release["month"]),
                "url": file_url,
                "published_at": release.get("publishedAt"),
                "source": "DMW CMS",
            }
        )
    return records


def download_philippines(
    session: requests.Session, refresh: bool
) -> list[dict[str, Any]]:
    directory = RAW_DIR / "philippines"
    records = dmw_static_2025_urls() + dmw_cms_urls(session)
    unique = {(record["year"], record["month"]): record for record in records}
    results = []
    for (_year, _month), record in sorted(unique.items()):
        destination = directory / (
            f"philippines_dmw_{record['year']}_{record['month']:02d}.xlsx"
        )
        save_download(session, record["url"], destination, b"PK", refresh)
        results.append(
            {**record, "path": str(destination.relative_to(DATA_DIR))}
        )
    return results


def strip_tags(value: str) -> str:
    value = re.sub(r"<[^>]+>", " ", value)
    return re.sub(r"\s+", " ", html.unescape(value)).strip()


def discover_nepal_reports(session: requests.Session) -> list[dict[str, Any]]:
    reports: dict[tuple[int, int], dict[str, Any]] = {}
    for page in range(1, 8):
        response = session.get(DOFE_MONTHLY_URL, params={"page": page}, timeout=90)
        if response.status_code == 404 and page > 1:
            break
        response.raise_for_status()
        matches = list(
            re.finditer(
                r'<h3[^>]*class="[^"]*card__title[^"]*"[^>]*>\s*'
                r'<a[^>]+href="\s*([^"]+?)\s*"[^>]*>(.*?)</a>',
                response.text,
                flags=re.IGNORECASE | re.DOTALL,
            )
        )
        found = 0
        for match in matches:
            title = strip_tags(match.group(2))
            parsed = re.search(
                r"([०-९]{4})\s+(\S+)\s+महिनाको\s+श्रम\s+स्वीकृति\s+विवरण",
                title,
            )
            if not parsed:
                continue
            year = int(parsed.group(1).translate(DEVANAGARI_DIGITS))
            month_name = parsed.group(2)
            month = NEPAL_MONTHS.get(month_name)
            if month is None:
                raise RuntimeError(f"Unknown Nepal month name: {month_name}")
            reports[(year, month)] = {
                "bs_year": year,
                "bs_month": month,
                "title": title,
                "detail_url": urljoin(DOFE_ROOT, match.group(1).strip()),
            }
            found += 1
        if not found:
            break

    for record in reports.values():
        response = session.get(record["detail_url"], timeout=90)
        response.raise_for_status()
        match = re.search(
            r'https://giwmscdnone\.gov\.np/media/pdf_upload/[^"\']+?\.pdf',
            response.text,
            flags=re.IGNORECASE,
        )
        if not match:
            raise RuntimeError(f"PDF link not found: {record['detail_url']}")
        record["url"] = html.unescape(match.group(0))
    return [reports[key] for key in sorted(reports)]


def download_nepal(
    session: requests.Session, refresh: bool
) -> list[dict[str, Any]]:
    directory = RAW_DIR / "nepal"
    results = []
    for record in discover_nepal_reports(session):
        destination = directory / (
            f"nepal_dofe_{record['bs_year']}_{record['bs_month']:02d}.pdf"
        )
        save_download(session, record["url"], destination, b"%PDF", refresh)
        results.append(
            {**record, "path": str(destination.relative_to(DATA_DIR))}
        )
    return results


def iter_complete_months(start_year: int, start_month: int):
    today = date.today()
    final_year = today.year if today.month > 1 else today.year - 1
    final_month = today.month - 1 if today.month > 1 else 12
    year, month = start_year, start_month
    while (year, month) <= (final_year, final_month):
        yield year, month
        if month == 12:
            year, month = year + 1, 1
        else:
            month += 1


def download_bangladesh(
    session: requests.Session, refresh: bool
) -> list[dict[str, Any]]:
    directory = RAW_DIR / "bangladesh"
    directory.mkdir(parents=True, exist_ok=True)
    results = []
    headers = {"X-Requested-With": "XMLHttpRequest", "Accept": "application/json"}
    for year, month in iter_complete_months(2023, 7):
        destination = directory / f"bangladesh_bmet_{year}_{month:02d}.json"
        start = date(year, month, 1)
        end = date(year, month, calendar.monthrange(year, month)[1])
        params = {
            "draw": 1,
            "start": 0,
            "length": 100,
            "approval_date_from": start.isoformat(),
            "approval_date_to": end.isoformat(),
            "country_id": UAE_OEP_COUNTRY_ID,
            "gender_id": "",
            "category_id": "",
            "all_skills": 0,
        }
        if not destination.exists() or refresh:
            response = session.get(
                BMET_OEP_URL, params=params, headers=headers, timeout=90
            )
            response.raise_for_status()
            payload = response.json()
            if "payload" not in payload:
                raise RuntimeError(
                    f"Unexpected OEP response for {year}-{month:02d}"
                )
            temporary = destination.with_suffix(".json.tmp")
            temporary.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            temporary.replace(destination)
            time.sleep(0.1)
        results.append(
            {
                "year": year,
                "month": month,
                "url": BMET_OEP_URL,
                "query": params,
                "path": str(destination.relative_to(DATA_DIR)),
            }
        )
    return results


def download_all(refresh: bool = False) -> None:
    session = build_session()
    manifest: dict[str, Any] = {
        "downloaded_on": date.today().isoformat(),
        "sources": {},
    }
    if MANIFEST_PATH.exists():
        existing = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
        manifest["sources"].update(existing.get("sources", {}))
    manifest["sources"]["philippines"] = download_philippines(session, refresh)
    manifest["sources"]["nepal"] = download_nepal(session, refresh)
    manifest["sources"]["bangladesh"] = download_bangladesh(session, refresh)
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    MANIFEST_PATH.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )


# ---------------------------------------------------------------------------
# 解析（平移自 process_data.py）
# ---------------------------------------------------------------------------


def month_end(year: int, month: int) -> date:
    return date(year, month, calendar.monthrange(year, month)[1])


def numeric(value: Any) -> int:
    if value in (None, "", "-"):
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    return int(str(value).replace(",", "").strip())


def find_uae_row(worksheet) -> tuple[Any, ...]:
    for row in worksheet.iter_rows(values_only=True):
        if any(
            "UNITED ARAB EMIRATES" in str(value).upper()
            for value in row
            if value is not None
        ):
            return row
    raise RuntimeError(f"UAE row not found in {worksheet.title}")


def parse_philippines() -> dict[str, dict[str, Any]]:
    ytd: dict[tuple[int, int], dict[str, int]] = {}
    for path in sorted((RAW_DIR / "philippines").glob("*.xlsx")):
        match = re.search(r"_(\d{4})_(\d{2})\.xlsx$", path.name)
        if not match:
            continue
        year, month = map(int, match.groups())
        workbook = load_workbook(path, data_only=True, read_only=True)
        worksheet = (
            workbook["Tab 11"]
            if "Tab 11" in workbook.sheetnames
            else workbook.worksheets[0]
        )
        row = find_uae_row(worksheet)
        ytd[(year, month)] = {
            "total": numeric(row[3]),
            "new_hires": numeric(row[7]),
            "rehires": numeric(row[11]),
        }

    monthly: dict[str, dict[str, Any]] = {}
    previous_by_year: dict[int, dict[str, int]] = defaultdict(
        lambda: {"total": 0, "new_hires": 0, "rehires": 0}
    )
    for (year, month), values in sorted(ytd.items()):
        previous = previous_by_year[year]
        deltas = {key: values[key] - previous[key] for key in values}
        if any(value < 0 for value in deltas.values()):
            monthly[f"{year}-{month:02d}"] = {
                "total": None,
                "new_hires": None,
                "rehires": None,
                "quality_flag": (
                    "DMW官方YTD较上月回落，无法可靠差分"
                    f"（总计 {previous['total']:,}→{values['total']:,}）"
                ),
            }
        elif deltas["total"] != deltas["new_hires"] + deltas["rehires"]:
            raise RuntimeError(
                f"DMW components do not add at {year}-{month:02d}: {deltas}"
            )
        else:
            monthly[f"{year}-{month:02d}"] = {
                **deltas,
                "quality_flag": None,
            }
        previous_by_year[year] = values
    return monthly


def parse_nepal() -> dict[str, dict[str, Any]]:
    monthly: dict[str, dict[str, Any]] = {}
    date_pattern = re.compile(
        r"from\s+(\d{4})-(\d{2})-(\d{2})\s+to\s+"
        r"(\d{4})-(\d{2})-(\d{2})",
        flags=re.IGNORECASE,
    )
    for path in sorted((RAW_DIR / "nepal").glob("*.pdf")):
        reader = PdfReader(path)
        page_texts = [page.extract_text() or "" for page in reader.pages]
        if not any(text.strip() for text in page_texts):
            raise RuntimeError(f"No extractable text in DoFE PDF: {path}")
        text = "\n".join(page_texts)
        dates = date_pattern.search(text)
        if not dates:
            raise RuntimeError(f"Report date range not found: {path}")
        end = date(*map(int, dates.groups()[3:]))
        uae_line = next(
            (
                line.strip()
                for line in text.splitlines()
                if re.match(
                    r"^\d+\s+UAE\s+", line.strip(), flags=re.IGNORECASE
                )
            ),
            None,
        )
        if not uae_line:
            raise RuntimeError(f"UAE row not found: {path}")
        values = [int(value) for value in re.findall(r"\d+", uae_line)][1:]
        if len(values) != 21:
            raise RuntimeError(
                f"Expected 21 DoFE measures in {path.name}, found {len(values)}"
            )
        with_reentry = values[17]
        without_reentry = values[20]
        if with_reentry < without_reentry:
            raise RuntimeError(f"Invalid DoFE re-entry totals in {path.name}")
        key = f"{end.year}-{end.month:02d}"
        if key in monthly:
            raise RuntimeError(f"Multiple DoFE periods assigned to {key}")
        monthly[key] = {
            "with_reentry": with_reentry,
            "without_reentry": without_reentry,
            "period_end": end.isoformat(),
        }
    return monthly


def parse_bangladesh() -> dict[str, int]:
    monthly: dict[str, int] = {}
    for path in sorted((RAW_DIR / "bangladesh").glob("*.json")):
        match = re.search(r"_(\d{4})_(\d{2})\.json$", path.name)
        if not match:
            continue
        year, month = map(int, match.groups())
        payload = json.loads(path.read_text(encoding="utf-8")).get(
            "payload", {}
        )
        rows = payload.get("data", [])
        if len(rows) != 1:
            raise RuntimeError(
                f"Expected one UAE OEP row in {path.name}, got {len(rows)}"
            )
        row = rows[0]
        if "EMIRATES" not in str(row.get("country_name", "")).upper():
            raise RuntimeError(f"Unexpected OEP country in {path.name}: {row}")
        value = numeric(row.get("total_employee"))
        if numeric(payload.get("totalEmployee")) != value:
            raise RuntimeError(f"OEP total mismatch in {path.name}")
        monthly[f"{year}-{month:02d}"] = value
    return monthly


def build_rows() -> list[dict[str, Any]]:
    philippines = parse_philippines()
    nepal = parse_nepal()
    bangladesh = parse_bangladesh()
    months = sorted(set(philippines) | set(nepal) | set(bangladesh))
    if not months:
        raise RuntimeError(
            "No processed observations found; run download first"
        )

    overlap = [
        month
        for month in months
        if BASELINE_START <= month <= BASELINE_END
        and month in philippines
        and month in nepal
        and month in bangladesh
    ]
    if len(overlap) < 6:
        raise RuntimeError(f"Insufficient common baseline months: {overlap}")
    baselines = {
        "philippines": mean(philippines[month]["total"] for month in overlap),
        "nepal": mean(nepal[month]["with_reentry"] for month in overlap),
        "bangladesh": mean(bangladesh[month] for month in overlap),
    }

    rows: list[dict[str, Any]] = []
    for month in months:
        year, month_number = map(int, month.split("-"))
        ph = philippines.get(month)
        np = nepal.get(month)
        bd = bangladesh.get(month)
        philippines_available = ph is not None and ph["total"] is not None
        available = sum((philippines_available, np is not None, bd is not None))
        all_three = available == 3
        proxy_sum = ph["total"] + np["with_reentry"] + bd if all_three else None
        proxy_index = None
        if all_three:
            proxy_index = 100 * mean(
                (
                    ph["total"] / baselines["philippines"],
                    np["with_reentry"] / baselines["nepal"],
                    bd / baselines["bangladesh"],
                )
            )
        rows.append(
            {
                "date": month_end(year, month_number),
                "philippines_total": ph["total"] if ph else None,
                "philippines_new_hires": ph["new_hires"] if ph else None,
                "philippines_rehires": ph["rehires"] if ph else None,
                "nepal_with_reentry": np["with_reentry"] if np else None,
                "nepal_without_reentry": np["without_reentry"] if np else None,
                "bangladesh_clearance": bd,
                "proxy_sum": proxy_sum,
                "proxy_index": (
                    round(proxy_index, 6) if proxy_index is not None else None
                ),
                "source_count": available,
                "quality_flag": ph["quality_flag"] if ph else None,
            }
        )
    return list(reversed(rows))


# ---------------------------------------------------------------------------
# 入库与写表
# ---------------------------------------------------------------------------


@contextmanager
def _transaction(con) -> Iterator[None]:
    transaction = con.begin()
    try:
        yield
    except BaseException:
        transaction.rollback()
        raise
    else:
        transaction.commit()


def _rows_to_long(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """宽行 → (period, series, value[, quality_flag]) 长行，仅存非空值。

    ``quality_flag`` 单独成行（series='quality_flag'，value 为空），因为
    质量标记出现的月份（如菲律宾 YTD 回落）恰好没有 philippines_total 值。
    """

    long_rows: list[dict[str, Any]] = []
    for row in rows:
        period = row["date"]
        for series in SERIES_KEYS:
            value = row.get(series)
            if value is None:
                continue
            long_rows.append(
                {
                    "period": period,
                    "series": series,
                    "value": value,
                    "quality_flag": None,
                }
            )
        if row.get("quality_flag"):
            long_rows.append(
                {
                    "period": period,
                    "series": "quality_flag",
                    "value": None,
                    "quality_flag": row["quality_flag"],
                }
            )
    return long_rows


def update(con, *, force: bool = False, skip_download: bool = False) -> dict:
    """下载（按需）→ 解析 → 事务内入库 foreign_labour_monthly。"""

    if not skip_download:
        download_all(refresh=force)
    rows = build_rows()
    long_rows = _rows_to_long(rows)
    with _transaction(con):
        db.replace(con, "foreign_labour_monthly", long_rows)
    note = f"{len(long_rows)} 条序列观测（{rows[-1]['date']} 至 {rows[0]['date']}）"
    return {"status": "ok", "rows": len(long_rows), "note": note}


def _rows_from_db(con) -> list[dict[str, Any]]:
    """长表 → 宽行（旧 process_data.py 的 rows 结构），最新在前。"""

    records = con.execute(
        "SELECT period, series, value, quality_flag "
        "FROM foreign_labour_monthly"
    ).fetchall()
    by_period: dict[date, dict[str, Any]] = defaultdict(dict)
    for period, series, value, quality_flag in records:
        if series == "quality_flag":
            by_period[period]["quality_flag"] = quality_flag
            continue
        by_period[period][series] = value
    rows = []
    for period in sorted(by_period, reverse=True):
        values = by_period[period]
        rows.append(
            {
                "date": period,
                "philippines_total": values.get("philippines_total"),
                "philippines_new_hires": values.get("philippines_new_hires"),
                "philippines_rehires": values.get("philippines_rehires"),
                "nepal_with_reentry": values.get("nepal_with_reentry"),
                "nepal_without_reentry": values.get("nepal_without_reentry"),
                "bangladesh_clearance": values.get("bangladesh_clearance"),
                "proxy_sum": values.get("proxy_sum"),
                "proxy_index": values.get("proxy_index"),
                "source_count": values.get("source_count"),
                "quality_flag": values.get("quality_flag"),
            }
        )
    return rows


def style_sheet(worksheet, rows: list[dict[str, Any]]) -> None:
    source_fill = PatternFill("solid", fgColor="1F4E78")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    metadata_fill = PatternFill("solid", fgColor="E2F0D9")
    white_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)

    headers = list(WORKBOOK_HEADERS)
    sources = [
        "各国官方劳工输出登记/部署/审批数据",
        "菲律宾 DMW Monthly Compendium Tab 11",
        "菲律宾 DMW Monthly Compendium Tab 11",
        "菲律宾 DMW Monthly Compendium Tab 11",
        "尼泊尔 DoFE monthly final labour approval",
        "尼泊尔 DoFE monthly final labour approval",
        "孟加拉国 BMET/OEP Country Clearance",
        "上述三国官方数据之和，仅三国均有值时计算",
    ]
    updated = max(row["date"] for row in rows)

    worksheet.append([sources[0]] + [None] * (len(headers) - 1))
    worksheet.append(["指标名称"] + headers[1:])
    worksheet.append(["频率"] + ["月"] * (len(headers) - 1))
    worksheet.append(["单位"] + ["人"] * (len(headers) - 1))
    worksheet.append(["来源"] + sources[1:])
    worksheet.append(["更新时间"] + [updated] * (len(headers) - 1))
    for row in rows:
        worksheet.append(
            [
                row["date"],
                row["philippines_total"],
                row["philippines_new_hires"],
                row["philippines_rehires"],
                row["nepal_with_reentry"],
                row["nepal_without_reentry"],
                row["bangladesh_clearance"],
                row["proxy_sum"],
            ]
        )

    worksheet.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=len(headers)
    )
    for cell in worksheet[1]:
        cell.fill = source_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for cell in worksheet[2]:
        cell.fill = header_fill
        cell.font = bold_font
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    for row_number in range(3, 7):
        for cell in worksheet[row_number]:
            cell.fill = metadata_fill
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
    for cell in worksheet[6]:
        cell.font = bold_font
    for cell in worksheet["A"][6:]:
        cell.number_format = "yyyy-mm-dd"
    for row_number in range(7, worksheet.max_row + 1):
        for column in range(2, 9):
            worksheet.cell(row_number, column).number_format = "#,##0"
    worksheet.freeze_panes = "B7"
    worksheet.auto_filter.ref = (
        f"A2:{get_column_letter(len(headers))}{worksheet.max_row}"
    )
    widths = [13, 22, 21, 21, 25, 27, 23, 22]
    for index, width in enumerate(widths, 1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.row_dimensions[1].height = 24
    worksheet.row_dimensions[2].height = 42


def workbook_is_open(path: Path) -> bool:
    return (path.parent / f"~${path.name}").exists()


def write_workbook(
    rows: list[dict[str, Any]], workbook_path: Path, allow_open: bool
) -> tuple[Path, Path | None]:
    output_path = workbook_path
    if not workbook_path.exists():
        workbook = Workbook()
        workbook.remove(workbook.active)
        backup = None
    else:
        if workbook_is_open(workbook_path) and not allow_open:
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_path = RAW_DIR / (
                f"{workbook_path.stem}_with_foreign_labour_{stamp}.xlsx"
            )
            backup = None
        else:
            BACKUP_DIR.mkdir(parents=True, exist_ok=True)
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup = BACKUP_DIR / (
                f"{workbook_path.stem}_before_foreign_labour_{stamp}.xlsx"
            )
            shutil.copy2(workbook_path, backup)
        workbook = load_workbook(workbook_path, keep_links=True)

    if SHEET_NAME in workbook.sheetnames:
        index = workbook.sheetnames.index(SHEET_NAME)
        workbook.remove(workbook[SHEET_NAME])
        worksheet = workbook.create_sheet(SHEET_NAME, index)
    else:
        worksheet = workbook.create_sheet(SHEET_NAME)
    style_sheet(worksheet, rows)

    temporary = output_path.with_suffix(".xlsx.tmp")
    workbook.save(temporary)
    # openpyxl keeps the source package open on Windows until the workbook is
    # explicitly closed.  Close it before replacing the canonical workbook so
    # the following GFS OOXML merge cannot receive WinError 5.
    workbook.close()
    os.replace(temporary, output_path)
    return output_path, backup


def validate(rows: list[dict[str, Any]], workbook_path: Path | None) -> None:
    dates = [row["date"] for row in rows]
    if dates != sorted(dates, reverse=True) or len(dates) != len(set(dates)):
        raise RuntimeError("Dates are not unique and newest-first")
    for row in rows:
        ph_total = row["philippines_total"]
        if ph_total is not None and ph_total != (
            row["philippines_new_hires"] + row["philippines_rehires"]
        ):
            raise RuntimeError(f"Philippines component mismatch: {row}")
        if row["proxy_sum"] is not None and row["source_count"] != 3:
            raise RuntimeError(f"Proxy emitted with partial coverage: {row}")
    if workbook_path is not None:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        if SHEET_NAME not in workbook.sheetnames:
            raise RuntimeError(f"Missing sheet {SHEET_NAME}")
        worksheet = workbook[SHEET_NAME]
        if worksheet.max_row != len(rows) + 6 or worksheet.max_column != 8:
            raise RuntimeError(
                f"Unexpected sheet shape: {worksheet.max_row}x{worksheet.max_column}"
            )


def merge(workbook_path: Path) -> dict:
    """把 foreign_labour_monthly 写回 月度_外籍劳动力 sheet。"""

    workbook_path = workbook_path.resolve()
    con = db.connect(read_only=True)
    try:
        rows = _rows_from_db(con)
    finally:
        con.close()
    if not rows:
        raise ValueError("foreign_labour_monthly is empty; nothing to merge")

    output_path, backup = write_workbook(rows, workbook_path, allow_open=False)
    validate(rows, output_path)
    note = (
        f"{len(rows)} 个月写入 {output_path.name}"
        + (f"（备份 {backup.name}）" if backup else "")
    )
    return {"status": "ok", "note": note}


if __name__ == "__main__":
    print("source_foreign_labour: 模块入口由 htfa.jobs.uae_data.update_data 与 htfa.jobs.uae_data.merge_workbook 调用")
