"""Shared DuckDB-to-workbook writer for source-owned indicator sheets."""

from __future__ import annotations

import os
import shutil
import tempfile
from calendar import monthrange
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db


def merge_source_columns(workbook_path: Path, spec_builder) -> dict[str, Any]:
    """Patch source-owned columns from DuckDB, backing up and atomically replacing Excel."""
    with db.connect(read_only=True) as con:
        spec = spec_builder(con)
        metadata = read_indicator_metadata(con, spec["indicators"])
    observations = spec["observations"]
    if not observations:
        raise ValueError("来源数据表为空")
    frequencies = {item["frequency"] for item in metadata.values()}
    if len(frequencies) != 1:
        raise ValueError("同一列更新任务的指标频率不一致")
    indicators = {name: (item["unit"], item["source"])
                  for name, item in metadata.items()}
    path = Path(workbook_path).resolve()
    if (path.parent / f"~${path.name}").exists():
        raise RuntimeError("请先关闭 Excel 工作簿")
    workbook = load_workbook(path, keep_links=True)
    temporary = None
    try:
        update_indicator_columns(
            workbook[spec["name"]], observations, indicators, frequencies.pop(),
        )
        backups = path.parent / "backups"
        backups.mkdir(exist_ok=True)
        shutil.copy2(path, backups / f"columns_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx")
        fd, name = tempfile.mkstemp(prefix=".source-columns-", suffix=".tmp.xlsx", dir=path.parent)
        os.close(fd)
        temporary = Path(name)
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return {"status": "ok", "rows": len(observations),
            "note": f"{spec['name']} 指定列已从 DuckDB 同步，保留其他指标与历史日期"}


def month_end(period: date) -> date:
    """Normalize a database observation's month to its calendar month end."""
    return date(period.year, period.month, monthrange(period.year, period.month)[1])


def pivot_indicator_rows(records, name_builder, *, period_transform=None):
    """Pivot long records using their period, dimensions and final value.

    ``name_builder`` maps the dimension columns to a workbook indicator name.
    ``period_transform`` optionally normalizes each observation period.
    """
    periods = defaultdict(dict)
    for record in records:
        period = period_transform(record[0]) if period_transform else record[0]
        periods[period][name_builder(*record[1:-1])] = record[-1]
    return list(periods.items())


def merge_source_sheet(workbook_path: Path, spec_builder) -> dict[str, Any]:
    """Read a source-owned sheet specification and metadata from DuckDB.

    ``spec_builder`` takes the read-only connection and returns a mapping with
    the sheet's name, title, indicators and wide rows. ``workbook_path`` is
    written atomically by the shared workbook writer.
    """
    with db.connect(read_only=True) as con:
        spec = spec_builder(con)
        spec["metadata"] = read_indicator_metadata(con, spec["indicators"])
    counts = write_indicator_sheets(workbook_path, [spec])
    return {
        "status": "ok",
        "rows": counts[spec["name"]],
        "note": f"{spec['name']} {counts[spec['name']]} 行",
    }


def read_indicator_metadata(con, names: list[str]) -> dict[str, dict[str, Any]]:
    """Read metadata from ``con`` for all indicator ``names``, rejecting missing entries."""

    records = con.execute(
        "SELECT indicator_name, frequency, unit, source "
        "FROM meta_indicator_dictionary"
    ).fetchall()
    by_name = {
        row[0]: {
            "frequency": row[1],
            "unit": row[2],
            "source": row[3],
        }
        for row in records
    }
    missing = [name for name in names if name not in by_name]
    if missing:
        raise ValueError(f"数据库指标字典缺少：{', '.join(missing)}")
    return {name: by_name[name] for name in names}


def write_indicator_sheets(
    workbook_path: Path,
    sheets: Iterable[Mapping[str, Any]],
) -> dict[str, int]:
    """Atomically replace several wide indicator sheets in one workbook.

    Each sheet mapping contains ``name``, ``title``, ``indicators``,
    ``metadata`` and ``rows``. Rows are ``(period, {indicator: value})``.
    """

    workbook_path = Path(workbook_path).resolve()
    if not workbook_path.is_file():
        raise FileNotFoundError(workbook_path)
    lock_path = workbook_path.parent / f"~${workbook_path.name}"
    if lock_path.exists():
        raise RuntimeError(f"请先关闭 Excel 工作簿：{lock_path}")

    sheet_specs = list(sheets)
    if not sheet_specs:
        return {}

    workbook = load_workbook(workbook_path, keep_links=True)
    counts: dict[str, int] = {}
    try:
        for spec in sheet_specs:
            sheet_name = str(spec["name"])
            indicators = [str(name) for name in spec["indicators"]]
            metadata = spec["metadata"]
            rows = sorted(spec["rows"], key=lambda item: item[0], reverse=True)
            if not rows:
                raise ValueError(f"{sheet_name} 没有可写入的数据")
            missing_metadata = [name for name in indicators if name not in metadata]
            if missing_metadata:
                raise ValueError(
                    f"{sheet_name} 缺少指标元数据：{', '.join(missing_metadata)}"
                )

            index = workbook.sheetnames.index(sheet_name) if sheet_name in workbook.sheetnames else len(workbook.sheetnames)
            if sheet_name in workbook.sheetnames:
                workbook.remove(workbook[sheet_name])
            worksheet = workbook.create_sheet(sheet_name, index)
            _write_sheet(worksheet, str(spec["title"]), indicators, metadata, rows)
            counts[sheet_name] = len(rows)

        temporary = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix=f".{workbook_path.stem}.extended-",
                suffix=".tmp.xlsx",
                dir=workbook_path.parent,
                delete=False,
            ) as handle:
                temporary = Path(handle.name)
            workbook.save(temporary)
            os.replace(temporary, workbook_path)
        finally:
            if temporary is not None and temporary.exists():
                temporary.unlink()
    finally:
        workbook.close()
    return counts


def _write_sheet(
    worksheet,
    title: str,
    indicators: list[str],
    metadata: Mapping[str, Mapping[str, Any]],
    rows: list[tuple[Any, Mapping[str, Any]]],
) -> None:
    column_count = len(indicators) + 1
    worksheet.cell(row=1, column=1, value=title)
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
    worksheet.cell(row=2, column=1, value="指标名称")
    worksheet.cell(row=3, column=1, value="频率")
    worksheet.cell(row=4, column=1, value="单位")
    worksheet.cell(row=5, column=1, value="来源")
    worksheet.cell(row=6, column=1, value="更新时间")

    updated_at = datetime.now()
    for column, name in enumerate(indicators, start=2):
        row_metadata = metadata[name]
        worksheet.cell(row=2, column=column, value=name)
        worksheet.cell(row=3, column=column, value=row_metadata["frequency"])
        worksheet.cell(row=4, column=column, value=row_metadata["unit"])
        worksheet.cell(row=5, column=column, value=row_metadata["source"])
        worksheet.cell(row=6, column=column, value=updated_at)

    for row_number, (period, values) in enumerate(rows, start=7):
        worksheet.cell(row=row_number, column=1, value=period)
        for column, name in enumerate(indicators, start=2):
            worksheet.cell(
                row=row_number,
                column=column,
                value=_excel_value(values.get(name)),
            )

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    metadata_fill = PatternFill("solid", fgColor="E2F0D9")
    title_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    header_font = Font(name="Microsoft YaHei", size=9, bold=True)
    body_font = Font(name="Microsoft YaHei", size=9)
    worksheet.cell(row=1, column=1).fill = PatternFill("solid", fgColor="1F4E78")
    worksheet.cell(row=1, column=1).font = title_font
    worksheet.cell(row=1, column=1).alignment = Alignment(horizontal="center")
    for row_number in range(2, 7):
        for cell in worksheet[row_number][:column_count]:
            cell.font = header_font
            cell.fill = header_fill if row_number == 2 else metadata_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_number in range(7, worksheet.max_row + 1):
        for cell in worksheet[row_number][:column_count]:
            cell.font = body_font
        worksheet.cell(row=row_number, column=1).number_format = "yyyy-mm-dd"

    worksheet.column_dimensions["A"].width = 13
    for column in range(2, column_count + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 28
    worksheet.freeze_panes = "B7"
    worksheet.auto_filter.ref = f"A2:{get_column_letter(column_count)}{worksheet.max_row}"


def _excel_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (date, datetime, int, float, str)) or value is None:
        return value
    return float(value)


def update_indicator_columns(
    sheet, observations: list[tuple], indicators: dict, frequency: str,
) -> None:
    """Patch DB-backed columns while preserving unrelated data and dates."""
    width, height = sheet.max_column, sheet.max_row
    columns = {sheet.cell(2, i).value: i for i in range(2, width + 1)}
    missing = set(indicators) - set(columns)
    if missing:
        raise ValueError(f"Legacy sheet lacks matched columns: {missing}")
    data = {sheet.cell(i, 1).value.date(): [sheet.cell(i, j).value
            for j in range(2, width + 1)]
            for i in range(7, height + 1)
            if isinstance(sheet.cell(i, 1).value, datetime)}
    for name in indicators:
        periods = [r[0] for r in observations if r[1] == name]
        if not periods:
            raise ValueError(f"No DB observations for {name}")
        first, last = min(periods), max(periods)
        for stamp, values in data.items():
            if first <= stamp <= last:
                values[columns[name] - 2] = None
        unit, source = indicators[name]
        for row, value in ((3, frequency), (4, unit), (5, source), (6, str(date.today()))):
            sheet.cell(row, columns[name], value)
    for stamp, name, value in observations:
        data.setdefault(stamp, [None] * (width - 1))[columns[name] - 2] = value
    for i, (stamp, values) in enumerate(sorted(data.items(), reverse=True), 7):
        sheet.cell(i, 1, datetime.combine(stamp, datetime.min.time()))
        for j, value in enumerate(values, 2):
            sheet.cell(i, j).value = value
