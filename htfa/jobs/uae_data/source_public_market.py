"""Public DFM, CBUAE and NY Fed daily data; DuckDB precedes workbook writes.

The legacy sheet name is retained for existing consumers. Only the explicitly
matched series are updated; licensed commodity quotes and Term SOFR remain
untouched rather than being silently replaced by different benchmarks.
"""

from __future__ import annotations

import html
import json
import math
import os
import re
import shutil
import tempfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import urljoin

from openpyxl import load_workbook

from . import db
from ._official_download import download_file
from .paths import DATA_DIR

RAW_DIR = DATA_DIR / "raw" / "public_market"
TABLE = "public_market_daily"
SHEET = "日度_Wind"
BASE = "https://www.centralbank.ae"
EIBOR_URL = BASE + "/umbraco/surface/eibor/GetEiborData"
FX_URL = BASE + "/umbraco/Surface/Exchange/GetExchangeRateData?lang=en"
FX_PAGE = BASE + "/en/forex-eibor/exchange-rates/"
DFM_PAGE = "https://www.dfm.ae/the-exchange/statistics-reports/historical-data/dfmgi"
EFFR_BASE = "https://markets.newyorkfed.org/api/rates/unsecured/effr/search.json"
EIBOR_ON = "阿联酋:银行间同业拆借利率(EIBOR):隔夜"
EIBOR_YEAR = "阿联酋:银行间同业拆借利率(EIBOR):1年"
EFFR = "美国:有效联邦基金利率(EFFR)"
USD = "美元兑阿联酋迪拉姆"
CNY = "人民币兑阿联酋迪拉姆"
DFM = "阿联酋DFM综合股票指数"
INDICATORS = {
    EIBOR_ON: ("%", "CBUAE EIBOR"),
    EIBOR_YEAR: ("%", "CBUAE EIBOR"),
    EFFR: ("%", "Federal Reserve Bank of New York"),
    USD: ("AED/USD", "CBUAE VAT reference exchange rates"),
    CNY: ("AED/CNY", "CBUAE VAT reference exchange rates"),
    DFM: ("点", "Dubai Financial Market"),
}
MONTHLY_CREDIT = {
    "阿联酋:银行存款:居民存款:政府部门": "阿联酋政府存款",
    "阿联酋:银行存款:居民存款:政府相关实体(政府持股超50%)": "阿联酋政府控股企业存款",
    "阿联酋:信贷总额:国内信贷:政府部门": "阿联酋政府信贷",
    "阿联酋:信贷总额:国内信贷:公共部门(政府相关实体)": "阿联酋政府控股企业信贷",
}


def _embedded_json(text: str, element_id: str) -> dict:
    match = re.search(
        rf'<script\b[^>]*id=["\']{re.escape(element_id)}["\'][^>]*>(.*?)</script>',
        text, re.S | re.I,
    )
    if match is None:
        raise ValueError(f"Official page is missing {element_id}")
    return json.loads(match.group(1))


def _row(period: date, indicator: str, value: object, url: str) -> dict:
    number = float(value)
    if not math.isfinite(number) or number <= 0:
        raise ValueError(f"Invalid public observation: {indicator}, {period}, {value}")
    if period > date.today():
        raise ValueError(f"Future public observation: {period}")
    return {"period": period, "indicator": indicator, "value": number,
            "source_url": url}


def parse_eibor(text: str) -> list[dict]:
    payload = _embedded_json(text, "eibor-chart-data")
    dates = payload["historicalFullDates"]
    result = []
    for key, indicator in (("historicalOnRates", EIBOR_ON),
                           ("historicalOneYear", EIBOR_YEAR)):
        values = payload[key]
        if len(values) != len(dates):
            raise ValueError("EIBOR date/value lengths differ")
        for stamp, value in zip(dates, values):
            if value is not None:
                result.append(_row(datetime.strptime(stamp, "%d %b %Y").date(),
                                   indicator, value, EIBOR_URL))
    return result


def parse_usd(text: str) -> list[dict]:
    payload = _embedded_json(text, "fxDataJson")["usd"]
    dates, values = payload["rowDates"], payload["values"]
    if len(dates) != len(values):
        raise ValueError("FX date/value lengths differ")
    return [_row(date.fromisoformat(stamp), USD, value, FX_URL)
            for stamp, value in zip(dates, values) if value is not None]


def parse_dfm(payload: dict) -> list[dict]:
    records = payload.get("daily_closing")
    if not isinstance(records, list) or not records:
        raise ValueError("DFM has no daily closing observations")
    return [_row(date.fromisoformat(r["d"][:10]), DFM, r["i"], DFM_PAGE)
            for r in records]


def parse_effr(payload: dict) -> list[dict]:
    records = payload.get("refRates")
    if not isinstance(records, list) or not records:
        raise ValueError("NY Fed has no reference rate observations")
    if any(r.get("type") != "EFFR" for r in records):
        raise ValueError("NY Fed returned a different reference rate")
    return [_row(date.fromisoformat(r["effectiveDate"]), EFFR,
                 r["percentRate"], EFFR_BASE) for r in records]


def _parse_cny_workbook(path: Path, url: str) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    result = []
    try:
        for sheet in workbook:
            for cells in sheet.values:
                if len(cells) < 3 or cells[1] != "Chinese Yuan":
                    continue
                stamp = cells[0]
                if not isinstance(stamp, (date, datetime)):
                    raise ValueError(f"Invalid FX workbook date: {stamp}")
                result.append(_row(stamp.date() if isinstance(stamp, datetime)
                                   else stamp, CNY, cells[2], url))
    finally:
        workbook.close()
    if not result:
        raise ValueError(f"FX workbook has no onshore Chinese Yuan: {path.name}")
    return result


def parse_cny_daily(text: str, period: date, url: str) -> list[dict]:
    for row in re.findall(r"<tr\b.*?</tr>", text, re.S | re.I):
        cells = [html.unescape(re.sub(r"<[^>]+>", "", c)).strip()
                 for c in re.findall(r"<td\b[^>]*>(.*?)</td>", row, re.S | re.I)]
        if len(cells) >= 2 and cells[-2] == "Chinese Yuan":
            return [_row(period, CNY, cells[-1], url)]
    if "Data for this date not available" in text:
        return []
    raise ValueError(f"FX daily response has no onshore Chinese Yuan: {period}")


def _download() -> list[dict]:
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    effr_url = f"{EFFR_BASE}?startDate=2023-01-01&endDate={date.today()}"
    plans = [("eibor.html", EIBOR_URL), ("usd.html", FX_URL),
             ("fx-page.html", FX_PAGE), ("dfm-page.html", DFM_PAGE),
             ("effr.json", effr_url)]

    def fetch(item):
        name, url = item
        download_file(url, RAW_DIR / name, force=True, browser_transport=True)

    with ThreadPoolExecutor(max_workers=5) as pool:
        list(pool.map(fetch, plans))
    # Discover the public API base from the official DFM client config.
    dfm_page = (RAW_DIR / "dfm-page.html").read_text(encoding="utf-8")
    match = re.search(r'apiBaseUrl:"(https://[^"\s]+)"', dfm_page)
    if match is None or not match.group(1).startswith("https://api2.dfm.ae/"):
        raise ValueError("DFM official API base is missing or unexpected")
    from curl_cffi import requests

    response = requests.post(
        match.group(1) + "/data", impersonate="chrome", timeout=60,
        data={"Command": "SearchDFMIndex", "Period": "custom",
              "FromDate": "2023-01-01", "ToDate": str(date.today()),
              "Language": "en", "type": "json"},
    )
    response.raise_for_status()
    payload = response.json()
    if isinstance(payload, str):
        payload = json.loads(payload)
    parse_dfm(payload)
    (RAW_DIR / "dfm.json").write_text(json.dumps(payload), encoding="utf-8")
    page = (RAW_DIR / "fx-page.html").read_text(encoding="utf-8")
    links = _latest_fx_year_links(page)

    def download_fx(link):
        url = urljoin(BASE, link)
        path = RAW_DIR / Path(link).name
        download_file(url, path, force=True, browser_transport=True)
        return _parse_cny_workbook(path, url)

    with ThreadPoolExecutor(max_workers=5) as pool:
        cny = [row for batch in pool.map(download_fx, links) for row in batch]
    latest = max(row["period"] for row in cny)
    days = [latest + timedelta(days=i)
            for i in range(1, (date.today() - latest).days + 1)]

    def fetch_day(period):
        url = (BASE + "/umbraco/Surface/Exchange/GetExchangeRateAllCurrencyDate"
               f"?dateTime={period}&lang=en")
        path = RAW_DIR / f"fx-{period}.html"
        download_file(url, path, force=True, browser_transport=True)
        return parse_cny_daily(path.read_text(encoding="utf-8"), period, url)

    with ThreadPoolExecutor(max_workers=4) as pool:
        cny.extend(row for batch in pool.map(fetch_day, days) for row in batch)
    (RAW_DIR / "cny.json").write_text(json.dumps(cny, default=str), encoding="utf-8")
    return cny


def _latest_fx_year_links(page: str) -> list[str]:
    """Use the latest published FX year, including at a calendar rollover."""
    candidates = {}
    for link in re.findall(r'href=["\']([^"\']+fx_[^"\']+_en\.xlsx)["\']', page, re.I):
        match = re.search(r"fx_[a-z]+(\d{2})_en", link, re.I)
        if match:
            candidates[html.unescape(link)] = 2000 + int(match.group(1))
    if not candidates:
        raise ValueError("CBUAE has no dated FX workbooks")
    latest_year = max(candidates.values())
    return sorted(link for link, year in candidates.items() if year == latest_year)


def update(con, *, force: bool = False, skip_download: bool = False) -> dict:
    """Refresh the six matched public series and upsert them in DuckDB."""
    if not skip_download:
        _download()
    cny = json.loads((RAW_DIR / "cny.json").read_text(encoding="utf-8"))
    for row in cny:
        row["period"] = date.fromisoformat(row["period"])
    rows = (parse_eibor((RAW_DIR / "eibor.html").read_text(encoding="utf-8"))
            + parse_usd((RAW_DIR / "usd.html").read_text(encoding="utf-8"))
            + parse_dfm(json.loads((RAW_DIR / "dfm.json").read_text(encoding="utf-8")))
            + parse_effr(json.loads((RAW_DIR / "effr.json").read_text(encoding="utf-8")))
            + cny)
    keys = [(r["period"], r["indicator"]) for r in rows]
    if len(keys) != len(set(keys)):
        raise ValueError("Public market inputs contain duplicate observation keys")
    if set(r["indicator"] for r in rows) != set(INDICATORS):
        raise ValueError("Public market inputs do not cover all six indicators")
    con.begin()
    try:
        con.execute(f"CREATE TABLE IF NOT EXISTS {TABLE} (period DATE, "
                    "indicator VARCHAR, value DOUBLE, source_url VARCHAR, "
                    "PRIMARY KEY(period, indicator))")
        existing = dict(con.execute(f"SELECT indicator,max(period) FROM {TABLE} GROUP BY 1").fetchall())
        for name, prior in existing.items():
            if max(r["period"] for r in rows if r["indicator"] == name) < prior:
                raise ValueError(f"Public market source coverage regressed: {name}")
        con.executemany(f"INSERT OR REPLACE INTO {TABLE} VALUES (?,?,?,?)",
                        [(r["period"], r["indicator"], r["value"], r["source_url"]) for r in rows])
        db.upsert_dictionary_rows(con, [
            {"indicator_name": name, "frequency": "日", "unit": unit,
             "source": source, "type": "市场", "industry": "金融",
             "updated_at": date.today()} for name, (unit, source) in INDICATORS.items()
        ])
        db.upsert_dictionary_rows(con, [
            {"indicator_name": name, "frequency": "月", "unit": "十亿阿联酋迪拉姆",
             "source": "CBUAE Statistical Bulletin", "type": "金额",
             "industry": "金融", "updated_at": date.today()}
            for name in MONTHLY_CREDIT
        ])
        con.commit()
    except BaseException:
        con.rollback()
        raise
    return {"status": "ok", "rows": len(rows),
            "note": "六个公开源指标已入库：DFM、EIBOR隔夜/1年、EFFR、USD/AED、CNY/AED"}


def merge(workbook_path: Path) -> dict:
    """Update matched columns from DuckDB, preserving all other workbook data."""
    with db.connect(read_only=True) as con:
        observations = con.execute(f"SELECT period,indicator,value FROM {TABLE} ORDER BY period").fetchall()
        credit = []
        if db.table_exists(con, "cbuae_monthly"):
            for name, original in MONTHLY_CREDIT.items():
                credit.extend((period, name, float(value) / 1000 if value is not None else None)
                              for period, value in con.execute(
                                  "SELECT period,value FROM cbuae_monthly WHERE indicator=?",
                                  [original],
                              ).fetchall())
    if not observations:
        raise ValueError("Public market table is empty")
    path = Path(workbook_path)
    if (path.parent / f"~${path.name}").exists():
        raise RuntimeError("请先关闭 Excel 工作簿")
    workbook = load_workbook(path, keep_links=True)
    temporary = None
    try:
        _update_columns(workbook[SHEET], observations, INDICATORS, "日")
        if credit:
            _update_columns(workbook["月度_Wind"], credit,
                            {name: ("十亿阿联酋迪拉姆", "CBUAE Statistical Bulletin")
                             for name in MONTHLY_CREDIT}, "月")
        backup_dir = DATA_DIR / "backups"
        backup_dir.mkdir(exist_ok=True)
        shutil.copy2(path, backup_dir / f"public_market_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
        fd, name = tempfile.mkstemp(prefix=".public-market-", suffix=".tmp.xlsx", dir=path.parent)
        os.close(fd)
        temporary = Path(name)
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return {"status": "ok", "rows": len(observations),
            "note": "Wind 表名保留兼容；六个日度公开源、四个央行月度存贷款列已更新"}


def _update_columns(sheet, observations: list[tuple], indicators: dict,
                    frequency: str) -> None:
    """Patch DB-backed columns while preserving unrelated data and dates."""
    width, height = sheet.max_column, sheet.max_row
    columns = {sheet.cell(2, i).value: i for i in range(2, width + 1)}
    missing = set(indicators) - set(columns)
    if missing:
        raise ValueError(f"Legacy sheet lacks matched columns: {missing}")
    data = {sheet.cell(i, 1).value.date(): [sheet.cell(i, j).value
            for j in range(2, width + 1)]
            for i in range(7, height + 1)
            if isinstance(sheet.cell(i, 1).value, datetime)}
    for name in indicators:
        periods = [r[0] for r in observations if r[1] == name]
        if not periods:
            raise ValueError(f"No DB observations for {name}")
        first, last = min(periods), max(periods)
        for stamp, values in data.items():
            if first <= stamp <= last:
                values[columns[name] - 2] = None
        unit, source = indicators[name]
        for row, value in ((3, frequency), (4, unit), (5, source), (6, str(date.today()))):
            sheet.cell(row, columns[name], value)
    for stamp, name, value in observations:
        data.setdefault(stamp, [None] * (width - 1))[columns[name] - 2] = value
    for i, (stamp, values) in enumerate(sorted(data.items(), reverse=True), 7):
        sheet.cell(i, 1, datetime.combine(stamp, datetime.min.time()))
        for j, value in enumerate(values, 2):
            sheet.cell(i, j).value = value
