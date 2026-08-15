"""Validate UAE headline PMI data and update ``data/阿联酋.xlsx``.

The managed worksheet is named ``月度_LSEG`` for compatibility with the
requested workbook schema.  Cell-level provenance identifies the actual public
source and does not represent the observations as an LSEG feed.
"""

from __future__ import annotations

import argparse
import sys
import csv
import json
import math
import re
import shutil
import tempfile
import warnings
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[3]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.data_sources._excel_helpers import (
    payload_json_file,
    run_powershell_sheet_writer,
)


TARGET_SHEET = "月度_LSEG"
DICTIONARY_SHEET = "指标字典"
INDICATOR_NAME = "阿联酋非油私营部门采购经理人指数(PMI)"
SOURCE_LABEL = "S&P Global / Trading Economics（公开样本）"
SOURCE_URL = "https://tradingeconomics.com/united-arab-emirates/manufacturing-pmi"
MINIMUM_HISTORY_MONTHS = 24
MAX_LATEST_LAG_MONTHS = 2


@dataclass(frozen=True)
class PmiObservation:
    """One reference month and its seasonally adjusted headline PMI value."""

    period: str
    value: float


def _month_number(period: str) -> int:
    parsed = datetime.strptime(period, "%Y-%m")
    return parsed.year * 12 + parsed.month


def _expected_periods(start: str, end: str) -> list[str]:
    current = datetime.strptime(start, "%Y-%m")
    end_date = datetime.strptime(end, "%Y-%m")
    periods: list[str] = []
    while current <= end_date:
        periods.append(current.strftime("%Y-%m"))
        year = current.year + (1 if current.month == 12 else 0)
        month = 1 if current.month == 12 else current.month + 1
        current = current.replace(year=year, month=month)
    return periods


def extract_series(payload: Any) -> dict[str, Any]:
    """Extract the sole public UAE PMI series from the raw chart payload."""

    try:
        return payload[0]["series"][0]["serie"]
    except (IndexError, KeyError, TypeError) as exc:
        raise ValueError("Raw file does not contain a UAE PMI chart series") from exc


def load_observations(path: Path) -> list[PmiObservation]:
    """Load, normalize, and validate observation-level values."""

    payload = json.loads(path.read_text(encoding="utf-8"))
    series = extract_series(payload)
    if series.get("country") != "United Arab Emirates":
        raise ValueError(f"Unexpected country: {series.get('country')!r}")
    if str(series.get("frequency", "")).lower() != "monthly":
        raise ValueError(f"Unexpected frequency: {series.get('frequency')!r}")
    if "S&P" not in str(series.get("source", "")):
        raise ValueError(f"Unexpected source: {series.get('source')!r}")

    observations: list[PmiObservation] = []
    seen: set[str] = set()
    for index, item in enumerate(series.get("data", []), start=1):
        if not isinstance(item, list) or len(item) < 4:
            raise ValueError(f"Malformed observation at source row {index}")
        raw_value, raw_date = item[0], item[3]
        try:
            reference_date = datetime.strptime(str(raw_date), "%Y-%m-%d")
        except ValueError as exc:
            raise ValueError(f"Invalid reference date: {raw_date!r}") from exc
        if reference_date.day != 1:
            raise ValueError(f"Reference date is not month-start: {raw_date!r}")
        period = reference_date.strftime("%Y-%m")
        if period in seen:
            raise ValueError(f"Duplicate UAE PMI period: {period}")
        seen.add(period)
        if isinstance(raw_value, bool):
            raise ValueError(f"Invalid UAE PMI value in {period}: {raw_value!r}")
        try:
            value = float(raw_value)
        except (TypeError, ValueError) as exc:
            raise ValueError(
                f"Invalid UAE PMI value in {period}: {raw_value!r}"
            ) from exc
        if not math.isfinite(value) or not 0 <= value <= 100:
            raise ValueError(f"Out-of-range UAE PMI value in {period}: {value}")
        observations.append(PmiObservation(period=period, value=value))
    if not observations:
        raise ValueError("Raw UAE PMI series contains no observations")
    return sorted(observations, key=lambda item: item.period)


def build_quality_report(
    observations: list[PmiObservation],
    metadata: dict[str, Any],
    *,
    today: date | None = None,
) -> dict[str, Any]:
    """Build reproducible history, continuity, range, and freshness checks."""

    as_of = today or date.today()
    periods = [item.period for item in observations]
    expected = _expected_periods(periods[0], periods[-1])
    missing = sorted(set(expected) - set(periods))
    latest_lag = as_of.year * 12 + as_of.month - _month_number(periods[-1])
    source_update = str(metadata.get("source_last_update") or "")
    source_update_period = None
    if re_match := re.fullmatch(r"(\d{4})(\d{2})\d{2}\d*", source_update):
        source_update_period = f"{re_match.group(1)}-{re_match.group(2)}"
    checks = [
        {
            "name": "has_historical_data",
            "passed": len(observations) >= MINIMUM_HISTORY_MONTHS,
            "detail": (
                f"{len(observations)} months available; minimum "
                f"{MINIMUM_HISTORY_MONTHS}"
            ),
        },
        {
            "name": "monthly_continuity",
            "passed": not missing,
            "detail": "no missing months" if not missing else ", ".join(missing),
        },
        {
            "name": "unique_periods",
            "passed": len(periods) == len(set(periods)),
            "detail": f"{len(periods)} unique periods",
        },
        {
            "name": "valid_pmi_range",
            "passed": all(0 <= item.value <= 100 for item in observations),
            "detail": "all values are within 0-100",
        },
        {
            "name": "latest_data_is_fresh",
            "passed": 0 <= latest_lag <= MAX_LATEST_LAG_MONTHS,
            "detail": (
                f"latest reference month is {periods[-1]}; "
                f"lag={latest_lag} months"
            ),
        },
        {
            "name": "latest_matches_source_update",
            "passed": source_update_period in (None, periods[-1]),
            "detail": (
                "source update date unavailable"
                if source_update_period is None
                else f"source update period={source_update_period}"
            ),
        },
        {
            "name": "metadata_count_matches",
            "passed": metadata.get("observation_count") in (None, len(observations)),
            "detail": (
                f"metadata={metadata.get('observation_count')}; "
                f"parsed={len(observations)}"
            ),
        },
    ]
    return {
        "generated_at": datetime.now().astimezone().isoformat(),
        "status": "passed" if all(item["passed"] for item in checks) else "failed",
        "indicator": INDICATOR_NAME,
        "source": SOURCE_LABEL,
        "coverage": {
            "start_period": periods[0],
            "end_period": periods[-1],
            "observation_count": len(observations),
            "missing_periods": missing,
            "public_sample_is_full_survey_history": False,
        },
        "latest": {
            "period": observations[-1].period,
            "value": observations[-1].value,
        },
        "checks": checks,
        "limitation": metadata.get("coverage_note"),
    }


def _atomic_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            prefix=f".{path.name}.",
            suffix=".tmp",
            dir=path.parent,
            encoding="utf-8",
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
            json.dump(payload, temporary, ensure_ascii=False, indent=2)
            temporary.write("\n")
        temporary_path.replace(path)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def write_csv(path: Path, observations: Iterable[PmiObservation]) -> None:
    """Write the normalized monthly series atomically."""

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            prefix=f".{path.name}.",
            suffix=".tmp",
            dir=path.parent,
            encoding="utf-8-sig",
            newline="",
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
            writer = csv.writer(temporary)
            writer.writerow(["period", "pmi", "source", "source_url"])
            for item in observations:
                writer.writerow(
                    [item.period, f"{item.value:.1f}", SOURCE_LABEL, SOURCE_URL]
                )
        temporary_path.replace(path)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def target_is_current(path: Path, observations: Iterable[PmiObservation]) -> bool:
    """Return whether the target sheet already exactly matches the input data."""

    if not path.is_file():
        return False
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if TARGET_SHEET not in workbook.sheetnames:
            return False
        sheet = workbook[TARGET_SHEET]
        if sheet.cell(2, 2).value != INDICATOR_NAME:
            return False
        actual: dict[str, float] = {}
        for raw_date, raw_value in sheet.iter_rows(
            min_row=7,
            max_col=2,
            values_only=True,
        ):
            if not isinstance(raw_date, (date, datetime)):
                continue
            try:
                actual[raw_date.strftime("%Y-%m")] = float(raw_value)
            except (TypeError, ValueError):
                return False
        expected = {item.period: item.value for item in observations}
        return actual == expected
    finally:
        workbook.close()


def _verify_workbook(path: Path, observations: list[PmiObservation]) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if TARGET_SHEET not in workbook.sheetnames:
            raise ValueError(f"Worksheet was not created: {TARGET_SHEET}")
        sheet = workbook[TARGET_SHEET]
        expected_rows = sorted(observations, key=lambda item: item.period, reverse=True)
        if sheet.cell(2, 2).value != INDICATOR_NAME:
            raise ValueError("Managed PMI indicator header is incorrect")
        if sheet.cell(5, 2).value != SOURCE_LABEL:
            raise ValueError("Managed PMI source metadata is incorrect")
        for row_number, expected in enumerate(expected_rows, start=7):
            raw_date = sheet.cell(row_number, 1).value
            raw_value = sheet.cell(row_number, 2).value
            if not isinstance(raw_date, (date, datetime)):
                raise ValueError(f"Invalid workbook date in row {row_number}")
            if raw_date.strftime("%Y-%m") != expected.period:
                raise ValueError(f"Workbook period mismatch in row {row_number}")
            if float(raw_value) != expected.value:
                raise ValueError(f"Workbook value mismatch in row {row_number}")
        dictionary = workbook[DICTIONARY_SHEET]
        matches = [
            row[0]
            for row in dictionary.iter_rows(min_row=2, max_col=1, values_only=True)
            if row[0] == INDICATOR_NAME
        ]
        if len(matches) != 1:
            raise ValueError(
                "Indicator dictionary should contain exactly one PMI row; "
                f"found {len(matches)}"
            )
    finally:
        workbook.close()


def write_workbook(path: Path, observations: list[PmiObservation]) -> None:
    """Update the managed sheet through Excel COM with rollback and verification."""

    if not path.is_file():
        raise FileNotFoundError(f"Destination workbook not found: {path}")
    helper_path = Path(__file__).with_name("write_uae_pmi_sheet.ps1")
    if not helper_path.is_file():
        raise FileNotFoundError(f"Excel writer helper not found: {helper_path}")
    payload = {
        "source_label": SOURCE_LABEL,
        "metadata_labels": ["指标名称", "频率", "单位", "来源", "更新时间"],
        "frequency": "月",
        "unit": "点",
        "source": SOURCE_LABEL,
        "updated_at": date.today().isoformat(),
        "indicator": {
            "name": INDICATOR_NAME,
            "type": "指数",
            "industry": "宏观",
        },
        "dictionary_sheet_name": DICTIONARY_SHEET,
        "records": [
            {"period": item.period, "value": item.value}
            for item in sorted(observations, key=lambda item: item.period, reverse=True)
        ],
    }
    backup_path: Path | None = None
    with payload_json_file(
        payload,
        prefix=".uae-pmi-sheet-",
        directory=Path(__file__).resolve().parent,
    ) as payload_path:
        with tempfile.NamedTemporaryFile(
            prefix=".uae-pmi-backup-",
            suffix=path.suffix,
            dir=path.parent,
            delete=False,
        ) as backup:
            backup_path = Path(backup.name)
        try:
            shutil.copy2(path, backup_path)
            run_powershell_sheet_writer(
                helper_path, path, payload_path, TARGET_SHEET
            )
            _verify_workbook(path, observations)
        except Exception:
            if backup_path is not None and backup_path.is_file():
                shutil.copy2(backup_path, path)
            raise
        finally:
            if backup_path is not None and backup_path.exists():
                backup_path.unlink()


def parse_args() -> argparse.Namespace:
    repo_root = Path(__file__).resolve().parents[3]
    data_dir = repo_root / "data" / "UAE_PMI"
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--raw",
        type=Path,
        default=data_dir / "raw" / "tradingeconomics_chart.json",
    )
    parser.add_argument(
        "--metadata",
        type=Path,
        default=data_dir / "raw" / "source_metadata.json",
    )
    parser.add_argument(
        "--csv",
        type=Path,
        default=data_dir / "uae_pmi_monthly.csv",
    )
    parser.add_argument(
        "--quality-report",
        type=Path,
        default=data_dir / "quality_report.json",
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=repo_root / "data" / "阿联酋.xlsx",
    )
    parser.add_argument(
        "--skip-workbook",
        action="store_true",
        help="Validate and write local outputs without changing the workbook.",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Rewrite the managed worksheet even if values already match.",
    )
    return parser.parse_args()


def main() -> int:
    warnings.filterwarnings(
        "ignore",
        message="Unknown extension is not supported and will be removed",
        module="openpyxl.worksheet._reader",
    )
    args = parse_args()
    raw_path = args.raw.resolve()
    metadata_path = args.metadata.resolve()
    if not raw_path.is_file():
        raise FileNotFoundError(f"Raw UAE PMI file not found: {raw_path}")
    if not metadata_path.is_file():
        raise FileNotFoundError(f"Source metadata file not found: {metadata_path}")
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    observations = load_observations(raw_path)
    report = build_quality_report(observations, metadata)
    _atomic_json(args.quality_report.resolve(), report)
    if report["status"] != "passed":
        failed = [item["name"] for item in report["checks"] if not item["passed"]]
        raise ValueError("UAE PMI quality checks failed: " + ", ".join(failed))
    write_csv(args.csv.resolve(), observations)

    workbook_path = args.workbook.resolve()
    workbook_status = "skipped"
    if not args.skip_workbook:
        if args.force or not target_is_current(workbook_path, observations):
            write_workbook(workbook_path, observations)
            workbook_status = "updated"
        else:
            _verify_workbook(workbook_path, observations)
            workbook_status = "already current"
    latest = observations[-1]
    print(
        f"Validated {len(observations)} UAE PMI months: "
        f"{observations[0].period} through {latest.period}"
    )
    print(f"Latest: {latest.period} = {latest.value:.1f}")
    print(f"CSV: {args.csv.resolve()}")
    print(f"Quality report: {args.quality_report.resolve()}")
    print(f"Workbook: {workbook_status} ({TARGET_SHEET})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
