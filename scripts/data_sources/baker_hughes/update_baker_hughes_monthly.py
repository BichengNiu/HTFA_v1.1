"""Extract UAE monthly rig counts and update ``data/阿联酋.xlsx``.

The newest Baker Hughes worldwide rig-count workbook in ``data/BakerHughes`` is
selected by its latest observation month. Oil-drilling rows from Abu Dhabi,
Dubai and Sharjah are aggregated and written to the managed
``月度_贝克休斯`` worksheet.
"""

from __future__ import annotations

import argparse
import json
import subprocess
import tempfile
import warnings
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


SOURCE_SHEET = "WW Monthly"
TARGET_SHEET = "月度_贝克休斯"
SOURCE_NAME = "Baker Hughes"
UAE_COUNTRY_PREFIX = "UAE - "
EXPECTED_STATUS = "Active Rigs"

INDICATORS = (("阿联酋石油活跃钻机数", "oil"),)

OBSOLETE_INDICATORS = (
    "阿联酋活跃钻机总数",
    "阿联酋石油钻机数",
    "阿联酋天然气钻机数",
    "阿联酋其他钻机数",
    "阿联酋陆地钻机数",
    "阿联酋海上钻机数",
)

REQUIRED_COLUMNS = (
    "Region",
    "Country",
    "DrillFor",
    "Location",
    "Rig Status",
    "Year",
    "Month",
    "Rig Count Value",
)


@dataclass(frozen=True)
class RigObservation:
    """Aggregated UAE oil rig count for one calendar month."""

    period: str
    oil: Decimal

    @property
    def values(self) -> tuple[Decimal, ...]:
        return tuple(getattr(self, key) for _, key in INDICATORS)


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _number(value: object) -> Decimal:
    if value is None or isinstance(value, bool):
        raise ValueError("Rig Count Value is blank or boolean")
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"Invalid Rig Count Value: {value!r}") from exc
    if not number.is_finite() or number < 0:
        raise ValueError(f"Invalid Rig Count Value: {value!r}")
    return number


def _find_header(sheet: object) -> tuple[int, dict[str, int]]:
    """Locate the normalized monthly table header in the source worksheet."""

    for row_number, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=min(40, sheet.max_row), values_only=True),
        start=1,
    ):
        columns = {
            _text(value): column
            for column, value in enumerate(row, start=1)
            if _text(value)
        }
        if all(name in columns for name in REQUIRED_COLUMNS):
            return row_number, columns
    raise ValueError(f"Required monthly-table header not found in {SOURCE_SHEET}")


def extract_uae_monthly(path: Path) -> list[RigObservation]:
    """Extract and aggregate all UAE monthly observations from one workbook."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if SOURCE_SHEET not in workbook.sheetnames:
            raise ValueError(f"Worksheet not found: {SOURCE_SHEET}")
        sheet = workbook[SOURCE_SHEET]
        header_row, columns = _find_header(sheet)
        monthly: dict[tuple[int, int], Decimal] = {}
        seen_rows: set[tuple[object, ...]] = set()

        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            country = _text(row[columns["Country"] - 1])
            if not country.startswith(UAE_COUNTRY_PREFIX):
                continue
            drill_for = _text(row[columns["DrillFor"] - 1])
            if drill_for != "Oil":
                continue
            status = _text(row[columns["Rig Status"] - 1])
            if status != EXPECTED_STATUS:
                raise ValueError(
                    f"Unexpected UAE rig status {status!r} in {path.name}"
                )
            year_value = row[columns["Year"] - 1]
            month_value = row[columns["Month"] - 1]
            try:
                year = int(year_value)
                month = int(month_value)
                datetime(year, month, 1)
            except (TypeError, ValueError) as exc:
                raise ValueError(
                    f"Invalid UAE observation period: {year_value!r}-{month_value!r}"
                ) from exc

            location = _text(row[columns["Location"] - 1])
            row_key = (year, month, country, drill_for, location, status)
            if row_key in seen_rows:
                raise ValueError(f"Duplicate UAE rig row: {row_key}")
            seen_rows.add(row_key)

            value = _number(row[columns["Rig Count Value"] - 1])
            if location not in {"Land", "Offshore"}:
                raise ValueError(f"Unexpected UAE Location value: {location!r}")
            monthly[(year, month)] = monthly.get(
                (year, month), Decimal("0")
            ) + value

        if not monthly:
            raise ValueError(f"No UAE monthly rig rows found in {path.name}")

        observations: list[RigObservation] = []
        for (year, month), oil_count in sorted(monthly.items()):
            observations.append(
                RigObservation(
                    period=f"{year:04d}-{month:02d}",
                    oil=oil_count,
                )
            )
        return observations
    finally:
        workbook.close()


def select_latest_workbook(
    directory: Path,
) -> tuple[Path, list[RigObservation], list[str]]:
    """Select the valid workbook containing the latest observation month."""

    candidates: list[tuple[str, int, str, Path, list[RigObservation]]] = []
    errors: list[str] = []
    for path in sorted(directory.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        try:
            observations = extract_uae_monthly(path)
        except (OSError, ValueError) as exc:
            errors.append(f"{path.name}: {exc}")
            continue
        candidates.append(
            (
                observations[-1].period,
                path.stat().st_mtime_ns,
                path.name,
                path,
                observations,
            )
        )
    if not candidates:
        details = "; ".join(errors) if errors else "no .xlsx files found"
        raise FileNotFoundError(f"No valid Baker Hughes workbook: {details}")
    _, _, _, selected_path, selected_observations = max(candidates)
    return selected_path, selected_observations, errors


def _normalized_number(value: object) -> Decimal | None:
    try:
        return _number(value)
    except ValueError:
        return None


def target_is_current(path: Path, observations: Iterable[RigObservation]) -> bool:
    """Return whether the managed target sheet already contains these values."""

    if not path.is_file():
        return False
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if TARGET_SHEET not in workbook.sheetnames:
            return False
        sheet = workbook[TARGET_SHEET]
        headers = {
            _text(value): column
            for column, value in enumerate(
                next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)),
                start=1,
            )
        }
        if any(name not in headers for name, _ in INDICATORS):
            return False
        actual: dict[str, tuple[Decimal | None, ...]] = {}
        for row in sheet.iter_rows(min_row=7, values_only=True):
            raw_date = row[0]
            if not isinstance(raw_date, (date, datetime)):
                continue
            actual[raw_date.strftime("%Y-%m")] = tuple(
                _normalized_number(row[headers[name] - 1])
                for name, _ in INDICATORS
            )
        expected = {item.period: item.values for item in observations}
        return actual == expected
    finally:
        workbook.close()


def _records_latest_first(
    observations: Iterable[RigObservation],
) -> list[dict[str, object]]:
    return [
        {
            "period": item.period,
            "values": [float(value) for value in item.values],
        }
        for item in sorted(observations, key=lambda item: item.period, reverse=True)
    ]


def write_monthly_sheet(path: Path, observations: Iterable[RigObservation]) -> None:
    """Write the managed sheet through Excel, preserving all other worksheets."""

    helper_path = Path(__file__).with_name("write_baker_hughes_sheet.ps1")
    if not helper_path.is_file():
        raise FileNotFoundError(f"Excel writer helper not found: {helper_path}")
    payload = {
        "dictionary_sheet_name": "指标字典",
        "metadata_labels": ["指标名称", "频率", "单位", "来源", "更新时间"],
        "frequency": "月",
        "unit": "台",
        "source": SOURCE_NAME,
        "updated_at": date.today().isoformat(),
        "indicators": [
            {"name": name, "type": "数量", "industry": "能源"}
            for name, _ in INDICATORS
        ],
        "obsolete_indicators": list(OBSOLETE_INDICATORS),
        "records": _records_latest_first(observations),
    }
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            prefix=".baker-hughes-sheet-",
            suffix=".json",
            dir=Path(__file__).resolve().parent,
            encoding="utf-8",
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
            json.dump(payload, temporary, ensure_ascii=True)
        completed = subprocess.run(
            (
                "powershell.exe",
                "-NoProfile",
                "-NonInteractive",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(helper_path),
                "-WorkbookPath",
                str(path),
                "-DataPath",
                str(temporary_path),
                "-SheetName",
                TARGET_SHEET,
            ),
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        if completed.returncode:
            message = (completed.stderr or completed.stdout).strip()
            raise RuntimeError(f"Excel sheet update failed: {message}")
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def parse_args() -> argparse.Namespace:
    repo_root = Path(__file__).resolve().parents[3]
    source_dir = repo_root / "data" / "BakerHughes"
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=source_dir)
    parser.add_argument(
        "--workbook",
        type=Path,
        default=repo_root / "data" / "阿联酋.xlsx",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Rewrite the managed sheet even when its values are current.",
    )
    return parser.parse_args()


def main() -> int:
    warnings.filterwarnings(
        "ignore",
        message="Unknown extension is not supported and will be removed",
        module="openpyxl.worksheet._reader",
    )
    args = parse_args()
    input_dir = args.input_dir.resolve()
    workbook_path = args.workbook.resolve()
    if not input_dir.is_dir():
        raise FileNotFoundError(f"Baker Hughes directory not found: {input_dir}")
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Destination workbook not found: {workbook_path}")

    source_path, observations, warnings_found = select_latest_workbook(input_dir)
    if not args.force and target_is_current(workbook_path, observations):
        print(
            f"Baker Hughes data already current through {observations[-1].period} "
            f"({source_path.name})"
        )
        return 0

    write_monthly_sheet(workbook_path, observations)
    latest = observations[-1]
    print(
        f"Updated {TARGET_SHEET} with {len(observations)} months in {workbook_path}"
    )
    print(
        f"Source: {source_path.name}; latest: {latest.period}; "
        f"UAE active oil rigs: {latest.oil}"
    )
    for warning in warnings_found:
        print(f"Source skipped: {warning}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
