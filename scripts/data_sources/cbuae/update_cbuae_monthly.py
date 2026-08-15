"""Update the CBUAE monthly-data sheet in the UAE source workbook.

The script scans every downloaded Statistical Bulletin workbook, reads the
all-banks domestic-credit and resident/non-resident deposit tables, and keeps
the latest available vintage for each observation month. PDF-only bulletins
are used as a fallback when no workbook exists for that bulletin month.
Wind data fill periods that are absent from the CBUAE bulletin series.

The normalized series are written directly to the ``月度_CBUAE`` worksheet in
``data/阿联酋.xlsx``. All measures are in millions of UAE dirhams (AED).
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from pypdf import PdfReader

PROJECT_ROOT = Path(__file__).resolve().parents[3]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.data_sources._excel_helpers import (
    payload_json_file,
    records_latest_first,
    run_powershell_sheet_writer,
)


MONTH_NUMBERS = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}

PERIOD_PATTERN = re.compile(
    r"\b(" + "|".join(MONTH_NUMBERS) + r")[a-z]*\s+(20\d{2})\b",
    re.IGNORECASE,
)
FILE_PERIOD_PATTERN = re.compile(r"^(20\d{2})-(0[1-9]|1[0-2])$")
NUMBER_PATTERN = re.compile(r"-?\d[\d,]*(?:\.\d+)?")

CBUAE_INDICATORS = (
    ("阿联酋政府存款", "存款"),
    ("阿联酋政府控股企业存款", "存款"),
    ("阿联酋政府信贷", "信贷"),
    ("阿联酋政府控股企业信贷", "信贷"),
)

WIND_FALLBACK_INDICATORS = (
    "阿联酋:银行存款:居民存款:政府部门",
    "阿联酋:银行存款:居民存款:政府相关实体(政府持股超50%)",
    "阿联酋:信贷总额:国内信贷:政府部门",
    "阿联酋:信贷总额:国内信贷:公共部门(政府相关实体)",
)

CBUAE_PRIMARY_START = "2020-01"


@dataclass(frozen=True)
class Observation:
    """One source vintage for one observation period."""

    period: str
    values: tuple[Decimal, Decimal, Decimal, Decimal]
    source_period: str
    source_file: str


def _normalize_label(value: object) -> str:
    """Normalize a table label without changing its semantic content."""

    if not isinstance(value, str):
        return ""
    value = value.replace("\u00a0", " ")
    value = re.sub(r"\*+", "", value)
    value = re.sub(r"\s+", " ", value).strip().casefold()
    return re.sub(r"\(\s*([^)]*?)\s*\)", r"(\1)", value)


def _parse_period(value: object) -> str | None:
    """Return YYYY-MM for English month headers used by the bulletins."""

    if not isinstance(value, str):
        return None
    match = PERIOD_PATTERN.search(value)
    if not match:
        return None
    month = MONTH_NUMBERS[match.group(1).casefold()]
    return f"{match.group(2)}-{month:02d}"


def _source_period(path: Path) -> str:
    """Read the bulletin period from a standardized YYYY-MM filename."""

    if not FILE_PERIOD_PATTERN.fullmatch(path.stem):
        raise ValueError(f"Unexpected bulletin filename: {path.name}")
    return path.stem


def _to_decimal(value: object) -> Decimal | None:
    """Convert an Excel numeric cell to a stable three-decimal value."""

    if value is None or isinstance(value, bool):
        return None
    try:
        number = Decimal(str(value).replace(",", ""))
    except (InvalidOperation, AttributeError):
        return None
    return number.quantize(Decimal("0.001"), rounding=ROUND_HALF_UP)


def _sheet_text(sheet: Worksheet, max_rows: int = 10) -> str:
    """Return normalized text from the title/header area of a worksheet."""

    values: list[str] = []
    for row in sheet.iter_rows(min_row=1, max_row=min(max_rows, sheet.max_row)):
        values.extend(str(cell.value) for cell in row if cell.value is not None)
    return _normalize_label(" ".join(values))


def _find_source_sheets(workbook: object) -> tuple[Worksheet, Worksheet]:
    """Find the all-banks domestic-credit and deposit worksheets."""

    credit_sheet: Worksheet | None = None
    deposit_sheet: Worksheet | None = None
    for sheet in workbook.worksheets:
        text = _sheet_text(sheet)
        if "domestic credit" in text and "all banks" in text:
            credit_sheet = sheet
        if (
            "deposits distributed residents / non residents" in text
            and "all banks" in text
        ):
            deposit_sheet = sheet
    if credit_sheet is None or deposit_sheet is None:
        raise ValueError("Required all-banks credit/deposit worksheets not found")
    return credit_sheet, deposit_sheet


def _header_columns(sheet: Worksheet) -> dict[str, int]:
    """Map observation period to worksheet column number."""

    columns: dict[str, int] = {}
    for row in sheet.iter_rows(min_row=1, max_row=min(12, sheet.max_row)):
        for cell in row:
            period = _parse_period(cell.value)
            if period:
                columns[period] = cell.column
    if not columns:
        raise ValueError(f"No English month headers found in {sheet.title}")
    return columns


def _find_row(sheet: Worksheet, accepted_labels: set[str]) -> int:
    """Find a row containing one of the exact normalized labels."""

    for row_number, row in enumerate(sheet.iter_rows(), start=1):
        if any(_normalize_label(cell.value) in accepted_labels for cell in row):
            return row_number
    labels = ", ".join(sorted(accepted_labels))
    raise ValueError(f"Labels not found in {sheet.title}: {labels}")


def _extract_sheet_rows(
    sheet: Worksheet,
    label_sets: tuple[set[str], set[str]],
) -> dict[str, tuple[Decimal, Decimal]]:
    """Extract two metric rows for every dated column in a worksheet."""

    columns = _header_columns(sheet)
    first_row = _find_row(sheet, label_sets[0])
    second_row = _find_row(sheet, label_sets[1])
    extracted: dict[str, tuple[Decimal, Decimal]] = {}
    for period, column in columns.items():
        first = _to_decimal(sheet.cell(first_row, column).value)
        second = _to_decimal(sheet.cell(second_row, column).value)
        if first is not None and second is not None:
            extracted[period] = (first, second)
    return extracted


def extract_workbook(path: Path) -> list[Observation]:
    """Extract all complete period observations from one workbook."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        credit_sheet, deposit_sheet = _find_source_sheets(workbook)
        deposits = _extract_sheet_rows(
            deposit_sheet,
            ({"government"}, {"gres"}),
        )
        credits = _extract_sheet_rows(
            credit_sheet,
            ({"government"}, {"public sector", "public sector (gres)"}),
        )
    finally:
        workbook.close()

    source_period = _source_period(path)
    observations: list[Observation] = []
    for period in sorted(deposits.keys() & credits.keys()):
        government_deposits, gre_deposits = deposits[period]
        government_credit, gre_credit = credits[period]
        observations.append(
            Observation(
                period=period,
                values=(
                    government_deposits,
                    gre_deposits,
                    government_credit,
                    gre_credit,
                ),
                source_period=source_period,
                source_file=path.name,
            )
        )
    return observations


def _last_number(line: str) -> Decimal | None:
    """Return the last numeric token in an extracted PDF table row."""

    matches = NUMBER_PATTERN.findall(line)
    if not matches:
        return None
    return _to_decimal(matches[-1])


def extract_pdf_fallback(path: Path) -> Observation:
    """Extract the bulletin-month values from a PDF-only bulletin."""

    deposit_values: tuple[Decimal, Decimal] | None = None
    credit_values: tuple[Decimal, Decimal] | None = None
    reader = PdfReader(path)

    for page in reader.pages:
        text = page.extract_text() or ""
        lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines()]
        if "Table 23 : Domestic Credit" in text or "Domestic Credit ( All Banks )" in text:
            government = next(
                (_last_number(line) for line in lines if re.match(r"^Government\s", line)),
                None,
            )
            public_sector = next(
                (_last_number(line) for line in lines if re.match(r"^Public Sector(?:\s|$)", line)),
                None,
            )
            if government is not None and public_sector is not None:
                credit_values = (government, public_sector)

        if (
            "Deposits distributed Residents / Non Residents ( All Banks )"
            in text
        ):
            government = next(
                (_last_number(line) for line in lines if re.match(r"^1\.3 Government\s", line)),
                None,
            )
            gres = next(
                (_last_number(line) for line in lines if re.match(r"^1\.4 GREs\s", line)),
                None,
            )
            if government is not None and gres is not None:
                deposit_values = (government, gres)

    if deposit_values is None or credit_values is None:
        raise ValueError(f"Required PDF tables could not be parsed: {path.name}")

    period = _source_period(path)
    return Observation(
        period=period,
        values=(*deposit_values, *credit_values),
        source_period=period,
        source_file=path.name,
    )


def extract_wind_fallback(path: Path) -> list[Observation]:
    """Read the four matching Wind series and convert billions to millions."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = (
            workbook["月度_Wind"]
            if "月度_Wind" in workbook.sheetnames
            else workbook.active
        )
        expected_names = {
            _normalize_label(name): name for name in WIND_FALLBACK_INDICATORS
        }
        matching_columns: dict[str, int] = {}
        header_row = next(
            sheet.iter_rows(min_row=2, max_row=2, values_only=True)
        )
        for column, value in enumerate(header_row[1:], start=2):
            normalized_name = _normalize_label(value)
            if normalized_name not in expected_names:
                continue
            if normalized_name in matching_columns:
                raise ValueError(
                    "Duplicate Wind fallback indicator: "
                    + expected_names[normalized_name]
                )
            matching_columns[normalized_name] = column

        missing = [
            name
            for name in WIND_FALLBACK_INDICATORS
            if _normalize_label(name) not in matching_columns
        ]
        if missing:
            raise ValueError(
                "Missing Wind fallback indicators: " + ", ".join(missing)
            )

        ordered_columns = [
            matching_columns[_normalize_label(name)]
            for name in WIND_FALLBACK_INDICATORS
        ]
        unit_row = next(
            sheet.iter_rows(min_row=4, max_row=4, values_only=True)
        )
        units = [
            _normalize_label(unit_row[column - 1])
            for column in ordered_columns
        ]
        if any(unit != _normalize_label("十亿阿联酋迪拉姆") for unit in units):
            raise ValueError("Wind indicators are not in billions of AED")

        observations: list[Observation] = []
        for row in sheet.iter_rows(min_row=7, values_only=True):
            period_value = row[0]
            if not isinstance(period_value, (date, datetime)):
                continue
            converted = tuple(
                _to_decimal(Decimal(str(value)) * Decimal("1000"))
                for value in (row[column - 1] for column in ordered_columns)
            )
            if any(value is None for value in converted):
                continue
            period = period_value.strftime("%Y-%m")
            observations.append(
                Observation(
                    period=period,
                    values=converted,
                    source_period=period,
                    source_file=path.name,
                )
            )
        return observations
    finally:
        workbook.close()


def _iter_months(start: str, end: str) -> Iterable[str]:
    """Yield inclusive YYYY-MM periods."""

    current = datetime.strptime(start, "%Y-%m")
    final = datetime.strptime(end, "%Y-%m")
    while current <= final:
        yield current.strftime("%Y-%m")
        year = current.year + (current.month == 12)
        month = 1 if current.month == 12 else current.month + 1
        current = current.replace(year=year, month=month)


def select_latest_vintages(
    observations: Iterable[Observation],
    start_period: str,
    end_period: str | None,
) -> tuple[list[Observation], list[str], int]:
    """Select the latest source vintage for each period and report revisions."""

    by_period: dict[str, list[Observation]] = defaultdict(list)
    for observation in observations:
        if observation.period >= start_period:
            by_period[observation.period].append(observation)

    if not by_period:
        raise ValueError("No observations were extracted")
    effective_end = end_period or max(by_period)
    selected: list[Observation] = []
    revised_periods = 0
    for period, candidates in sorted(by_period.items()):
        if period > effective_end:
            continue
        if len({candidate.values for candidate in candidates}) > 1:
            revised_periods += 1
        selected.append(max(candidates, key=lambda item: (item.source_period, item.source_file)))

    available = {observation.period for observation in selected}
    missing = [
        period
        for period in _iter_months(start_period, effective_end)
        if period not in available
    ]
    return selected, missing, revised_periods


def add_missing_fallbacks(
    primary: Iterable[Observation],
    fallback: Iterable[Observation],
    start_period: str,
    end_period: str | None,
) -> tuple[list[Observation], list[str]]:
    """Add fallback observations only where the primary series has no period."""

    merged = {observation.period: observation for observation in primary}
    added: list[str] = []
    for observation in fallback:
        if observation.period < start_period:
            continue
        if end_period is not None and observation.period > end_period:
            continue
        if observation.period not in merged:
            merged[observation.period] = observation
            added.append(observation.period)
    return [merged[period] for period in sorted(merged)], sorted(added)


def write_monthly_sheet(
    path: Path,
    observations: Iterable[Observation],
    sheet_name: str = "月度_CBUAE",
) -> None:
    """Update the managed sheet through Excel without rewriting other sheets."""

    helper_path = Path(__file__).with_name("write_cbuae_sheet.ps1")
    if not helper_path.is_file():
        raise FileNotFoundError(f"Excel writer helper not found: {helper_path}")

    payload = {
        "dictionary_sheet_name": "指标字典",
        "metadata_labels": ["指标名称", "频率", "单位", "来源", "更新时间"],
        "frequency": "月",
        "unit": "百万迪拉姆",
        "source": "CBUAE",
        "updated_at": date.today().isoformat(),
        "indicators": [
            {"name": name, "type": indicator_type, "industry": "金融"}
            for name, indicator_type in CBUAE_INDICATORS
        ],
        "records": records_latest_first(observations),
    }
    with payload_json_file(
        payload,
        prefix=".cbuae-sheet-",
        directory=Path(__file__).resolve().parent,
    ) as payload_path:
        run_powershell_sheet_writer(
            helper_path, path, payload_path, sheet_name
        )


def parse_args() -> argparse.Namespace:
    """Parse command-line arguments."""

    repo_root = Path(__file__).resolve().parents[3]
    cbuae_dir = repo_root / "data" / "CBUAE"
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input-dir",
        type=Path,
        default=cbuae_dir / "reports",
        help="Directory containing YYYY-MM.xlsx/PDF bulletin files.",
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=repo_root / "data" / "阿联酋.xlsx",
        help="Destination workbook containing the managed monthly sheet.",
    )
    parser.add_argument(
        "--wind-fallback",
        type=Path,
        default=None,
        help=(
            "Wind workbook used only for periods absent from the CBUAE series; "
            "defaults to the destination workbook."
        ),
    )
    parser.add_argument("--start-period", default="2019-12")
    parser.add_argument("--end-period", default=None)
    return parser.parse_args()


def main() -> int:
    """Run extraction and emit a compact quality summary."""

    args = parse_args()
    input_dir = args.input_dir.resolve()
    if not input_dir.is_dir():
        raise FileNotFoundError(f"Bulletin directory not found: {input_dir}")

    observations: list[Observation] = []
    workbook_errors: list[str] = []
    workbook_paths = sorted(input_dir.glob("20??-??.xlsx"))
    for path in workbook_paths:
        try:
            observations.extend(extract_workbook(path))
        except (OSError, ValueError) as exc:
            workbook_errors.append(f"{path.name}: {exc}")

    workbook_periods = {path.stem for path in workbook_paths}
    for path in sorted(input_dir.glob("20??-??.pdf")):
        if path.stem in workbook_periods:
            continue
        try:
            observations.append(extract_pdf_fallback(path))
        except (OSError, ValueError) as exc:
            workbook_errors.append(f"{path.name}: {exc}")

    primary_start = max(args.start_period, CBUAE_PRIMARY_START)
    selected, _, revised_periods = select_latest_vintages(
        observations,
        start_period=primary_start,
        end_period=args.end_period,
    )
    fallback_added: list[str] = []
    wind_path = (args.wind_fallback or args.workbook).resolve()
    if wind_path.is_file():
        try:
            selected, fallback_added = add_missing_fallbacks(
                selected,
                extract_wind_fallback(wind_path),
                start_period=args.start_period,
                end_period=args.end_period,
            )
        except (OSError, ValueError) as exc:
            workbook_errors.append(f"{wind_path.name}: {exc}")
    else:
        workbook_errors.append(f"{wind_path.name}: Wind fallback file not found")

    effective_end = args.end_period or max(observation.period for observation in selected)
    available = {observation.period for observation in selected}
    missing = [
        period
        for period in _iter_months(args.start_period, effective_end)
        if period not in available
    ]
    workbook_path = args.workbook.resolve()
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Destination workbook not found: {workbook_path}")
    write_monthly_sheet(workbook_path, selected)

    print(f"Updated 月度_CBUAE with {len(selected)} rows in {workbook_path}")
    print(f"Latest-vintage selection affected {revised_periods} periods")
    print(
        "Wind fallback periods: "
        + (", ".join(fallback_added) if fallback_added else "none")
    )
    print("Missing periods: " + (", ".join(missing) if missing else "none"))
    if workbook_errors:
        print("Source warnings:", file=sys.stderr)
        for error in workbook_errors:
            print(f"- {error}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
