"""Aggregate classified equipment trade and merge it into the UAE workbook."""

from __future__ import annotations

import csv
import json
import os
import shutil
import subprocess
import sys
from calendar import monthrange
from collections import Counter, defaultdict
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))

from scope_config import (  # noqa: E402
    CODE_DESCRIPTIONS,
    EXCLUDED_RELATED_CODES,
    SERIES,
    UAE_CODE,
    series_for_code,
)

BASE_DIR = Path(__file__).resolve().parent
RAW_DIR = BASE_DIR / "raw"
REFERENCE_DIR = BASE_DIR / "reference"
PROCESSED_DIR = BASE_DIR / "processed"
REPORTERS_PATH = REFERENCE_DIR / "reporters.json"
START_YEAR = 2017
ITEM_UNIT_CODE = 5


def load_reporters() -> dict[int, dict[str, Any]]:
    payload = json.loads(REPORTERS_PATH.read_text(encoding="utf-8"))
    records = payload.get("results", payload)
    return {int(record["reporterCode"]): record for record in records}


def iter_raw_rows(source: str):
    for path in sorted((RAW_DIR / source).glob("*.json")):
        payload = json.loads(path.read_text(encoding="utf-8"))
        for row in payload.get("data", []):
            yield row


def month_grid(latest_period: str) -> list[str]:
    """Return months newest-first from the latest observation to 2017-01."""
    result = []
    for year in range(START_YEAR, int(latest_period[:4]) + 1):
        for month in range(1, 13):
            period = f"{year}{month:02d}"
            if period <= latest_period:
                result.append(period)
    return list(reversed(result))


def item_quantity(row: dict[str, Any]) -> tuple[float, bool, str] | None:
    """Return one item count without double counting qty and altQty.

    UN Comtrade can expose the same supplementary quantity in both fields.  A
    non-estimated alternate quantity is closest to the reporter submission, so
    it takes precedence over the harmonized/estimated quantity field.
    """
    candidates = [
        ("altQty", "altQtyUnitCode", "isAltQtyEstimated", False),
        ("qty", "qtyUnitCode", "isQtyEstimated", False),
        ("altQty", "altQtyUnitCode", "isAltQtyEstimated", True),
        ("qty", "qtyUnitCode", "isQtyEstimated", True),
    ]
    for value_field, unit_field, estimate_field, estimated in candidates:
        if row.get(unit_field) != ITEM_UNIT_CODE:
            continue
        value = row.get(value_field)
        if value is None or float(value) <= 0:
            continue
        row_estimated = bool(row.get(estimate_field))
        if row_estimated != estimated:
            continue
        return float(value), row_estimated, value_field
    return None


def aggregate() -> tuple[
    list[dict[str, Any]],
    list[dict[str, Any]],
    dict[tuple[str, str], str],
    list[dict[str, Any]],
]:
    """Aggregate with UAE imports first and full-country mirror fallback."""
    reporters = load_reporters()
    direct_values: dict[tuple[str, str], float] = defaultdict(float)
    mirror_values: dict[tuple[str, str], float] = defaultdict(float)
    mirror_detail: dict[tuple[str, str, int], float] = defaultdict(float)
    direct_observed: set[tuple[str, str]] = set()
    mirror_observed: set[tuple[str, str]] = set()
    direct_units: dict[tuple[str, str, str], float] = defaultdict(float)
    direct_estimated_units: dict[tuple[str, str, str], float] = defaultdict(float)
    mirror_units: dict[tuple[str, str, str], float] = defaultdict(float)
    mirror_estimated_units: dict[tuple[str, str, str], float] = defaultdict(float)
    mirror_unit_reporters: dict[tuple[str, str, str], set[int]] = defaultdict(set)

    for row in iter_raw_rows("equipment_uae_reported"):
        period = str(row.get("period", ""))
        categories = series_for_code(str(row.get("cmdCode", "")))
        if len(period) != 6 or not categories:
            continue
        code = str(row.get("cmdCode", "")).zfill(6)
        quantity = item_quantity(row)
        for category in categories:
            if row.get("primaryValue") is not None:
                direct_observed.add((period, category))
                direct_values[(period, category)] += float(row["primaryValue"])
            if quantity is not None:
                value, estimated, _field = quantity
                unit_key = (period, category, code)
                direct_units[unit_key] += value
                if estimated:
                    direct_estimated_units[unit_key] += value

    unknown_reporters: set[int] = set()
    for row in iter_raw_rows("equipment_all_mirror"):
        period = str(row.get("period", ""))
        categories = series_for_code(str(row.get("cmdCode", "")))
        if len(period) != 6 or not categories:
            continue
        reporter = int(row.get("reporterCode") or 0)
        metadata = reporters.get(reporter)
        if metadata is None:
            unknown_reporters.add(reporter)
            continue
        if bool(metadata.get("isGroup")) or reporter == UAE_CODE:
            continue
        code = str(row.get("cmdCode", "")).zfill(6)
        quantity = item_quantity(row)
        for category in categories:
            if row.get("primaryValue") is not None:
                value = float(row["primaryValue"])
                mirror_observed.add((period, category))
                mirror_values[(period, category)] += value
                mirror_detail[(period, category, reporter)] += value
            if quantity is not None:
                value, estimated, _field = quantity
                unit_key = (period, category, code)
                mirror_units[unit_key] += value
                mirror_unit_reporters[unit_key].add(reporter)
                if estimated:
                    mirror_estimated_units[unit_key] += value

    if unknown_reporters:
        raise RuntimeError(
            "Unknown Comtrade reporter codes: "
            + ", ".join(map(str, sorted(unknown_reporters)))
        )

    available_months = {
        period for period, _category in direct_observed
    } | {
        period for period, _category in mirror_observed
    } | {
        period for period, _category, _code in direct_units
    } | {
        period for period, _category, _code in mirror_units
    }
    if not available_months:
        raise RuntimeError("No monthly equipment observations found")

    latest_period = max(available_months)
    sources: dict[tuple[str, str], str] = {}
    quantity_detail: list[dict[str, Any]] = []
    output: list[dict[str, Any]] = []
    for period in month_grid(latest_period):
        year = int(period[:4])
        month = int(period[4:])
        result: dict[str, Any] = {
            "month": period,
            "date": date(year, month, monthrange(year, month)[1]),
        }
        for category in SERIES:
            key = (period, category)
            if key in direct_observed:
                value: float | None = direct_values[key]
                sources[key] = "uae_reported"
            elif key in mirror_observed:
                value = mirror_values[key]
                sources[key] = "all_partner_mirror"
            else:
                value = None
                sources[key] = "missing"
            result[f"{category}_usd"] = (
                round(value, 3) if value is not None else None
            )
            result[f"{category}_usd_mn"] = (
                round(value / 1_000_000, 6) if value is not None else None
            )
            units = 0.0
            unit_codes = 0
            for code in sorted(SERIES[category]["codes"]):
                unit_key = (period, category, code)
                if unit_key in direct_units:
                    unit_value = direct_units[unit_key]
                    estimated_units = direct_estimated_units[unit_key]
                    unit_source = "uae_reported"
                    reporter_count = 1
                elif unit_key in mirror_units:
                    unit_value = mirror_units[unit_key]
                    estimated_units = mirror_estimated_units[unit_key]
                    unit_source = "all_partner_mirror"
                    reporter_count = len(mirror_unit_reporters[unit_key])
                else:
                    continue
                units += unit_value
                unit_codes += 1
                quantity_detail.append(
                    {
                        "month": period,
                        "series": category,
                        "hs6": code,
                        "source": unit_source,
                        "units": round(unit_value, 3),
                        "reported_units": round(
                            unit_value - estimated_units,
                            3,
                        ),
                        "estimated_units": round(estimated_units, 3),
                        "reporter_count": reporter_count,
                    }
                )
            result[f"{category}_units"] = (
                round(units, 3) if unit_codes else None
            )
        output.append(result)

    detail = []
    for (period, category, reporter), value in sorted(
        mirror_detail.items(),
        key=lambda item: item[0],
        reverse=True,
    ):
        if (period, category) in direct_observed:
            continue
        detail.append(
            {
                "month": period,
                "series": category,
                "reporter_code": reporter,
                "reporter_name": reporters[reporter].get("reporterDesc", ""),
                "trade_value_usd": round(value, 3),
            }
        )
    return output, detail, sources, quantity_detail


def write_csvs(
    rows: list[dict[str, Any]],
    detail: list[dict[str, Any]],
    quantity_detail: list[dict[str, Any]],
) -> None:
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    monthly_path = PROCESSED_DIR / "uae_imports_monthly.csv"
    with monthly_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)

    detail_path = PROCESSED_DIR / "mirror_partner_detail.csv"
    fields = [
        "month",
        "series",
        "reporter_code",
        "reporter_name",
        "trade_value_usd",
    ]
    with detail_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(detail)

    quantity_path = PROCESSED_DIR / "quantity_source_detail.csv"
    quantity_fields = [
        "month",
        "series",
        "hs6",
        "source",
        "units",
        "reported_units",
        "estimated_units",
        "reporter_count",
    ]
    with quantity_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=quantity_fields)
        writer.writeheader()
        writer.writerows(quantity_detail)


def write_methodology_markdown(
    rows: list[dict[str, Any]],
    detail: list[dict[str, Any]],
    sources: dict[tuple[str, str], str],
) -> Path:
    latest_period = rows[0]["month"]
    latest_label = f"{latest_period[:4]}-{latest_period[4:]}"
    periods = list(reversed([row["month"] for row in rows]))
    source_names = {
        "uae_reported": "阿联酋直报",
        "all_partner_mirror": "出口镜像",
        "missing": "无数据",
    }
    reporter_counts: dict[tuple[str, str], set[int]] = defaultdict(set)
    for row in detail:
        reporter_counts[(row["month"], row["series"])].add(
            int(row["reporter_code"])
        )

    def period_label(period: str) -> str:
        return f"{period[:4]}-{period[4:]}"

    def source_ranges(category: str) -> list[tuple[str, str, str]]:
        ranges: list[tuple[str, str, str]] = []
        start = periods[0]
        previous = periods[0]
        state = sources[(start, category)]
        for period in periods[1:]:
            current = sources[(period, category)]
            if current != state:
                ranges.append((start, previous, state))
                start = period
                state = current
            previous = period
        ranges.append((start, previous, state))
        return ranges

    lines = [
        "# UN Comtrade 重点设备进口数据口径",
        "",
        "## 输出定义",
        "",
        "- 工作簿：`../阿联酋.xlsx`",
        "- 工作表：`月度_UNComtrade`",
        f"- 时间范围：2017-01 至 {latest_label}，按从新到旧排列。",
        "- 数值为单月值，不做滚动平均或滚动合计。",
        "- 金额字段为 UN Comtrade `primaryValue`，工作表单位为百万美元。",
        "- 台数字段仅使用数量单位代码 5（台/件），五个分类内的 HS6 台数直接合计。",
        "- 日期写为月末；不做季节调整、通胀调整或汇率换算。",
        "",
        "## 数据查询和处理方法",
        "",
        "- 分类体系：HS，月度频率，使用下列精确 HS6 编码。",
        "- 阿联酋直报查询：申报国 UAE（784）、伙伴 World（0）、流向 Import。",
        "- 出口镜像查询：所有申报国、伙伴 UAE（784）、流向 Export。",
        "- 镜像保留 `reporters.json` 中 `isGroup=false` 的独立申报国，并排除 UAE 自身。",
        "- 不采用固定伙伴名单，也不设置前 N 国；当月有记录的全部独立申报国均参与求和。",
        "- 每个月、每个分类独立选择来源。分类内任一 HS6 有阿联酋 `primaryValue` 时，",
        "  该分类当月只汇总阿联酋直报；否则汇总全体出口国镜像。",
        "- 镜像中某分类只要当月有记录即汇总；所有出口国均无记录时保持空白。",
        "- 分类金额等于该分类所有 HS6、所选来源全部记录的 `primaryValue` 之和。",
        "- 台数按每个 HS6、每个月独立选择来源：有阿联酋台/件数量就用阿联酋，",
        "  否则汇总全部独立申报国对阿联酋出口的台/件数量，再加总为分类台数。",
        "- `qty` 与 `altQty` 同为台/件时只取一项；优先取非估算的 `altQty`，",
        "  其次取非估算的 `qty`，最后才保留带估算标识的台数。",
        "- 台数来源与估算数量保存在 `processed/quantity_source_detail.csv`。",
        "- 最新月份等于任一分类、任一来源存在记录的最新月份，之后月份不写入。",
        "",
        "## 分类定义",
        "",
    ]
    for definition in SERIES.values():
        lines.extend(
            [
                f"### {definition['name_zh']}",
                "",
                f"口径：{definition['basis']}。",
                "",
                "| HS6 | 商品说明 |",
                "|---|---|",
            ]
        )
        for code in sorted(definition["codes"]):
            lines.append(f"| {code} | {CODE_DESCRIPTIONS[code]} |")
        lines.append("")

    lines.extend(
        [
            "## 纳入边界和排除项",
            "",
            "本次核查把与原方案用途一致、名称可明确归类的整机补入。制造业口径",
            "补齐自动化金属加工和橡塑加工整机；施工口径补齐土方、隧道、移动",
            "起重、混凝土和沥青施工整机；能源口径补齐较小功率段蒸汽轮机和",
            "燃气轮机；钻探口径补入流动钻探机；港口铁路口径补齐港口搬运、",
            "机车和铁路养护设备。",
            "零部件和用途无法唯一判定的兜底编码不纳入。",
            "",
            "| 排除编码/范围 | 排除原因 |",
            "|---|---|",
        "",
        ]
    )
    for code, reason in EXCLUDED_RELATED_CODES.items():
        lines.append(f"| {code} | {reason} |")

    lines.extend(
        [
            "",
            "## 实际来源覆盖",
            "",
            "下表的数字为该年份各分类分别使用阿联酋直报、出口镜像和无数据的月份数。",
            "2026 年只统计输出表已有的 1—6 月。",
            "",
            "| 年份 | 制造业设备 | 土木施工 | 能源项目 | 钻探设备 | 港口铁路 |",
            "|---|---|---|---|---|---|",
        ]
    )
    for year in range(START_YEAR, int(latest_period[:4]) + 1):
        year_periods = [period for period in periods if period.startswith(str(year))]
        cells = []
        for category in SERIES:
            counts = Counter(sources[(period, category)] for period in year_periods)
            cells.append(
                f"直{counts['uae_reported']}/镜{counts['all_partner_mirror']}/"
                f"空{counts['missing']}"
            )
        lines.append(f"| {year} | " + " | ".join(cells) + " |")

    lines.extend(
        [
            "",
            "### 各分类连续来源区间",
            "",
            "| 分类 | 起始月 | 截止月 | 实际来源 |",
            "|---|---|---|---|",
        ]
    )
    for category, definition in SERIES.items():
        for start, end, source in source_ranges(category):
            lines.append(
                f"| {definition['name_zh']} | {period_label(start)} | "
                f"{period_label(end)} | {source_names[source]} |"
            )

    total_source_counts = Counter(sources.values())
    lines.extend(
        [
            "",
            "### 来源汇总",
            "",
            f"- 阿联酋直报单元：{total_source_counts['uae_reported']} 个。",
            f"- 全体出口国镜像单元：{total_source_counts['all_partner_mirror']} 个。",
            f"- 无数据单元：{total_source_counts['missing']} 个。",
            "",
            "### 完整月度来源台账",
            "",
            "`镜像(n国)` 中的 n 是该分类当月参与汇总的独立申报国数量。",
            "",
            "| 月份 | 制造业设备 | 土木施工 | 能源项目 | 钻探设备 | 港口铁路 |",
            "|---|---|---|---|---|---|",
        ]
    )
    for period in reversed(periods):
        cells = []
        for category in SERIES:
            source = sources[(period, category)]
            if source == "all_partner_mirror":
                count = len(reporter_counts[(period, category)])
                cells.append(f"镜像({count}国)")
            else:
                cells.append(source_names[source])
        lines.append(f"| {period_label(period)} | " + " | ".join(cells) + " |")

    lines.extend(
        [
            "## 注意事项",
            "",
            "- 阿联酋直报进口通常按 CIF、伙伴出口通常按 FOB，来源切换存在口径断点。",
            "- 镜像数据存在申报时滞；最近月份即使有值，也可能只覆盖少数申报国。",
            "- 台/件数量可含 UN Comtrade 估算值，因此可能出现小数；不强制取整。",
            "- 分类台数是异质设备件数之和，适合监测总件数，不代表等质量设备或产能。",
            "- 能源、钻探、港口和铁路设备可能受单笔大型项目影响，月度波动较大。",
            "- 本指标是重点设备监测口径，不代表全部工程或机械进口。",
            "",
        ]
    )
    destination = BASE_DIR / "UNComtrade_口径.md"
    destination.write_text("\n".join(lines), encoding="utf-8")
    return destination


def locate_workbook() -> Path:
    candidates = [
        path
        for path in BASE_DIR.parent.glob("*.xlsx")
        if not path.name.startswith("~$")
    ]
    if len(candidates) != 1:
        raise RuntimeError(
            f"Expected one workbook in data/, found {len(candidates)}"
        )
    return candidates[0]


def apply_header_style(worksheet, row: int, max_column: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[row][:max_column]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def write_data_sheet(workbook, rows: list[dict[str, Any]]) -> None:
    title = "月度_UNComtrade"
    if title in workbook.sheetnames:
        del workbook[title]
    worksheet = workbook.create_sheet(title)
    headers = ["日期"] + [
        f"{definition['name_zh']}:百万美元"
        for definition in SERIES.values()
    ] + [
        f"{definition['name_zh']}:台/件"
        for definition in SERIES.values()
    ]
    worksheet.append(headers)
    for row in rows:
        worksheet.append(
            [row["date"]]
            + [row[f"{category}_usd_mn"] for category in SERIES]
            + [row[f"{category}_units"] for category in SERIES]
        )
    apply_header_style(worksheet, 1, len(headers))
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.column_dimensions["A"].width = 13
    for column in range(2, len(headers) + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 28
    for cell in worksheet["A"][1:]:
        cell.number_format = "yyyy-mm"
    for row_cells in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
        min_col=2,
        max_col=len(headers),
    ):
        for cell in row_cells:
            cell.number_format = "#,##0.000"


def update_dictionary(workbook) -> None:
    worksheet = workbook.worksheets[0]
    obsolete = {
        "阿联酋:进口:资本品:当月值",
        "阿联酋:进口:钢材及结构金属:当月值",
        "阿联酋:进口:工程及工业机械:当月值",
        "阿联酋:进口:水泥:当月值",
        "阿联酋:进口:能源大型项目设备:当月值",
    }
    for row in range(worksheet.max_row, 1, -1):
        if str(worksheet.cell(row, 1).value) in obsolete:
            worksheet.delete_rows(row)
    existing = {
        str(worksheet.cell(row, 1).value)
        for row in range(2, worksheet.max_row + 1)
    }
    themes = {
        "manufacturing_equipment": "工业",
        "civil_construction_equipment": "建筑",
        "energy_project_equipment": "能源",
        "drilling_equipment": "能源",
        "port_rail_equipment": "基建",
    }
    amount_entries = [
        (
            f"阿联酋:进口:{definition['name_zh'].removesuffix('进口')}:当月值",
            "金额",
            themes[category],
            "UN Comtrade",
        )
        for category, definition in SERIES.items()
    ]
    quantity_entries = [
        (
            f"阿联酋:进口:{definition['name_zh'].removesuffix('进口')}:台数:当月值",
            "数量",
            themes[category],
            "UN Comtrade",
        )
        for category, definition in SERIES.items()
    ]
    entries = amount_entries + quantity_entries
    template_row = worksheet.max_row
    for entry in entries:
        if entry[0] in existing:
            continue
        worksheet.append([*entry, None])
        new_row = worksheet.max_row
        for column in range(1, worksheet.max_column + 1):
            source = worksheet.cell(template_row, column)
            target = worksheet.cell(new_row, column)
            if source.has_style:
                target._style = copy(source._style)
            if source.number_format:
                target.number_format = source.number_format


def merge_workbook(rows: list[dict[str, Any]]) -> Path:
    source = locate_workbook()
    backup_dir = BASE_DIR / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = backup_dir / f"{source.stem}_before_uncomtrade_{stamp}.xlsx"
    shutil.copy2(source, backup)

    workbook = load_workbook(source, keep_links=True)
    write_data_sheet(workbook, rows)
    if "UNComtrade_口径" in workbook.sheetnames:
        del workbook["UNComtrade_口径"]
    update_dictionary(workbook)
    temporary = BASE_DIR / f"{source.stem}_updated_{stamp}.xlsx"
    workbook.save(temporary)
    workbook.close()

    check = load_workbook(temporary, read_only=True, data_only=False)
    worksheet = check["月度_UNComtrade"]
    actual_size = (worksheet.max_row, worksheet.max_column)
    check.close()
    expected_size = (len(rows) + 1, len(SERIES) * 2 + 1)
    if actual_size != expected_size:
        raise RuntimeError(
            f"Workbook validation failed: {actual_size} != {expected_size}"
        )
    try:
        os.replace(temporary, source)
        return source
    except PermissionError:
        alternate = BASE_DIR / f"{source.stem}_含UNComtrade设备数据.xlsx"
        os.replace(temporary, alternate)
        return alternate


def write_quality_report(
    rows: list[dict[str, Any]],
    detail: list[dict[str, Any]],
    sources: dict[str, str],
    quantity_detail: list[dict[str, Any]],
    workbook_path: Path,
) -> None:
    months = [row["month"] for row in rows]
    reporter_counts: dict[str, dict[str, set[int]]] = defaultdict(
        lambda: defaultdict(set)
    )
    for row in detail:
        if float(row["trade_value_usd"]) != 0:
            reporter_counts[row["month"]][row["series"]].add(
                int(row["reporter_code"])
            )
    report = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "row_count": len(rows),
        "unique_month_count": len(set(months)),
        "first_month": min(months),
        "last_month": max(months),
        "source_cell_counts": {
            source: sum(value == source for value in sources.values())
            for source in {
                "uae_reported",
                "all_partner_mirror",
                "missing",
            }
        },
        "missing_by_series": {
            category: sum(row[f"{category}_usd"] is None for row in rows)
            for category in SERIES
        },
        "quantity_missing_by_series": {
            category: sum(row[f"{category}_units"] is None for row in rows)
            for category in SERIES
        },
        "quantity_source_rows": dict(
            Counter(row["source"] for row in quantity_detail)
        ),
        "estimated_quantity_share": (
            round(
                sum(float(row["estimated_units"]) for row in quantity_detail)
                / sum(float(row["units"]) for row in quantity_detail),
                6,
            )
            if sum(float(row["units"]) for row in quantity_detail)
            else None
        ),
        "latest_month_reporter_counts": {
            category: len(reporter_counts[max(months)][category])
            for category in SERIES
        },
        "negative_value_count": sum(
            value < 0
            for row in rows
            for key, value in row.items()
            if key.endswith("_usd") and value is not None
        ),
        "workbook_output": str(workbook_path),
    }
    (PROCESSED_DIR / "quality_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def main() -> None:
    rows, detail, sources, quantity_detail = aggregate()
    write_csvs(rows, detail, quantity_detail)
    write_methodology_markdown(rows, detail, sources)
    native_script = BASE_DIR / "merge_with_excel.ps1"
    if os.name == "nt" and native_script.exists():
        subprocess.run(
            [
                "powershell",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(native_script),
            ],
            check=True,
            cwd=BASE_DIR,
        )
        workbook_path = locate_workbook()
    else:
        workbook_path = merge_workbook(rows)
    write_quality_report(
        rows,
        detail,
        sources,
        quantity_detail,
        workbook_path,
    )
    print(f"Monthly rows: {len(rows)}")
    print(f"Workbook: {workbook_path}")


if __name__ == "__main__":
    main()
