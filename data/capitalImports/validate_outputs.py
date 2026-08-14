"""Validate equipment trade caches, processed files and merged workbook."""

from __future__ import annotations

import csv
import json
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE_DIR))

from process_and_merge import aggregate, month_grid  # noqa: E402
from scope_config import SERIES  # noqa: E402


def load_monthly_rows() -> list[dict[str, str]]:
    path = BASE_DIR / "processed" / "uae_imports_monthly.csv"
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def validate_raw_batches(errors: list[str]) -> dict[str, int]:
    expected_years = date.today().year - 2017 + 1
    counts = {
        source: len(list((BASE_DIR / "raw" / source).glob("*.json")))
        for source in ("equipment_uae_reported", "equipment_all_mirror")
    }
    expected = {source: expected_years for source in counts}
    if counts != expected:
        errors.append(f"Unexpected raw batch counts: {counts}; expected {expected}")
    for folder in counts:
        for path in (BASE_DIR / "raw" / folder).glob("*.json"):
            payload = json.loads(path.read_text(encoding="utf-8"))
            if payload.get("error"):
                errors.append(f"API error in {path.name}: {payload['error']}")
            if len(payload.get("data", [])) >= 100000:
                errors.append(f"Possible truncation in {path.name}")
    return counts


def validate_monthly(
    rows: list[dict[str, str]],
    errors: list[str],
) -> dict[str, object]:
    if not rows:
        errors.append("Monthly CSV is empty")
        return {"row_count": 0}
    months = [row["month"] for row in rows]
    duplicates = [month for month, count in Counter(months).items() if count > 1]
    expected_months = month_grid(months[0])
    if months != expected_months:
        errors.append(
            f"Expected contiguous months {months[0]} through 201701"
        )
    if duplicates:
        errors.append(f"Duplicate months: {duplicates}")

    expected_fields = {"month", "date"}
    for category in SERIES:
        expected_fields.add(f"{category}_usd")
        expected_fields.add(f"{category}_usd_mn")
        expected_fields.add(f"{category}_units")
    if set(rows[0]) != expected_fields:
        errors.append(f"Unexpected monthly fields: {set(rows[0]) ^ expected_fields}")

    expected_rows, detail, sources, quantity_detail = aggregate()
    for index, (actual, expected) in enumerate(zip(rows, expected_rows), start=2):
        if actual["month"] != expected["month"]:
            errors.append(f"Month mismatch at CSV row {index}")
            continue
        for category in SERIES:
            field = f"{category}_usd"
            actual_value = float(actual[field]) if actual[field] else None
            expected_value = expected[field]
            if actual_value != expected_value:
                errors.append(f"Aggregation mismatch at {actual['month']} {category}")
            if actual_value is not None and actual_value < 0:
                errors.append(f"Negative value at {actual['month']} {category}")
            units_field = f"{category}_units"
            actual_units = (
                float(actual[units_field]) if actual[units_field] else None
            )
            expected_units = expected[units_field]
            if actual_units != expected_units:
                errors.append(
                    f"Quantity mismatch at {actual['month']} {category}"
                )
            if actual_units is not None and actual_units < 0:
                errors.append(
                    f"Negative quantity at {actual['month']} {category}"
                )

    latest_month = months[0]
    latest_reporters = {
        category: {
            int(row["reporter_code"])
            for row in detail
            if row["month"] == latest_month
            and row["series"] == category
            and float(row["trade_value_usd"]) != 0
        }
        for category in SERIES
    }
    return {
        "row_count": len(rows),
        "unique_month_count": len(set(months)),
        "duplicate_months": duplicates,
        "first_row_month": months[0],
        "last_row_month": months[-1],
        "missing_by_series": {
            category: sum(row[f"{category}_usd"] == "" for row in rows)
            for category in SERIES
        },
        "quantity_missing_by_series": {
            category: sum(row[f"{category}_units"] == "" for row in rows)
            for category in SERIES
        },
        "quantity_detail_rows": len(quantity_detail),
        "quantity_source_rows": dict(
            Counter(row["source"] for row in quantity_detail)
        ),
        "estimated_quantity_share": (
            round(
                sum(float(row["estimated_units"]) for row in quantity_detail)
                / sum(float(row["units"]) for row in quantity_detail),
                6,
            )
            if sum(float(row["units"]) for row in quantity_detail)
            else None
        ),
        "source_cell_counts": dict(Counter(sources.values())),
        "latest_month_reporter_counts": {
            category: len(reporters)
            for category, reporters in latest_reporters.items()
        },
    }


def formula_count(worksheet) -> int:
    return sum(
        isinstance(cell.value, str) and cell.value.startswith("=")
        for row in worksheet.iter_rows()
        for cell in row
    )


def locate_workbook() -> Path:
    candidates = [
        path
        for path in BASE_DIR.parent.glob("*.xlsx")
        if not path.name.startswith("~$")
    ]
    if len(candidates) != 1:
        raise RuntimeError(f"Expected one data workbook, found {len(candidates)}")
    return candidates[0]


def validate_workbook(
    rows: list[dict[str, str]],
    errors: list[str],
) -> dict[str, object]:
    path = locate_workbook()
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=True)
    if "月度_UNComtrade" not in workbook.sheetnames:
        errors.append("Missing 月度_UNComtrade worksheet")
        workbook.close()
        return {"path": str(path), "exists": False}
    if "UNComtrade_口径" in workbook.sheetnames:
        errors.append("Methodology sheet remains in workbook")
    worksheet = workbook["月度_UNComtrade"]
    expected_size = (len(rows) + 1, len(SERIES) * 2 + 1)
    actual_size = (worksheet.max_row, worksheet.max_column)
    if actual_size != expected_size:
        errors.append(f"Unexpected data sheet size: {actual_size} != {expected_size}")
    for index, row in enumerate(rows, start=2):
        for column, category in enumerate(SERIES, start=2):
            text = row[f"{category}_usd_mn"]
            csv_value = float(text) if text else None
            excel_value = worksheet.cell(index, column).value
            if csv_value is None and excel_value is None:
                continue
            if (
                csv_value is None
                or excel_value is None
                or abs(csv_value - float(excel_value)) > 0.000001
            ):
                errors.append(
                    f"CSV/workbook mismatch at row {index}, column {column}"
                )
                break
        for column, category in enumerate(
            SERIES,
            start=len(SERIES) + 2,
        ):
            text = row[f"{category}_units"]
            csv_value = float(text) if text else None
            excel_value = worksheet.cell(index, column).value
            if csv_value is None and excel_value is None:
                continue
            if (
                csv_value is None
                or excel_value is None
                or abs(csv_value - float(excel_value)) > 0.000001
            ):
                errors.append(
                    f"CSV/workbook quantity mismatch at row {index}, "
                    f"column {column}"
                )
                break

    dictionary_values = {
        str(workbook.worksheets[0].cell(row, 1).value)
        for row in range(2, workbook.worksheets[0].max_row + 1)
    }
    expected_dictionary = {
        f"阿联酋:进口:{definition['name_zh'].removesuffix('进口')}:当月值"
        for definition in SERIES.values()
    }
    expected_dictionary.update(
        {
            f"阿联酋:进口:{definition['name_zh'].removesuffix('进口')}:台数:当月值"
            for definition in SERIES.values()
        }
    )
    obsolete_dictionary = {
        "阿联酋:进口:资本品:当月值",
        "阿联酋:进口:钢材及结构金属:当月值",
        "阿联酋:进口:工程及工业机械:当月值",
        "阿联酋:进口:水泥:当月值",
        "阿联酋:进口:能源大型项目设备:当月值",
    }
    if not expected_dictionary.issubset(dictionary_values):
        errors.append("Equipment indicator dictionary entries are missing")
    if obsolete_dictionary.intersection(dictionary_values):
        errors.append("Obsolete UN Comtrade dictionary entries remain")

    original_structure_ok = True
    backup_dir = BASE_DIR / "backups"
    native_backups = list(backup_dir.glob("*_before_native_merge_*.xlsx"))
    legacy_backups = list(backup_dir.glob("*_before_uncomtrade_*.xlsx"))
    backups = native_backups or legacy_backups
    baseline_path = (
        max(backups, key=lambda item: (item.stat().st_mtime_ns, item.name))
        if backups
        else None
    )
    if baseline_path is not None:
        baseline = load_workbook(
            baseline_path,
            read_only=True,
            data_only=False,
            keep_links=True,
        )
        for sheet_name in baseline.sheetnames:
            if (
                sheet_name == baseline.sheetnames[0]
                or sheet_name == "月度_UNComtrade"
                or sheet_name.endswith("_Wind")
            ):
                continue
            if sheet_name not in workbook.sheetnames:
                original_structure_ok = False
                errors.append(f"Original sheet missing: {sheet_name}")
                continue
            before = baseline[sheet_name]
            after = workbook[sheet_name]
            if (
                before.max_row != after.max_row
                or before.max_column != after.max_column
                or formula_count(before) != formula_count(after)
            ):
                original_structure_ok = False
                errors.append(f"Original sheet changed structurally: {sheet_name}")
        baseline.close()
    sheet_count = len(workbook.sheetnames)
    workbook.close()
    return {
        "path": str(path),
        "sheet_count": sheet_count,
        "data_sheet_rows": actual_size[0],
        "data_sheet_columns": actual_size[1],
        "original_structure_ok": original_structure_ok,
        "structure_baseline": str(baseline_path) if baseline_path else None,
    }


def validate_methodology(errors: list[str]) -> dict[str, object]:
    path = BASE_DIR / "UNComtrade_口径.md"
    if not path.exists():
        errors.append("Methodology Markdown is missing")
        return {"path": str(path), "exists": False}
    text = path.read_text(encoding="utf-8")
    required = [definition["name_zh"] for definition in SERIES.values()]
    required.extend(
        [
            "每个月、每个分类独立选择来源",
            "不采用固定伙伴名单",
            "完整月度来源台账",
        ]
    )
    missing = [term for term in required if term not in text]
    if missing:
        errors.append(f"Methodology is missing: {missing}")
    return {"path": str(path), "exists": True, "characters": len(text)}


def validate_secret_isolation(errors: list[str]) -> bool:
    env_path = BASE_DIR / ".env"
    if not env_path.exists():
        return True
    key = ""
    for line in env_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("COMTRADE_API_KEY="):
            key = line.split("=", 1)[1].strip()
            break
    if not key:
        errors.append("API key is missing from .env")
        return False
    needle = key.encode("utf-8")
    for path in BASE_DIR.rglob("*"):
        allowed_extensions = {".py", ".md", ".csv", ".json"}
        if not path.is_file() or path.suffix.lower() not in allowed_extensions:
            continue
        if path.resolve() == env_path.resolve():
            continue
        if needle in path.read_bytes():
            errors.append(f"API key leaked into {path}")
            return False
    return True


def main() -> None:
    errors: list[str] = []
    rows = load_monthly_rows()
    report = {
        "validated_at": datetime.now().isoformat(timespec="seconds"),
        "raw_batches": validate_raw_batches(errors),
        "monthly": validate_monthly(rows, errors),
        "workbook": validate_workbook(rows, errors),
        "methodology_markdown": validate_methodology(errors),
        "secret_isolation_ok": validate_secret_isolation(errors),
        "errors": errors,
        "status": "passed" if not errors else "failed",
    }
    report_path = BASE_DIR / "processed" / "quality_report.json"
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
