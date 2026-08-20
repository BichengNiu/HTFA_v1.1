"""Complete DuckDB metadata and synchronize workbook indicator metadata.

The observation tables are the source of truth for values.  This module only
maintains metadata: numeric business series go to ``meta_indicator_dictionary``
and raw/detail fields go to ``meta_column_dictionary``.  Workbook metadata is
updated only for indicators that are actually present in a workbook data sheet;
future or detail-only database variables therefore do not make the workbook
parser reject an otherwise valid workbook.
"""

from __future__ import annotations

import copy
import os
import sys
import tempfile
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SCRIPTS_DIR = Path(__file__).resolve().parent
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import db


def _indicator(
    name: str,
    frequency: str,
    unit: str | None,
    source: str,
    type_: str,
    industry: str,
) -> dict[str, Any]:
    return {
        "indicator_name": name,
        "frequency": frequency,
        "unit": unit,
        "source": source,
        "type": type_,
        "industry": industry,
    }


# These are the numeric series described in section 2 of
# data/UAE/变量说明_uae_duckdb.md but previously absent from the database
# indicator dictionary.  Workbook-facing names intentionally match the
# existing sheets where a sheet already exists.
INDICATOR_DEFINITIONS: tuple[dict[str, Any], ...] = (
    _indicator(
        "菲律宾_DMW部署_总计",
        "月",
        "人",
        "菲律宾 DMW Monthly Compendium Tab 11",
        "就业流量",
        "劳工",
    ),
    _indicator(
        "菲律宾_DMW部署_新雇",
        "月",
        "人",
        "菲律宾 DMW Monthly Compendium Tab 11",
        "就业流量",
        "劳工",
    ),
    _indicator(
        "菲律宾_DMW部署_再雇",
        "月",
        "人",
        "菲律宾 DMW Monthly Compendium Tab 11",
        "就业流量",
        "劳工",
    ),
    _indicator(
        "尼泊尔_DoFE批准_含再入境",
        "月",
        "人",
        "尼泊尔 DoFE monthly final labour approval",
        "就业流量",
        "劳工",
    ),
    _indicator(
        "尼泊尔_DoFE批准_不含再入境",
        "月",
        "人",
        "尼泊尔 DoFE monthly final labour approval",
        "就业流量",
        "劳工",
    ),
    _indicator(
        "孟加拉国_BMET出境许可",
        "月",
        "人",
        "孟加拉国 BMET/OEP Country Clearance",
        "就业流量",
        "劳工",
    ),
    _indicator(
        "重点三国合计_可比月",
        "月",
        "人",
        "上述三国官方数据之和，仅三国均有值时计算",
        "就业流量",
        "劳工",
    ),
    _indicator(
        "阿联酋:外籍劳动力:三国综合代理指数",
        "月",
        "指数",
        "上述三国官方数据标准化代理指标",
        "代理指标",
        "劳工",
    ),
    _indicator(
        "阿联酋:外籍劳动力:来源覆盖数",
        "月",
        "个",
        "上述三国官方数据覆盖计数",
        "质量指标",
        "劳工",
    ),
    _indicator(
        "迪拜:DED:当月快照记录数",
        "月",
        "条",
        "data.dubai Commerce Registry 原始快照",
        "记录数",
        "商业注册",
    ),
    _indicator(
        "迪拜:当月有发证活动的企业数(企业号筛重)",
        "月",
        "家",
        "data.dubai Commerce Registry (commerce_number 按发照月去重)",
        "企业数",
        "商业注册",
    ),
    _indicator(
        "迪拜:当月新发执照数(执照号筛重)",
        "月",
        "张",
        "data.dubai Commerce Registry (main_license_number 按发照月去重)",
        "执照数",
        "商业注册",
    ),
    _indicator(
        "迪拜:新发执照数(官方口径)",
        "月",
        "张",
        "DET/迪拜媒体办新闻稿",
        "执照数",
        "商业注册",
    ),
    _indicator(
        "阿联酋:Google搜索热度(work in dubai)",
        "月",
        "指数",
        "Google Trends/工作搜索热度 CSV",
        "搜索热度",
        "劳工",
    ),
    _indicator(
        "阿联酋:Google搜索热度(work in uae)",
        "月",
        "指数",
        "Google Trends/工作搜索热度 CSV",
        "搜索热度",
        "劳工",
    ),
    # 这些 Wind 指标早于统一 DuckDB 元数据表，曾只在工作表第 5 行保留
    # 来源，导致“指标字典”同步时没有可回填的 source。
    _indicator(
        "阿联酋DFM综合股票指数",
        "日",
        "点",
        "迪拜金融市场",
        "指数",
        "金融",
    ),
    _indicator(
        "全球:现货均价:原油(阿联酋穆尔班)",
        "周",
        "美元/桶",
        "金联创",
        "价格",
        "能源",
    ),
    _indicator(
        "期货结算价(连续):布伦特原油",
        "日",
        "美元/桶",
        "ICE",
        "价格",
        "能源",
    ),
    _indicator(
        "全球:现货价:原油(英国布伦特Dtd)",
        "日",
        "美元/桶",
        "金联创",
        "价格",
        "能源",
    ),
    _indicator(
        "全球:现货价:原油(阿联酋迪拜)",
        "日",
        "美元/桶",
        "金联创",
        "价格",
        "能源",
    ),
    _indicator(
        "阿联酋:现货价(CIF,低端价):石脑油",
        "日",
        "美元/吨",
        "隆众资讯",
        "价格",
        "能源",
    ),
    _indicator(
        "阿联酋:现货价(CIF,低端价):汽油(优质无铅)",
        "日",
        "美元/吨",
        "隆众资讯",
        "价格",
        "能源",
    ),
    _indicator(
        "阿联酋:现货价(CIF,低端价):柴油(Gasoil.1)",
        "日",
        "美元/吨",
        "隆众资讯",
        "价格",
        "能源",
    ),
    _indicator(
        "阿联酋:现货价(CIF,低端价):煤油(航空)",
        "日",
        "美元/吨",
        "隆众资讯",
        "价格",
        "能源",
    ),
    _indicator(
        "中国:广州市场:市场价(主流价):线型低密度聚乙烯(FB2230膜料):阿联酋博禄化工",
        "日",
        "美元/吨",
        "隆众资讯",
        "价格",
        "化工",
    ),
    _indicator(
        "阿联酋:现货价(CIF,低端价):柴油(10ppm)",
        "日",
        "美元/吨",
        "隆众资讯",
        "价格",
        "能源",
    ),
    _indicator(
        "印度:出口数量:咖啡:阿联酋:累计值",
        "日",
        "公吨",
        "印度咖啡委员会",
        "物量",
        "农产品",
    ),
    _indicator(
        "阿联酋:银行间同业拆借利率(EIBOR):1年",
        "月",
        "%",
        "阿联酋央行",
        "利率",
        "金融",
    ),
    _indicator(
        "美元兑阿联酋迪拉姆",
        "日",
        None,
        "阿联酋央行",
        "汇率",
        "汇率",
    ),
    _indicator(
        "人民币兑阿联酋迪拉姆",
        "日",
        None,
        "阿联酋央行",
        "汇率",
        "汇率",
    ),
    _indicator(
        "阿联酋:银行间同业拆借利率(EIBOR):隔夜",
        "月",
        None,
        "阿联酋央行",
        "利率",
        "金融",
    ),
    # PortWatch 的月度宽表由日度明细聚合生成，13 个指标共用同一来源。
    _indicator(
        "阿联酋:港口到港总次数:当月值",
        "月",
        "艘次",
        "IMF PortWatch (HDX mirror)",
        "数量",
        "海运",
    ),
    _indicator(
        "阿联酋:港口集装箱到港次数:当月值",
        "月",
        "艘次",
        "IMF PortWatch (HDX mirror)",
        "数量",
        "海运",
    ),
    _indicator(
        "阿联酋:港口油轮到港次数:当月值",
        "月",
        "艘次",
        "IMF PortWatch (HDX mirror)",
        "数量",
        "海运",
    ),
    _indicator(
        "阿联酋:港口进口总量:当月值",
        "月",
        "吨",
        "IMF PortWatch (HDX mirror)",
        "数量",
        "海运",
    ),
    _indicator(
        "阿联酋:港口集装箱进口量:当月值",
        "月",
        "吨",
        "IMF PortWatch (HDX mirror)",
        "数量",
        "海运",
    ),
    _indicator(
        "阿联酋:港口油轮进口量:当月值",
        "月",
        "吨",
        "IMF PortWatch (HDX mirror)",
        "数量",
        "海运",
    ),
    _indicator(
        "阿联酋:港口出口总量:当月值",
        "月",
        "吨",
        "IMF PortWatch (HDX mirror)",
        "数量",
        "海运",
    ),
    _indicator(
        "阿联酋:港口集装箱出口量:当月值",
        "月",
        "吨",
        "IMF PortWatch (HDX mirror)",
        "数量",
        "海运",
    ),
    _indicator(
        "阿联酋:港口油轮出口量:当月值",
        "月",
        "吨",
        "IMF PortWatch (HDX mirror)",
        "数量",
        "海运",
    ),
    _indicator(
        "霍尔木兹:过境总次数:当月值",
        "月",
        "艘次",
        "IMF PortWatch (HDX mirror)",
        "数量",
        "海运",
    ),
    _indicator(
        "霍尔木兹:载货容量:当月值",
        "月",
        "吨",
        "IMF PortWatch (HDX mirror)",
        "数量",
        "海运",
    ),
    _indicator(
        "霍尔木兹:油轮过境次数:当月值",
        "月",
        "艘次",
        "IMF PortWatch (HDX mirror)",
        "数量",
        "海运",
    ),
    _indicator(
        "霍尔木兹:油轮载货容量:当月值",
        "月",
        "吨",
        "IMF PortWatch (HDX mirror)",
        "数量",
        "海运",
    ),
)


# Raw/detail fields are not business indicators.  They belong in the column
# dictionary so that dimensions such as reporter, partner, portid, and
# transaction_id are documented without becoming selectable workbook series.
_COLUMN_MEANINGS: dict[str, tuple[str, str, str | None, str | None]] = {
    "period": ("时间", "观测期或月末日期", None, None),
    "month": ("时间", "观测月份或月末日期", None, None),
    "date": ("时间", "观测日期", None, None),
    "series": ("指标维度", "长表中的原始序列键", "对应 value 数值列", "序列键保留源脚本名称"),
    "value": ("数值", "指标观测值", None, "缺失值保留为 NULL 或缺行"),
    "quality_flag": ("质量标记", "源数据质量或不可差分原因", None, "文本字段，不作为数值指标"),
    "keyword": ("搜索维度", "Google Trends 搜索关键词", "employment_search_index.value", None),
    "legal_form": ("法律形式", "企业法律形式分类", "ded_monthly_by_type.records", None),
    "records": ("记录数", "该月份/法律形式的原始记录数", None, "缺行不等于零"),
    "category": ("分类维度", "资本品或设备类别", None, None),
    "reporter": ("申报国", "UN Comtrade 申报国", None, None),
    "partner": ("贸易伙伴", "UN Comtrade 贸易伙伴", None, None),
    "hs6": ("HS6 分类", "UN Comtrade 六位商品编码", None, None),
    "item_count": ("数量", "HS6 商品数量", None, "单位随源文件口径，需结合 source_type 解读"),
    "source_type": ("来源口径", "直报或镜像来源类型", None, None),
    "reported_estimated": ("估算标记", "直报/估算状态标志", None, None),
    "mirror_reporter_count": ("镜像覆盖", "参与镜像汇总的申报国数量", None, None),
    "dataset": ("数据集", "Eurostat 数据集代码", None, None),
    "geo": ("地理维度", "Eurostat 地理实体代码", None, None),
    "schedule": ("航班计划", "Eurostat 航班计划口径", None, None),
    "unit": ("单位维度", "Eurostat 原始单位代码", None, None),
    "tra_meas": ("运输度量", "Eurostat 运输度量代码", None, None),
    "portid": ("港口/海峡标识", "PortWatch 港口或海峡 ID", None, None),
    "portname": ("港口/海峡名称", "PortWatch 港口或海峡名称", None, None),
    "country": ("国家维度", "港口所在国家名称", None, None),
    "iso3": ("国家代码", "港口所在国家 ISO3 代码", None, None),
    "transaction_id": ("主键", "交易记录唯一编号", None, "由官方源保证唯一"),
    "instance_date": ("时间", "交易登记日期", None, "保留早于 1900 年的异常日期供质量视图审计"),
    "trans_group_en": ("交易分类", "Sales/Mortgages/Gifts 等交易大类", None, None),
    "reg_type_en": ("登记类型", "Existing 或 Off-Plan 登记类型", None, None),
    "project_number": ("项目维度", "交易数据中的项目编号", None, "允许 NULL"),
    "actual_worth": ("金额面积", "交易金额或实际价值", None, "通常按 AED 理解；源文件未单列单位"),
    "property_type_en": ("房产分类", "房产大类英文名", None, None),
    "property_usage_en": ("用途分类", "住宅/商业/酒店/其他用途英文名", None, None),
    "load_timestamp": ("时间", "Data Dubai 生成 bulk 快照的时间", None, "不是交易发生时间"),
}

_DETAIL_OBJECTS = (
    "foreign_labour_monthly",
    "ded_monthly",
    "ded_monthly_by_type",
    "employment_search_index",
    "comtrade_partner_detail",
    "comtrade_quantity_detail",
    "eurostat_air_monthly",
    "portwatch_uae_daily",
    "portwatch_chokepoint_daily",
    "dld.transactions",
)

_PORTWATCH_VALUE_PREFIXES = {
    "portcalls": "港口挂靠次数",
    "import": "进口货运量",
    "export": "出口货运量",
    "n_": "海峡过境次数",
    "capacity": "海峡载货容量",
}


def _portwatch_meaning(column_name: str) -> tuple[str, str, str | None, str | None]:
    for prefix, meaning in _PORTWATCH_VALUE_PREFIXES.items():
        if column_name.startswith(prefix):
            return ("运输流量", meaning, None, "按货种后缀拆分")
    return _COLUMN_MEANINGS.get(
        column_name,
        ("原始字段", f"PortWatch 原始字段 {column_name}", None, None),
    )


def _column_rows(con) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for object_name in _DETAIL_OBJECTS:
        if "." in object_name:
            schema, table = object_name.split(".", 1)
        else:
            schema, table = "main", object_name
        columns = con.execute(
            "SELECT column_name, data_type "
            "FROM information_schema.columns "
            "WHERE table_schema = ? AND table_name = ? "
            "ORDER BY ordinal_position",
            [schema, table],
        ).fetchall()
        for column_name, data_type in columns:
            if table.startswith("portwatch_"):
                group, meaning, relates_to, caveat = _portwatch_meaning(column_name)
            else:
                group, meaning, relates_to, caveat = _COLUMN_MEANINGS.get(
                    column_name,
                    ("原始字段", f"数据库字段 {column_name}", None, None),
                )
            rows.append(
                {
                    "object_name": object_name,
                    "column_name": column_name,
                    "data_type": data_type,
                    "group_cn": group,
                    "meaning_cn": meaning,
                    "relates_to": relates_to,
                    "caveat_cn": caveat,
                }
            )
    return rows


def complete_metadata(con, *, updated_at: date | None = None) -> dict[str, int]:
    """Upsert all extended indicator and detail-column metadata.

    The operation is idempotent and only changes metadata tables.  It does not
    create or rewrite observation rows.
    """

    db.init_schema(con)
    stamp = updated_at or date.today()
    indicator_rows = [dict(row, updated_at=stamp) for row in INDICATOR_DEFINITIONS]
    column_rows = _column_rows(con)
    con.execute("BEGIN TRANSACTION")
    try:
        db.upsert_dictionary_rows(con, indicator_rows)
        for row in column_rows:
            con.execute(
                """
                INSERT OR REPLACE INTO meta_column_dictionary
                    (object_name, column_name, data_type, group_cn, meaning_cn,
                     relates_to, caveat_cn)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    row["object_name"],
                    row["column_name"],
                    row["data_type"],
                    row["group_cn"],
                    row["meaning_cn"],
                    row["relates_to"],
                    row["caveat_cn"],
                ],
            )
        con.execute("COMMIT")
    except BaseException:
        con.execute("ROLLBACK")
        raise
    return {"indicator_rows": len(indicator_rows), "column_rows": len(column_rows)}


def _row_values(worksheet, row_number: int) -> list[Any]:
    return [cell.value for cell in worksheet[row_number]]


def _workbook_indicator_names(workbook) -> set[str]:
    names: set[str] = set()
    for worksheet in workbook.worksheets[1:]:
        for row_number in (1, 2):
            values = _row_values(worksheet, row_number)
            if not values:
                continue
            first = str(values[0]).strip() if values[0] is not None else ""
            if first not in {"指标名称", "日期"}:
                continue
            for value in values[1:]:
                if isinstance(value, str) and value.strip():
                    names.add(value.strip())
    return names


def sync_workbook_dictionary(workbook_path: Path) -> dict[str, int]:
    """Sync DB metadata for indicators present in workbook data sheets.

    Dictionary rows without a matching indicator column in any data sheet are
    removed. Existing sheet-backed rows receive authoritative
    type/industry/source values; sheet-backed DB indicators absent from the
    dictionary are appended.
    """

    workbook_path = Path(workbook_path).resolve()
    if not workbook_path.is_file():
        raise FileNotFoundError(workbook_path)
    lock_path = workbook_path.parent / f"~${workbook_path.name}"
    if lock_path.exists():
        raise RuntimeError(f"请先关闭 Excel 工作簿：{lock_path}")

    con = db.connect(read_only=True)
    try:
        rows = con.execute(
            """
            SELECT indicator_name, type, industry, source
            FROM meta_indicator_dictionary
            ORDER BY indicator_name
            """
        ).fetchall()
    finally:
        con.close()

    workbook = load_workbook(workbook_path, keep_links=True)
    dictionary = workbook.worksheets[0]
    headers = [dictionary.cell(1, column).value for column in range(1, 6)]
    if headers != ["指标名称", "类型", "行业", "数据来源", "预测变量"]:
        raise ValueError("指标字典表头不符合工作簿协议")

    sheet_names = _workbook_indicator_names(workbook)
    dictionary_rows: dict[str, int] = {}
    for row_number in range(2, dictionary.max_row + 1):
        value = dictionary.cell(row_number, 1).value
        if value is None or not str(value).strip():
            continue
        name = str(value).strip()
        if name in dictionary_rows:
            raise ValueError(f"指标字典包含重复指标：{name}")
        dictionary_rows[name] = row_number

    orphan_rows = sorted(
        (row_number for name, row_number in dictionary_rows.items() if name not in sheet_names),
        reverse=True,
    )
    for row_number in orphan_rows:
        dictionary.delete_rows(row_number, 1)
    removed = len(orphan_rows)

    dictionary_rows = {}
    for row_number in range(2, dictionary.max_row + 1):
        value = dictionary.cell(row_number, 1).value
        if value is None or not str(value).strip():
            continue
        dictionary_rows[str(value).strip()] = row_number

    updated = 0
    added = 0
    for name, type_, industry, source in rows:
        if name not in sheet_names:
            continue
        row_number = dictionary_rows.get(name)
        if row_number is None:
            row_number = dictionary.max_row + 1
            dictionary_rows[name] = row_number
            if row_number > 2:
                source_row = row_number - 1
                for column in range(1, dictionary.max_column + 1):
                    source_cell = dictionary.cell(source_row, column)
                    target_cell = dictionary.cell(row_number, column)
                    target_cell._style = copy.copy(source_cell._style)
                    if source_cell.has_style:
                        target_cell.number_format = source_cell.number_format
                    if source_cell.alignment:
                        target_cell.alignment = copy.copy(source_cell.alignment)
                    if source_cell.protection:
                        target_cell.protection = copy.copy(source_cell.protection)
            added += 1
        if (
            dictionary.cell(row_number, 2).value != type_
            or dictionary.cell(row_number, 3).value != industry
            or dictionary.cell(row_number, 4).value != source
        ):
            updated += 1
        dictionary.cell(row_number, 1).value = name
        dictionary.cell(row_number, 2).value = type_
        dictionary.cell(row_number, 3).value = industry
        dictionary.cell(row_number, 4).value = source

    temporary = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{workbook_path.stem}.dictionary-",
            suffix=".tmp.xlsx",
            dir=workbook_path.parent,
            delete=False,
        ) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        os.replace(temporary, workbook_path)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()
        workbook.close()
    return {
        "added": added,
        "updated": updated,
        "removed": removed,
        "sheet_indicators": len(sheet_names),
    }


if __name__ == "__main__":
    connection = db.connect()
    try:
        result = complete_metadata(connection)
    finally:
        connection.close()
    print(
        "[uae_metadata] completed: "
        f"{result['indicator_rows']} indicator rows, "
        f"{result['column_rows']} column rows"
    )
