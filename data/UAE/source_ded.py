"""Data Dubai DED 企业注册主表数据源：下载 → 月度聚合 → 入库 → 回写工作簿。

逻辑源自 data/ded_license/fetch_ded_license.py（下载+首现筛重聚合）与
data/ded_license/build_monthly_ded_sheet.py（合并写入月度_DED），整体移植。
唯一介质差异：
- 输入/审计输出目录改为 data/UAE/raw/ded/DED_data/（备份在 data/UAE/raw/ded/backups/）；
- 聚合结果先入 DuckDB（ded_monthly / ded_monthly_by_type），
  merge() 再从库查询回写「月度_DED」sheet。

口径保持与旧脚本 100% 一致：
- records      ：快照记录行数（按发照日期归月）；
- enterprises  ：新增企业数（commerce_number 首次出现月份，剔除占位 '0'/空）；
- licences     ：新增执照数（main_license_number 首次出现月份，同号筛除）；
- official     ：官方公布数字（official_ded_monthly.csv，人工维护）；
- 单期存量快照无法给出流量，历史月份低估、近期可信（与旧脚本同样口径）。
"""

from __future__ import annotations

import bz2
import calendar
import csv
import gzip
import hashlib
import http.cookiejar
import io
import json
import re
import shutil
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter
from contextlib import contextmanager
from datetime import datetime, timedelta
from pathlib import Path

try:  # Windows 控制台 GBK 下避免中文乱码
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

DATA_DIR = Path(__file__).resolve().parent
if str(DATA_DIR) not in sys.path:
    sys.path.insert(0, str(DATA_DIR))

from db import (  # noqa: E402
    connect,
    replace,
)

DATA_DUBAI = "https://data.dubai"
DATASET_PAGE = (
    "https://data.dubai/en/l/460521?com_dda_data_and_statistics_ThemeId=36962"
)
DATASET_ID = "460521"
DOWNLOAD_API = f"{DATA_DUBAI}/o/dda/data-services/dataset-download"
UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 Chrome/151 Safari/537.36"
)
CHUNK_SIZE = 1024 * 1024
RETRYABLE_STATUS = {429, 500, 502, 503, 504}

# 输入与审计输出固定在 data/UAE/raw/ded/DED_data，不跟随 cwd
RAW_DED = DATA_DIR / "raw" / "ded"
OUT = RAW_DED / "DED_data"
RAW = OUT / "raw"
OFFICIAL_CSV = RAW_DED / "official_ded_monthly.csv"
BACKUP_DIR = RAW_DED / "backups"

SHEET_NAME = "月度_DED"
SOURCE_NAME = "DED"

import openpyxl  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

HEADER_FONT = Font(name="等线", size=11, bold=True)
BODY_FONT = Font(name="等线", size=11)

MONTHLY_TABLE = "ded_monthly"
BY_TYPE_TABLE = "ded_monthly_by_type"


class AuthRequiredError(RuntimeError):
    """资源需要登录（UAE Pass / Dubai ID），携带会话后重试。"""


class IncompleteDownloadError(RuntimeError):
    """下载字节数与 Content-Length 不符（连接被截断），需要重试。"""


# --------------------------------------------------------------------------
# 下载基础设施（与 fetch_ded_license.py 一致）
# --------------------------------------------------------------------------

def format_bytes(size):
    size = int(size or 0)
    value = float(size)
    for unit in ("B", "KB", "MB", "GB", "TB"):
        if value < 1024 or unit == "TB":
            return f"{value:.1f} {unit}"
        value /= 1024
    return "0 B"


def sha256_file(path):
    digest = hashlib.sha256()
    with open(path, "rb") as source:
        while True:
            block = source.read(CHUNK_SIZE)
            if not block:
                break
            digest.update(block)
    return digest.hexdigest()


def parse_netscape_cookies(path):
    """读取 Netscape cookies.txt（curl -c / 浏览器插件导出），返回 CookieJar。"""
    jar = http.cookiejar.CookieJar()
    with open(path, "r", encoding="utf-8", errors="replace") as source:
        for line in source:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            parts = line.split("\t")
            if len(parts) < 7:
                continue
            domain, _, path_part, secure, expires, name, value = parts[:7]
            try:
                secure = secure.lower() == "true"
                expires = int(expires) if expires and expires != "0" else None
            except ValueError:
                continue
            jar.set_cookie(http.cookiejar.Cookie(
                version=0, name=name, value=value,
                port=None, port_specified=False,
                domain=domain, domain_specified=bool(domain),
                domain_initial_dot=domain.startswith("."),
                path=path_part or "/", path_specified=bool(path_part),
                secure=secure, expires=expires, discard=False,
                comment=None, comment_url=None, rest={},
            ))
    return jar


def build_opener(cookies_file=None):
    jar = http.cookiejar.CookieJar()
    if cookies_file:
        jar = parse_netscape_cookies(cookies_file)
    return urllib.request.build_opener(
        urllib.request.HTTPCookieProcessor(jar),
        urllib.request.HTTPRedirectHandler(),
    )


def make_request(url, headers=None):
    if isinstance(url, urllib.request.Request):
        return url
    headers = dict(headers or {})
    headers.setdefault("User-Agent", UA)
    return urllib.request.Request(url, headers=headers)


def raise_if_sso_url(response):
    """最终 URL 被 SSO 拦截时抛出 AuthRequiredError。"""
    final = getattr(response, "url", "") or response.geturl()
    if "smartsso.dubai.gov.ae" in final or "/isam/sps/" in final:
        raise AuthRequiredError(
            "请求被重定向到迪拜政府统一登录（smartsso.dubai.gov.ae，需要登录）。"
            "处理方式：浏览器登录后，把会话 Cookie 通过 --header \"Cookie: ...\" 传入；"
            "或把数据文件直链用 --url 直接指定。详见 README.md。"
        )
    return response


def raise_if_sso_page(text):
    """HTML 内容为 SSO 登录页时抛出 AuthRequiredError。"""
    if "smartsso.dubai.gov.ae" in text or "/isam/sps/" in text:
        raise AuthRequiredError(
            "返回的是登录页（smartsso.dubai.gov.ae），该资源需要登录。"
            "详见 README.md 的认证说明。"
        )


def read_text(opener, url, headers=None, timeout=90, retries=4):
    req = make_request(url, headers)
    last = None
    for attempt in range(retries + 1):
        try:
            with opener.open(req, timeout=timeout) as response:
                raise_if_sso_url(response)
                text = response.read().decode("utf-8", errors="replace")
                raise_if_sso_page(text)
                return text
        except AuthRequiredError:
            raise
        except urllib.error.HTTPError as exc:
            last = exc
            if exc.code in RETRYABLE_STATUS and attempt < retries:
                time.sleep(max(int(exc.headers.get("Retry-After", 0) or 0), 2 ** attempt))
                continue
            raise
        except (urllib.error.URLError, TimeoutError, OSError) as exc:
            last = exc
            if attempt < retries:
                time.sleep(2 ** attempt)
                continue
            raise
    raise last


def stream_to_file(opener, url, dest, headers=None, timeout=180, retries=4):
    """流式下载到 dest，返回 (最终响应 URL, Content-Type)。"""
    req = make_request(url, headers)
    last = None
    for attempt in range(retries + 1):
        try:
            with opener.open(req, timeout=timeout) as response:
                raise_if_sso_url(response)
                final_url = response.geturl()
                ctype = response.headers.get("Content-Type", "")
                expected = int(response.headers.get("Content-Length") or 0)
                total = 0
                last_report = time.time()
                head = response.read(CHUNK_SIZE)
                if b"smartsso.dubai.gov.ae" in head or b"/isam/sps/" in head:
                    raise AuthRequiredError(
                        "返回的是登录页（smartsso.dubai.gov.ae），该资源需要登录。"
                        "详见 README.md 的认证说明。"
                    )
                with open(dest, "wb") as target:
                    if head:
                        target.write(head)
                        total += len(head)
                    while True:
                        block = response.read(CHUNK_SIZE)
                        if not block:
                            break
                        target.write(block)
                        total += len(block)
                        now = time.time()
                        if now - last_report >= 10:  # 每 10 秒报一次进度
                            print(f"  ... {format_bytes(total)}", flush=True)
                            last_report = now
                if expected and total != expected:
                    raise IncompleteDownloadError(
                        f"下载不完整: 期望 {expected} 字节，实际 {total} 字节"
                    )
                return final_url, ctype
        except IncompleteDownloadError as exc:
            last = exc
            print(f"  ! {exc}，重试 ({attempt + 1}/{retries + 1})", flush=True)
            dest.unlink(missing_ok=True)
            if attempt < retries:
                time.sleep(2 ** attempt)
                continue
            raise
        except AuthRequiredError:
            dest.unlink(missing_ok=True)
            raise
        except urllib.error.HTTPError as exc:
            last = exc
            if exc.code in RETRYABLE_STATUS and attempt < retries:
                time.sleep(max(int(exc.headers.get("Retry-After", 0) or 0), 2 ** attempt))
                continue
            dest.unlink(missing_ok=True)
            raise
        except (urllib.error.URLError, TimeoutError, OSError) as exc:
            last = exc
            if attempt < retries:
                time.sleep(2 ** attempt)
                continue
            dest.unlink(missing_ok=True)
            raise
    raise last


def discover_legacy_datafile(page_url, page_html, prefer="csv"):
    """旧门户 dubaipulse 页面：解析 Data Files 区块的数据文件直链。"""
    hrefs = re.findall(r'href="([^"]*datafiles/[^"]*)"', page_html, re.I)
    if not hrefs:
        return None
    urls = [urllib.parse.urljoin(page_url, h) for h in dict.fromkeys(hrefs)]

    def score(u):
        low = u.lower()
        if prefer == "json":
            score = 0 if ".json" in low else 2
        else:
            score = 0 if (".csv" in low or ".csv.gz" in low) else 2
        return score

    urls.sort(key=score)
    return urls[0]


def fetch_data_dubai_files(opener, dataset_id, headers=None, prefer="csv"):
    """通过官方 dataset-download 接口获取全部快照文件（签名直链）。

    返回 [{"file_name":..., "file_extension":..., "file_size":..., "file_url":...}]
    """
    page_url = (
        f"{DATA_DUBAI}/en/l/{dataset_id}"
        "?com_dda_data_and_statistics_ThemeId=36962"
    )
    print(f"抓取数据集页面: {page_url}")
    html_text = read_text(opener, page_url, headers, timeout=60, retries=2)
    token_match = re.search(r"Liferay\.authToken\s*=\s*'([^']+)'", html_text)
    if not token_match:
        raise RuntimeError("Data Dubai 页面未提供 CSRF token")

    files = []
    page = 1
    while True:
        query = urllib.parse.urlencode({
            "datasetId": dataset_id,
            "page": str(page),
            "pageSize": "100",
            "sortDir": "desc",
        })
        api_req = urllib.request.Request(
            f"{DOWNLOAD_API}?{query}",
            headers={
                "User-Agent": UA,
                "Accept": "application/json",
                "X-CSRF-Token": token_match.group(1),
                "Referer": page_url,
            },
        )
        payload = json.loads(read_text(opener, api_req, headers={}, timeout=60, retries=2))
        if not payload.get("success"):
            raise RuntimeError(f"dataset-download 接口失败: {payload.get('message')}")
        data = payload.get("data") or {}
        for folder in data.get("metadata") or []:
            for item in folder.get("files") or []:
                name = str(item.get("file_name") or "")
                ext = str(item.get("file_extension") or "").lower()
                if prefer == "json":
                    pick = ext == "json"
                else:
                    pick = ext == "csv"
                if pick and item.get("file_url"):
                    files.append({
                        "file_name": name,
                        "file_extension": ext,
                        "file_size": int(item.get("file_size") or 0),
                        "file_url": item["file_url"],
                    })
        pagination = data.get("pagination") or {}
        page_count = int(pagination.get("page_count") or 1)
        if page >= page_count:
            break
        page += 1
    if not files:
        raise RuntimeError(f"数据集 {dataset_id} 没有可用的 CSV/JSON 文件")
    return files


def detect_format(path):
    name = path.name.lower()
    if name.endswith(".gz") or name.endswith(".bz2"):
        name = name.rsplit(".", 1)[0]
    return "json" if name.endswith(".json") else "csv"


def open_text(path):
    """打开（必要时解压）文本流，自动探测编码与换行。"""
    raw = open(path, "rb")
    if path.name.lower().endswith(".gz"):
        raw = gzip.GzipFile(fileobj=raw)
    elif path.name.lower().endswith(".bz2"):
        raw = bz2.BZ2File(raw)
    sample = raw.read(65536)
    raw.seek(0)
    encoding = "utf-8"
    for candidate in ("utf-8-sig", "utf-8", "cp1256", "latin-1"):
        try:
            sample.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    return io.TextIOWrapper(raw, encoding=encoding, errors="replace", newline="")


def sniff_delimiter(text):
    first = text.readline()
    text.seek(0)
    counts = {d: first.count(d) for d in ("\t", ";", "|", ",")}
    delim = max(counts, key=counts.get)
    return delim if counts[delim] > 0 else ","


def iter_rows(paths, limit=None):
    """按行产出 dict；paths 可为单文件或多文件列表（多分片）。"""
    if isinstance(paths, (str, Path)):
        paths = [paths]
    produced = 0
    for path in paths:
        fmt = detect_format(path)
        if fmt == "csv":
            text = open_text(path)
            delim = sniff_delimiter(text)
            reader = csv.DictReader(text, delimiter=delim)
            for row in reader:
                if row and any(v is not None and v != "" for v in row.values()):
                    yield row
                    produced += 1
                    if limit and produced >= limit:
                        return
            text.close()
        else:
            size = path.stat().st_size
            text = open_text(path)
            if size < 500 * 1024 * 1024:
                try:
                    payload = json.load(text)
                except json.JSONDecodeError:
                    text.seek(0)
                    for line in text:
                        line = line.strip()
                        if not line:
                            continue
                        yield json.loads(line)
                        produced += 1
                        if limit and produced >= limit:
                            return
                    continue
                if isinstance(payload, dict):
                    for key in ("data", "records", "results", "items", "rows"):
                        if isinstance(payload.get(key), list):
                            payload = payload[key]
                            break
                    else:
                        payload = [payload]
                if not isinstance(payload, list):
                    raise ValueError(f"无法识别的 JSON 结构: {path.name}")
                for row in payload:
                    if isinstance(row, dict):
                        yield row
                        produced += 1
                        if limit and produced >= limit:
                            return
            else:
                for line in text:
                    line = line.strip()
                    if not line:
                        continue
                    yield json.loads(line)
                    produced += 1
                    if limit and produced >= limit:
                        return
            text.close()


DATE_PATTERNS = [
    re.compile(r"issue.*date", re.I),
    re.compile(r"licen[cs]e.*date", re.I),
    re.compile(r"date.*issu", re.I),
    re.compile(r"^issued", re.I),
    re.compile(r"registr.*date", re.I),
    re.compile(r"establishment.*date", re.I),
    re.compile(r"^created", re.I),
    re.compile(r"^date$", re.I),
    re.compile(r"date", re.I),
]
TYPE_PATTERNS = [
    re.compile(r"legal_type.*desc_en", re.I),
    re.compile(r"commerce_register_type.*desc_en", re.I),
    re.compile(r"licen[cs]e.*type", re.I),
    re.compile(r"^type", re.I),
    re.compile(r"^activit", re.I),
    re.compile(r"sector", re.I),
    re.compile(r"categor", re.I),
]


def guess_column(columns, patterns):
    normalized = {c: re.sub(r"[^a-z0-9]+", "_", c.lower()).strip("_") for c in columns}
    for pattern in patterns:
        for col in columns:
            if pattern.search(normalized[col]):
                return col
    return None


def parse_date(value):
    """解析常见日期格式，返回 (year, month) 或 None。"""
    if value is None:
        return None
    s = str(value).strip().strip('"')
    if not s:
        return None
    if s.isdigit():
        n = int(s)
        if n > 1e12:
            n /= 1000.0
        if 1e9 < n < 4.1e9:
            dt = datetime.fromtimestamp(n)
            return dt.year, dt.month
    date_part = re.split(r"[ T]", s.strip())[0][:10]
    if len(date_part) == 8 and date_part.isdigit():
        try:
            dt = datetime.strptime(date_part, "%Y%m%d")
            return dt.year, dt.month
        except ValueError:
            return None
    formats = ["%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"]
    if "/" in date_part:
        formats += ["%d/%m/%Y", "%m/%d/%Y"]
    elif "-" in date_part:
        formats += ["%d-%m-%Y"]
    elif "." in date_part:
        formats += ["%d.%m.%Y"]
    for fmt in formats:
        try:
            dt = datetime.strptime(date_part, fmt)
            return dt.year, dt.month
        except ValueError:
            continue
    return None


def aggregate(rows, date_col, type_col):
    """月度聚合（与 fetch_ded_license.py 完全相同）。

    - monthly            ：记录行数（每条登记记录一张执照，按发照日期归月）
    - monthly_enterprise ：新增企业数——企业号(commerce_number)在快照中首次
                           出现的月份（剔除占位值 '0'/空）
    - monthly_licence    ：新增执照数——执照号(main_license_number)首次出现的月份
    """
    total = parsed = bad = 0
    monthly = Counter()
    ent_first = {}    # commerce_number -> 最早有效发照月份
    lic_first = {}    # main_license_number -> 最早有效发照月份
    by_type = Counter()
    date_min = date_max = None
    sample_rows = []
    for row in rows:
        total += 1
        if len(sample_rows) < 3:
            sample_rows.append(row)
        month_key = None
        if date_col:
            result = parse_date(row.get(date_col))
            if result and 1900 <= result[0] <= 2100:  # 过滤占位/脏年份
                y, m = result
                parsed += 1
                month_key = f"{y:04d}-{m:02d}"
                monthly[month_key] += 1
                date_min = month_key if date_min is None else min(date_min, month_key)
                date_max = month_key if date_max is None else max(date_max, month_key)
                ent = str(row.get("commerce_number") or "").strip()
                if ent and ent != "0":
                    first = ent_first.get(ent)
                    if first is None or month_key < first:
                        ent_first[ent] = month_key
                lic = str(row.get("main_license_number") or "").strip()
                if lic:
                    first = lic_first.get(lic)
                    if first is None or month_key < first:
                        lic_first[lic] = month_key
            else:
                bad += 1
        if type_col:
            tval = str(row.get(type_col) or "").strip()
            by_type[(month_key or "unknown", tval or "(空)")] += 1
    monthly_enterprise = Counter(ent_first.values())
    monthly_licence = Counter(lic_first.values())
    return {
        "total": total,
        "parsed": parsed,
        "bad": bad,
        "monthly": monthly,
        "monthly_enterprise": monthly_enterprise,
        "monthly_licence": monthly_licence,
        "enterprise_total": len(ent_first),
        "licence_total": len(lic_first),
        "by_type": by_type,
        "date_min": date_min,
        "date_max": date_max,
        "sample_rows": sample_rows,
    }


def aggregate_snapshots(paths, *, date_col=None, type_col=None, limit=None):
    """读取一个或多个快照文件并完成列识别与月度聚合（纯函数）。

    返回 {"columns", "date_col", "type_col", "stats", "warnings"}。
    供 update() 与测试复用；结构与旧 main 的处理段一致。
    """

    rows_iter = iter_rows(paths, limit=limit)
    first = next(rows_iter, None)
    if first is None:
        raise ValueError("数据文件为空，无法处理。")
    columns_seen = list(first.keys())
    used_date_col = date_col or guess_column(columns_seen, DATE_PATTERNS)
    used_type_col = None if type_col is False else (type_col or guess_column(columns_seen, TYPE_PATTERNS))

    def all_rows():
        yield first
        yield from rows_iter

    stats = aggregate(all_rows(), used_date_col, used_type_col)
    warnings = []
    if not used_date_col:
        warnings.append("未识别到日期列，无法按月聚合（可用 --date-col 指定）。")
    elif stats["bad"]:
        warnings.append(f"{stats['bad']:,} 行日期无法解析，已跳过（可用 --date-col 换列）。")
    if not used_type_col:
        warnings.append("未识别到类型列，跳过按类型明细。")
    return {
        "columns": columns_seen,
        "date_col": used_date_col,
        "type_col": used_type_col,
        "stats": stats,
        "warnings": warnings,
    }


# --------------------------------------------------------------------------
# 审计文件输出（与 fetch_ded_license.py 相同结构）
# --------------------------------------------------------------------------

def write_monthly(path, monthly):
    with open(path, "w", encoding="utf-8-sig", newline="") as target:
        writer = csv.writer(target)
        writer.writerow(["month", "new_licenses"])
        for key in sorted(monthly):
            writer.writerow([key, monthly[key]])


def write_monthly_counts(path, monthly, monthly_enterprise, monthly_licence):
    """三种口径的月度序列：记录数 / 新增企业数 / 新增执照数。"""
    months = sorted(set(monthly) | set(monthly_enterprise) | set(monthly_licence))
    with open(path, "w", encoding="utf-8-sig", newline="") as target:
        writer = csv.writer(target)
        writer.writerow(["month", "records", "enterprises", "licences"])
        for key in months:
            writer.writerow(
                [key, monthly.get(key, 0), monthly_enterprise.get(key, 0), monthly_licence.get(key, 0)]
            )


def write_monthly_by_type(path, by_type):
    with open(path, "w", encoding="utf-8-sig", newline="") as target:
        writer = csv.writer(target)
        writer.writerow(["month", "license_type", "count"])
        for (month, tval), count in sorted(by_type.items()):
            writer.writerow([month, tval, count])


def write_schema_report(path, columns, stats, date_col, type_col, warnings):
    lines = [
        "DED 企业注册主表 字段与数据质量报告",
        "=" * 50,
        f"总行数: {stats['total']:,}",
        (f"日期可解析: {stats['parsed']:,} "
         f"({(stats['parsed'] / stats['total'] * 100):.1f}%)")
        if stats["total"] else "日期可解析: 0",
        f"日期不可解析: {stats['bad']:,}",
        f"日期范围: {stats['date_min']} ~ {stats['date_max']}",
        f"日期列(自动识别): {date_col or '(未识别)'}",
        f"类型列(自动识别): {type_col or '(未识别)'}",
        "",
        "字段列表:",
    ]
    for col in columns:
        lines.append(f"  - {col}")
    if warnings:
        lines += ["", "警告:"]
        lines += [f"  - {w}" for w in warnings]
    lines += ["", "样例数据(前3行):"]
    for row in stats["sample_rows"]:
        lines.append("  " + json.dumps(row, ensure_ascii=False)[:500])
    with open(path, "w", encoding="utf-8", newline="") as target:
        target.write("\n".join(lines) + "\n")


def build_manifest(out_dir, **kwargs):
    manifest = {"generated_at": datetime.now().isoformat(timespec="seconds")}
    manifest.update(kwargs)
    with open(out_dir / "fetch_manifest.json", "w", encoding="utf-8") as target:
        json.dump(manifest, target, ensure_ascii=False, indent=2)


# --------------------------------------------------------------------------
# 官方数字与库行转换
# --------------------------------------------------------------------------

def _month_end_date(month_key: str) -> object:
    """'YYYY-MM' -> 月末 date（DuckDB DATE 列用）。"""
    year, month = int(month_key[:4]), int(month_key[5:7])
    return (
        datetime(year, month, calendar.monthrange(year, month)[1]).date()
    )


def read_official(path) -> dict[str, int]:
    """读官方 CSV -> {YYYY-MM: int}（与 build_monthly_ded_sheet.py 相同）。"""
    data = {}
    if not Path(path).exists():
        return data
    with open(path, "r", encoding="utf-8-sig") as source:
        reader = csv.DictReader(source)
        for row in reader:
            month = (row.get("month") or "").strip()
            try:
                data[month] = int(float(row["official_new_licenses"]))
            except (ValueError, TypeError, KeyError):
                continue
    return data


def _ded_monthly_rows(stats, official: dict[str, int]) -> list[dict[str, object]]:
    """把聚合结果与官方数字摊平成 ded_monthly 长表行。"""

    rows: list[dict[str, object]] = []
    series_map = {
        "records": stats["monthly"],
        "enterprises": stats["monthly_enterprise"],
        "licences": stats["monthly_licence"],
    }
    for series, counter in series_map.items():
        for month_key, value in sorted(counter.items()):
            rows.append(
                {
                    "month": _month_end_date(month_key),
                    "series": series,
                    "value": float(value),
                }
            )
    for month_key, value in sorted(official.items()):
        rows.append(
            {
                "month": _month_end_date(month_key),
                "series": "official_new_licenses",
                "value": float(value),
            }
        )
    return rows


def _ded_by_type_rows(stats) -> list[dict[str, object]]:
    """by_type 摊平成 ded_monthly_by_type 行。

    'unknown' 月份（日期无法解析的行）无法映射到 DATE 列，只留在审计 CSV 中。
    """

    rows: list[dict[str, object]] = []
    for (month_key, legal_form), count in sorted(stats["by_type"].items()):
        if month_key == "unknown":
            continue
        rows.append(
            {
                "month": _month_end_date(month_key),
                "legal_form": legal_form,
                "records": count,
            }
        )
    return rows


# --------------------------------------------------------------------------
# 事务上下文
# --------------------------------------------------------------------------

@contextmanager
def _transaction(con):
    """显式事务上下文（BEGIN/COMMIT/ROLLBACK）。

    注意：duckdb 1.5.5 的 con.begin() 上下文管理器在退出时会关闭连接，
    与 update_data.py 在同一连接上继续 log_run 的约定冲突，因此自行管理事务。
    """

    con.execute("BEGIN TRANSACTION")
    try:
        yield con
        con.execute("COMMIT")
    except BaseException:
        con.execute("ROLLBACK")
        raise


# --------------------------------------------------------------------------
# 月度_DED 工作簿写入（与 build_monthly_ded_sheet.py 一致）
# --------------------------------------------------------------------------

def month_to_date(month):
    """'YYYY-MM' -> 月末 datetime（工作簿单元格用）。"""
    year, m = int(month[:4]), int(month[5:7])
    if m == 12:
        return datetime(year, 12, 31)
    return datetime(year, m + 1, 1) - timedelta(days=1)


def build_rows(counts, official):
    """合并为行字典，降序排列。"""
    months = sorted(set(counts) | set(official), reverse=True)
    rows = []
    for month in months:
        c = counts.get(month) or {}
        rows.append({
            "month": month,
            "date": month_to_date(month),
            "enterprises": c.get("enterprises"),
            "licences": c.get("licences"),
            "official": official.get(month),
        })
    return rows


def update_sheet(workbook, rows):
    """写/替换 月度_DED sheet（其他 sheet 不动）。"""
    if SHEET_NAME in workbook.sheetnames:
        del workbook[SHEET_NAME]
    ws = workbook.create_sheet(SHEET_NAME)

    ws.cell(row=1, column=1, value=SOURCE_NAME)
    ws.cell(row=2, column=1, value="指标名称")
    ws.cell(row=2, column=2, value="迪拜:新增企业数(企业号首现筛重)")
    ws.cell(row=2, column=3, value="迪拜:新增执照数(执照号首现筛重)")
    ws.cell(row=2, column=4, value="迪拜:新发执照数(官方口径)")
    ws.cell(row=3, column=1, value="频率")
    ws.cell(row=3, column=2, value="月")
    ws.cell(row=3, column=3, value="月")
    ws.cell(row=3, column=4, value="月")
    ws.cell(row=4, column=1, value="单位")
    ws.cell(row=4, column=2, value="家")
    ws.cell(row=4, column=3, value="张")
    ws.cell(row=4, column=4, value="张")
    ws.cell(row=5, column=1, value="来源")
    ws.cell(row=5, column=2, value="data.dubai Commerce Registry (commerce_number 首现筛重)")
    ws.cell(row=5, column=3, value="data.dubai Commerce Registry (main_license_number 首现筛重)")
    ws.cell(row=5, column=4, value="DET/迪拜媒体办新闻稿")
    ws.cell(row=6, column=1, value="更新时间")
    now = datetime.now()
    for col in (2, 3, 4):
        ws.cell(row=6, column=col, value=now)

    for i, row in enumerate(rows):
        r = 7 + i
        ws.cell(row=r, column=1, value=row["date"])
        ws.cell(row=r, column=2, value=row["enterprises"])
        ws.cell(row=r, column=3, value=row["licences"])
        ws.cell(row=r, column=4, value=row["official"])

    for row_cells in ws.iter_rows(min_row=1, max_row=6, max_col=4):
        for cell in row_cells:
            cell.font = HEADER_FONT
    for row_cells in ws.iter_rows(min_row=7, max_row=ws.max_row, max_col=4):
        for cell in row_cells:
            cell.font = BODY_FONT

    for col, width in (("A", 13), ("B", 46), ("C", 46), ("D", 24)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A7"
    return len(rows)


# --------------------------------------------------------------------------
# 库查询 → 月度_DED 输入（merge 用）
# --------------------------------------------------------------------------

def _month_key(date_value) -> str:
    """date/datetime/str -> 'YYYY-MM'。"""
    if isinstance(date_value, str):
        return date_value[:7]
    return f"{date_value.year:04d}-{date_value.month:02d}"


def _counts_from_db(con) -> dict[str, dict[str, int]]:
    """从 ded_monthly 读取筛重序列 -> {YYYY-MM: {"enterprises": int, "licences": int}}。"""
    data: dict[str, dict[str, int]] = {}
    for month, series, value in con.execute(
        "SELECT month, series, value FROM ded_monthly WHERE series IN ('enterprises', 'licences')"
    ).fetchall():
        key = _month_key(month)
        data.setdefault(key, {})[series] = int(round(float(value)))
    return data


def _official_from_db(con) -> dict[str, int]:
    """从 ded_monthly 读取官方系列 -> {YYYY-MM: int}。"""
    data: dict[str, int] = {}
    for month, value in con.execute(
        "SELECT month, value FROM ded_monthly WHERE series = 'official_new_licenses'"
    ).fetchall():
        data[_month_key(month)] = int(round(float(value)))
    return data


# --------------------------------------------------------------------------
# 统一入口
# --------------------------------------------------------------------------

def _latest_snapshot() -> Path:
    """--skip-download 路径：取 raw/ 下最新快照（与原脚本一致）。"""
    snaps = sorted(RAW.glob("commerce_registry_*"))
    if not snaps:
        snaps = sorted(RAW.glob("ded_license_master_*"))
    if not snaps:
        raise FileNotFoundError(f"raw/ 下没有快照，无法跳过下载: {RAW}")
    return snaps[-1]


def _download_snapshots() -> tuple[list[Path], list[str]]:
    """完整下载流程（data.dubai dataset 460521）。

    返回 (本地快照路径列表, 源文件直链列表)（直链供清单记录）。
    """

    opener = build_opener(None)
    files = fetch_data_dubai_files(opener, DATASET_ID, headers={}, prefer="csv")
    for item in files:
        print(f"  {item['file_name']} ({format_bytes(item['file_size'])})")
    snapshots = []
    file_urls = []
    for item in files:
        name = Path(urllib.parse.urlparse(item["file_url"]).path).name or "datafile"
        dest = RAW / name
        print(f"下载到: {name}")
        final_url, ctype = stream_to_file(
            opener, item["file_url"], dest, headers={}, timeout=180, retries=4
        )
        print(f"下载完成: {format_bytes(dest.stat().st_size)}")
        snapshots.append(dest)
        file_urls.append(item["file_url"])
    if not snapshots:
        raise RuntimeError("没有可处理的快照文件。")
    return snapshots, file_urls


def update(
    con, *, force: bool = False, skip_download: bool = False
) -> dict[str, object]:
    """下载（或复用）快照 → 月度聚合 → 审计文件 → 事务内入库。

    异常直接抛出由总控记录；skip_download=True 时使用 raw/ 最新快照重算。
    """

    OUT.mkdir(parents=True, exist_ok=True)
    RAW.mkdir(parents=True, exist_ok=True)

    snapshots: list[Path]
    if skip_download:
        snapshots = [_latest_snapshot()]
        print(f"使用已有快照: {snapshots[0].name}")
        file_urls = ["local"]
    else:
        snapshots, file_urls = _download_snapshots()

    result = aggregate_snapshots(snapshots)
    stats = result["stats"]
    warnings = result["warnings"]

    # 审计文件（本地，结构照旧）
    write_monthly(OUT / "monthly_new_licenses.csv", stats["monthly"])
    write_monthly_counts(
        OUT / "monthly_counts.csv",
        stats["monthly"],
        stats["monthly_enterprise"],
        stats["monthly_licence"],
    )
    if result["type_col"]:
        write_monthly_by_type(OUT / "monthly_by_type.csv", stats["by_type"])
    write_schema_report(
        OUT / "schema_report.txt",
        result["columns"],
        stats,
        result["date_col"],
        result["type_col"],
        warnings,
    )
    build_manifest(
        OUT,
        dataset_id=DATASET_ID,
        source_urls=file_urls,
        snapshot_files=[s.name for s in snapshots],
        snapshot_sha256=[sha256_file(s) for s in snapshots],
        snapshot_sizes=[s.stat().st_size for s in snapshots],
        row_count=stats["total"],
        parsed_dates=stats["parsed"],
        unparsed_dates=stats["bad"],
        enterprise_count=stats["enterprise_total"],
        licence_count=stats["licence_total"],
        date_col_used=result["date_col"],
        type_col_used=result["type_col"],
        date_min=stats["date_min"],
        date_max=stats["date_max"],
        warnings=warnings,
    )

    official = read_official(OFFICIAL_CSV)
    monthly_rows = _ded_monthly_rows(stats, official)
    by_type_rows = _ded_by_type_rows(stats)
    with _transaction(con):
        replace(con, MONTHLY_TABLE, monthly_rows)
        replace(con, BY_TYPE_TABLE, by_type_rows)

    print()
    print("=" * 60)
    print(f"总行数      : {stats['total']:,}")
    print(f"日期列      : {result['date_col'] or '(未识别)'}")
    print(f"类型列      : {result['type_col'] or '(未识别)'}")
    print(f"日期范围    : {stats['date_min']} ~ {stats['date_max']}")
    print("=" * 60)

    return {
        "status": "ok",
        "rows": len(monthly_rows) + len(by_type_rows),
        "note": (
            f"快照 {len(snapshots)} 个，{stats['total']:,} 行；月度三口径 "
            f"{len(stats['monthly'])} 个月、分类型 {len(by_type_rows)} 行；"
            f"官方数字 {len(official)} 个月"
        ),
    }


def merge(workbook_path: Path) -> dict[str, object]:
    """从库查询 DED 数据，按旧脚本协议写「月度_DED」sheet（openpyxl 直写）。

    写入前在 data/UAE/raw/ded/backups/ 留时间戳备份；Excel 占用时抛出明确错误。
    """

    target = Path(workbook_path).resolve()
    if not target.exists():
        raise FileNotFoundError(f"工作簿不存在: {target}")

    con = connect(read_only=True)
    try:
        counts = _counts_from_db(con)
        official = _official_from_db(con)
    finally:
        con.close()
    if not counts and not official:
        raise RuntimeError("ded_monthly 数据为空，无可写入内容。")
    rows = build_rows(counts, official)

    # 写入前备份（遵循项目惯例；备份目录改到 data/UAE/raw/ded/backups/）
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = BACKUP_DIR / f"阿联酋_before_ded_{stamp}.xlsx"
    shutil.copy2(target, backup)

    wb = openpyxl.load_workbook(str(target))
    n = update_sheet(wb, rows)
    try:
        wb.save(str(target))
    except PermissionError as exc:
        raise RuntimeError(
            f"无法写入 {target.name}：文件被占用。\n"
            "请先关闭 Excel 中打开的 阿联酋.xlsx，然后重新运行本模块。"
        ) from exc
    note = (
        f"sheet「{SHEET_NAME}」更新 {n} 行（{rows[0]['month']} ~ {rows[-1]['month']}）；"
        f"企业列有值 {sum(1 for r in rows if r['enterprises'] is not None)} 个月，"
        f"执照列有值 {sum(1 for r in rows if r['licences'] is not None)} 个月，"
        f"官方列有值 {sum(1 for r in rows if r['official'] is not None)} 个月"
    )
    return {"status": "ok", "note": note}


if __name__ == "__main__":
    # 自检提示（不执行网络/工作簿写入）
    _snap = _latest_snapshot()
    _result = aggregate_snapshots([_snap])
    _stats = _result["stats"]
    print(
        f"[source_ded] 自检（skip-download 路径）：{_snap.name}，"
        f"{_stats['total']:,} 行，日期可解析 {_stats['parsed']:,}，"
        f"月份范围 {_stats['date_min']} ~ {_stats['date_max']}。"
    )
    print("[source_ded] 运行 data\\update_data.py --source ded 执行完整入库；"
          "data\\merge_workbook.py --source ded 回写工作簿。")