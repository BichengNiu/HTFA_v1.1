"""Process official UAE-bound labour data and update the UAE workbook."""

from __future__ import annotations

import argparse
import calendar
import csv
import json
import os
import re
import shutil
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from statistics import mean
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader


BASE_DIR = Path(__file__).resolve().parent
RAW_DIR = BASE_DIR / "raw"
PROCESSED_DIR = BASE_DIR / "processed"
REPO_ROOT = BASE_DIR.parents[1]
DEFAULT_WORKBOOK = REPO_ROOT / "data" / "阿联酋.xlsx"
SHEET_NAME = "月度_外籍劳动力"
BASELINE_START = "2025-05"
BASELINE_END = "2025-11"


def month_end(year: int, month: int) -> date:
    return date(year, month, calendar.monthrange(year, month)[1])


def numeric(value: Any) -> int:
    if value in (None, "", "-"):
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    return int(str(value).replace(",", "").strip())


def find_uae_row(worksheet) -> tuple[Any, ...]:
    for row in worksheet.iter_rows(values_only=True):
        if any(
            "UNITED ARAB EMIRATES" in str(value).upper()
            for value in row
            if value is not None
        ):
            return row
    raise RuntimeError(f"UAE row not found in {worksheet.title}")


def parse_philippines() -> dict[str, dict[str, Any]]:
    ytd: dict[tuple[int, int], dict[str, int]] = {}
    for path in sorted((RAW_DIR / "philippines").glob("*.xlsx")):
        match = re.search(r"_(\d{4})_(\d{2})\.xlsx$", path.name)
        if not match:
            continue
        year, month = map(int, match.groups())
        workbook = load_workbook(path, data_only=True, read_only=True)
        worksheet = workbook["Tab 11"] if "Tab 11" in workbook.sheetnames else workbook.worksheets[0]
        row = find_uae_row(worksheet)
        ytd[(year, month)] = {
            "total": numeric(row[3]),
            "new_hires": numeric(row[7]),
            "rehires": numeric(row[11]),
        }

    monthly: dict[str, dict[str, Any]] = {}
    previous_by_year: dict[int, dict[str, int]] = defaultdict(
        lambda: {"total": 0, "new_hires": 0, "rehires": 0}
    )
    for (year, month), values in sorted(ytd.items()):
        previous = previous_by_year[year]
        deltas = {key: values[key] - previous[key] for key in values}
        if any(value < 0 for value in deltas.values()):
            monthly[f"{year}-{month:02d}"] = {
                "total": None,
                "new_hires": None,
                "rehires": None,
                "quality_flag": (
                    "DMW官方YTD较上月回落，无法可靠差分"
                    f"（总计 {previous['total']:,}→{values['total']:,}）"
                ),
            }
        elif deltas["total"] != deltas["new_hires"] + deltas["rehires"]:
            raise RuntimeError(
                f"DMW components do not add at {year}-{month:02d}: {deltas}"
            )
        else:
            monthly[f"{year}-{month:02d}"] = {
                **deltas,
                "quality_flag": None,
            }
        previous_by_year[year] = values
    return monthly


def parse_nepal() -> dict[str, dict[str, Any]]:
    monthly: dict[str, dict[str, Any]] = {}
    date_pattern = re.compile(
        r"from\s+(\d{4})-(\d{2})-(\d{2})\s+to\s+"
        r"(\d{4})-(\d{2})-(\d{2})",
        flags=re.IGNORECASE,
    )
    for path in sorted((RAW_DIR / "nepal").glob("*.pdf")):
        reader = PdfReader(path)
        page_texts = [page.extract_text() or "" for page in reader.pages]
        if not any(text.strip() for text in page_texts):
            raise RuntimeError(f"No extractable text in DoFE PDF: {path}")
        text = "\n".join(page_texts)
        dates = date_pattern.search(text)
        if not dates:
            raise RuntimeError(f"Report date range not found: {path}")
        end = date(*map(int, dates.groups()[3:]))
        uae_line = next(
            (
                line.strip()
                for line in text.splitlines()
                if re.match(r"^\d+\s+UAE\s+", line.strip(), flags=re.IGNORECASE)
            ),
            None,
        )
        if not uae_line:
            raise RuntimeError(f"UAE row not found: {path}")
        values = [int(value) for value in re.findall(r"\d+", uae_line)][1:]
        if len(values) != 21:
            raise RuntimeError(
                f"Expected 21 DoFE measures in {path.name}, found {len(values)}"
            )
        with_reentry = values[17]
        without_reentry = values[20]
        if with_reentry < without_reentry:
            raise RuntimeError(f"Invalid DoFE re-entry totals in {path.name}")
        key = f"{end.year}-{end.month:02d}"
        if key in monthly:
            raise RuntimeError(f"Multiple DoFE periods assigned to {key}")
        monthly[key] = {
            "with_reentry": with_reentry,
            "without_reentry": without_reentry,
            "period_end": end.isoformat(),
        }
    return monthly


def parse_bangladesh() -> dict[str, int]:
    monthly: dict[str, int] = {}
    for path in sorted((RAW_DIR / "bangladesh").glob("*.json")):
        match = re.search(r"_(\d{4})_(\d{2})\.json$", path.name)
        if not match:
            continue
        year, month = map(int, match.groups())
        payload = json.loads(path.read_text(encoding="utf-8")).get("payload", {})
        rows = payload.get("data", [])
        if len(rows) != 1:
            raise RuntimeError(f"Expected one UAE OEP row in {path.name}, got {len(rows)}")
        row = rows[0]
        if "EMIRATES" not in str(row.get("country_name", "")).upper():
            raise RuntimeError(f"Unexpected OEP country in {path.name}: {row}")
        value = numeric(row.get("total_employee"))
        if numeric(payload.get("totalEmployee")) != value:
            raise RuntimeError(f"OEP total mismatch in {path.name}")
        monthly[f"{year}-{month:02d}"] = value
    return monthly


def build_rows() -> list[dict[str, Any]]:
    philippines = parse_philippines()
    nepal = parse_nepal()
    bangladesh = parse_bangladesh()
    months = sorted(set(philippines) | set(nepal) | set(bangladesh))
    if not months:
        raise RuntimeError("No processed observations found; run download_data.py first")

    overlap = [
        month
        for month in months
        if BASELINE_START <= month <= BASELINE_END
        and month in philippines
        and month in nepal
        and month in bangladesh
    ]
    if len(overlap) < 6:
        raise RuntimeError(f"Insufficient common baseline months: {overlap}")
    baselines = {
        "philippines": mean(philippines[month]["total"] for month in overlap),
        "nepal": mean(nepal[month]["with_reentry"] for month in overlap),
        "bangladesh": mean(bangladesh[month] for month in overlap),
    }

    rows: list[dict[str, Any]] = []
    for month in months:
        year, month_number = map(int, month.split("-"))
        ph = philippines.get(month)
        np = nepal.get(month)
        bd = bangladesh.get(month)
        philippines_available = ph is not None and ph["total"] is not None
        available = sum(
            (philippines_available, np is not None, bd is not None)
        )
        all_three = available == 3
        proxy_sum = (
            ph["total"] + np["with_reentry"] + bd if all_three else None
        )
        proxy_index = None
        if all_three:
            proxy_index = 100 * mean(
                (
                    ph["total"] / baselines["philippines"],
                    np["with_reentry"] / baselines["nepal"],
                    bd / baselines["bangladesh"],
                )
            )
        rows.append(
            {
                "date": month_end(year, month_number),
                "philippines_total": ph["total"] if ph else None,
                "philippines_new_hires": ph["new_hires"] if ph else None,
                "philippines_rehires": ph["rehires"] if ph else None,
                "nepal_with_reentry": np["with_reentry"] if np else None,
                "nepal_without_reentry": np["without_reentry"] if np else None,
                "bangladesh_clearance": bd,
                "proxy_sum": proxy_sum,
                "proxy_index": round(proxy_index, 6) if proxy_index is not None else None,
                "source_count": available,
                "quality_flag": ph["quality_flag"] if ph else None,
            }
        )
    return list(reversed(rows))


def write_csv(rows: list[dict[str, Any]]) -> Path:
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    path = PROCESSED_DIR / "uae_foreign_labour_monthly.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    return path


def style_sheet(worksheet, rows: list[dict[str, Any]]) -> None:
    source_fill = PatternFill("solid", fgColor="1F4E78")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    metadata_fill = PatternFill("solid", fgColor="E2F0D9")
    white_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)

    headers = [
        "日期",
        "菲律宾_DMW部署_总计",
        "菲律宾_DMW部署_新雇",
        "菲律宾_DMW部署_再雇",
        "尼泊尔_DoFE批准_含再入境",
        "尼泊尔_DoFE批准_不含再入境",
        "孟加拉国_BMET出境许可",
        "重点三国合计_可比月",
    ]
    frequencies = ["月"] * len(headers)
    units = ["日期"] + ["人"] * 7
    sources = [
        "各国官方劳工输出登记/部署/审批数据",
        "菲律宾 DMW Monthly Compendium Tab 11",
        "菲律宾 DMW Monthly Compendium Tab 11",
        "菲律宾 DMW Monthly Compendium Tab 11",
        "尼泊尔 DoFE monthly final labour approval",
        "尼泊尔 DoFE monthly final labour approval",
        "孟加拉国 BMET/OEP Country Clearance",
        "上述三国官方数据之和，仅三国均有值时计算",
    ]
    updated = max(row["date"] for row in rows)

    worksheet.append([sources[0]] + [None] * (len(headers) - 1))
    worksheet.append(headers)
    worksheet.append(frequencies)
    worksheet.append(units)
    worksheet.append(sources)
    worksheet.append(["更新时间"] + [updated] * (len(headers) - 1))
    for row in rows:
        worksheet.append(
            [
                row["date"],
                row["philippines_total"],
                row["philippines_new_hires"],
                row["philippines_rehires"],
                row["nepal_with_reentry"],
                row["nepal_without_reentry"],
                row["bangladesh_clearance"],
                row["proxy_sum"],
            ]
        )

    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    for cell in worksheet[1]:
        cell.fill = source_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for cell in worksheet[2]:
        cell.fill = header_fill
        cell.font = bold_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_number in range(3, 7):
        for cell in worksheet[row_number]:
            cell.fill = metadata_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in worksheet[6]:
        cell.font = bold_font
    for cell in worksheet["A"][6:]:
        cell.number_format = "yyyy-mm-dd"
    for row_number in range(7, worksheet.max_row + 1):
        for column in range(2, 9):
            worksheet.cell(row_number, column).number_format = "#,##0"
    worksheet.freeze_panes = "B7"
    worksheet.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{worksheet.max_row}"
    widths = [13, 22, 21, 21, 25, 27, 23, 22]
    for index, width in enumerate(widths, 1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.row_dimensions[1].height = 24
    worksheet.row_dimensions[2].height = 42


def workbook_is_open(path: Path) -> bool:
    return (path.parent / f"~${path.name}").exists()


def write_workbook(
    rows: list[dict[str, Any]], workbook_path: Path, allow_open: bool
) -> tuple[Path, Path | None]:
    output_path = workbook_path
    if not workbook_path.exists():
        workbook = Workbook()
        workbook.remove(workbook.active)
        backup = None
    else:
        if workbook_is_open(workbook_path) and not allow_open:
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_path = BASE_DIR / (
                f"{workbook_path.stem}_with_foreign_labour_{stamp}.xlsx"
            )
            backup = None
        else:
            backup_dir = BASE_DIR / "backup"
            backup_dir.mkdir(parents=True, exist_ok=True)
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup = backup_dir / (
                f"{workbook_path.stem}_before_foreign_labour_{stamp}.xlsx"
            )
            shutil.copy2(workbook_path, backup)
        workbook = load_workbook(workbook_path, keep_links=True)

    if SHEET_NAME in workbook.sheetnames:
        index = workbook.sheetnames.index(SHEET_NAME)
        workbook.remove(workbook[SHEET_NAME])
        worksheet = workbook.create_sheet(SHEET_NAME, index)
    else:
        worksheet = workbook.create_sheet(SHEET_NAME)
    style_sheet(worksheet, rows)

    temporary = output_path.with_suffix(".xlsx.tmp")
    workbook.save(temporary)
    os.replace(temporary, output_path)
    return output_path, backup


def validate(rows: list[dict[str, Any]], workbook_path: Path | None) -> None:
    dates = [row["date"] for row in rows]
    if dates != sorted(dates, reverse=True) or len(dates) != len(set(dates)):
        raise RuntimeError("Dates are not unique and newest-first")
    for row in rows:
        ph_total = row["philippines_total"]
        if ph_total is not None and ph_total != (
            row["philippines_new_hires"] + row["philippines_rehires"]
        ):
            raise RuntimeError(f"Philippines component mismatch: {row}")
        if row["proxy_sum"] is not None and row["source_count"] != 3:
            raise RuntimeError(f"Proxy emitted with partial coverage: {row}")
    if workbook_path is not None:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        if SHEET_NAME not in workbook.sheetnames:
            raise RuntimeError(f"Missing sheet {SHEET_NAME}")
        worksheet = workbook[SHEET_NAME]
        if worksheet.max_row != len(rows) + 6 or worksheet.max_column != 8:
            raise RuntimeError(
                f"Unexpected sheet shape: {worksheet.max_row}x{worksheet.max_column}"
            )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument(
        "--no-workbook",
        action="store_true",
        help="Only write the processed CSV.",
    )
    parser.add_argument(
        "--allow-open-workbook",
        action="store_true",
        help="Attempt replacement even when an Excel lock file exists.",
    )
    args = parser.parse_args()

    rows = build_rows()
    csv_path = write_csv(rows)
    workbook_path: Path | None = None
    backup: Path | None = None
    if not args.no_workbook:
        workbook_path, backup = write_workbook(
            rows, args.workbook.resolve(), args.allow_open_workbook
        )
    validate(rows, workbook_path)
    print(f"Rows: {len(rows)}")
    print(f"CSV: {csv_path}")
    if workbook_path:
        print(f"Workbook: {workbook_path}")
    if backup:
        print(f"Backup: {backup}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
