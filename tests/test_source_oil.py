"""Tests for the public OPEC/EIA oil-data source."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from htfa.jobs.uae_data import db, source_oil
from htfa.monitoring.uae.oil.data import load_oil_market_data


def _price_payload() -> bytes:
    content = BytesIO()
    frame = pd.DataFrame(
        [
            ["Sourcekey", "RBRTE"],
            ["Date", "Europe Brent Spot Price FOB (Dollars per Barrel)"],
            [pd.Timestamp("2026-01-02"), 70.5],
            [pd.Timestamp("2026-01-05"), 71.25],
        ]
    )
    with pd.ExcelWriter(content, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="Data 1", index=False, header=False)
    return content.getvalue()


def _opec_payload() -> bytes:
    return b"""
    <html><h4>Table 5 - 7</h4>
    <table>
      <tr><td>Secondary sources</td><td>Dec 25</td><td>Jan 26</td><td>Jun 26</td><td>Jul 26</td><td>Aug 26</td></tr>
      <tr><td>UAE</td><td>3200</td><td>3210</td><td>3810</td><td>3780</td><td>3835</td></tr>
    </table></html>
    """


def _dubai_payload() -> bytes:
    return b"""<?xml version='1.0' encoding='UTF-8'?>
<message:StructureSpecificData xmlns:message='http://www.sdmx.org/resources/sdmxml/schemas/v2_1/message'
    xmlns:ss='http://www.sdmx.org/resources/sdmxml/schemas/v2_1/data/structurespecific'>
  <message:DataSet>
    <message:Series>
      <message:Obs TIME_PERIOD='2026-M01' OBS_VALUE='62.72818181818182'/>
      <message:Obs TIME_PERIOD='2026-M02' OBS_VALUE='68.5085'/>
      <message:Obs TIME_PERIOD='2026-M03' OBS_VALUE='.'/>
    </message:Series>
  </message:DataSet>
</message:StructureSpecificData>
"""


def test_parse_opec_and_eia_payloads_normalize_units() -> None:
    production = source_oil.parse_opec_latest_production(_opec_payload())
    prices = source_oil.parse_brent_xls(_price_payload())
    dubai_prices = source_oil.parse_imf_dubai_xml(_dubai_payload())

    assert [(item.period, item.production_bpd) for item in production] == [
        (pd.Timestamp("2025-12-31").date(), 3_200_000.0),
        (pd.Timestamp("2026-01-31").date(), 3_210_000.0),
        (pd.Timestamp("2026-06-30").date(), 3_810_000.0),
        (pd.Timestamp("2026-07-31").date(), 3_780_000.0),
        (pd.Timestamp("2026-08-31").date(), 3_835_000.0),
    ]
    assert [(item.period, item.price_usd_bbl) for item in prices] == [
        (pd.Timestamp("2026-01-02").date(), 70.5),
        (pd.Timestamp("2026-01-05").date(), 71.25),
    ]
    assert [(item.period, item.price_usd_bbl) for item in dubai_prices] == [
        (pd.Timestamp("2026-01-31").date(), 62.72818181818182),
        (pd.Timestamp("2026-02-28").date(), 68.5085),
    ]


def test_opec_archive_supports_legacy_table_number_and_newest_revision(tmp_path) -> None:
    latest_path = tmp_path / "momr_supply_161.html"
    archive_path = tmp_path / "momr_supply_160.html"
    latest_path.write_bytes(_opec_payload())
    archive_path.write_bytes(
        _opec_payload()
        .replace(b"Table 5 - 7", b"Table 5 - 8")
        .replace(b"3835", b"3800")
    )

    observations = source_oil.parse_opec_production_reports(
        [latest_path, archive_path]
    )

    august = next(item for item in observations if item.source_period == "2026-08")
    assert august.production_bpd == 3_835_000.0
    assert august.source_file == latest_path.name


def test_update_merge_and_oil_loader_use_duckdb_as_source_of_truth(tmp_path, monkeypatch) -> None:
    price_path = tmp_path / "brent_spot_daily.xls"
    dubai_path = tmp_path / "imf_dubai_crude_monthly.xml"
    database_path = tmp_path / "uae.duckdb"
    monkeypatch.setattr(source_oil, "PRICE_RAW_PATH", price_path)
    monkeypatch.setattr(source_oil, "DUBAI_RAW_PATH", dubai_path)
    opec_path = tmp_path / "momr_world_oil_supply.html"
    home_path = tmp_path / "momr_home.html"
    monkeypatch.setattr(source_oil, "OPEC_RAW_PATH", opec_path)
    monkeypatch.setattr(source_oil, "OPEC_HOME_RAW_PATH", home_path)
    monkeypatch.setattr(source_oil, "RAW_DIR", tmp_path)
    monkeypatch.setattr(source_oil.db, "DB_PATH", database_path)

    def fake_download(url, destination, **kwargs):
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(
            _price_payload()
            if destination == price_path
            else _dubai_payload()
            if destination == dubai_path
            else b'<a href="/momr/chapter/1/2">World oil supply</a>'
            if destination == home_path
            else _opec_payload()
        )
        return "downloaded"

    monkeypatch.setattr(source_oil, "download_file", fake_download)
    con = db.connect(database_path)
    db.init_schema(con)
    try:
        outcome = source_oil.update(con, force=True)
    finally:
        con.close()

    assert outcome["status"] == "ok"
    read_con = db.connect(database_path, read_only=True)
    try:
        assert read_con.execute(
            f"SELECT period, production_bpd FROM {source_oil.OPEC_PRODUCTION_TABLE} ORDER BY period"
        ).fetchall() == [
            (pd.Timestamp("2025-12-31").date(), 3_200_000.0),
            (pd.Timestamp("2026-01-31").date(), 3_210_000.0),
            (pd.Timestamp("2026-06-30").date(), 3_810_000.0),
            (pd.Timestamp("2026-07-31").date(), 3_780_000.0),
            (pd.Timestamp("2026-08-31").date(), 3_835_000.0),
        ]
        assert read_con.execute(
            f"SELECT count(*) FROM {source_oil.EIA_PRICE_TABLE}"
        ).fetchone()[0] == 2
        assert read_con.execute(
            f"SELECT period, price_usd_bbl FROM {source_oil.DUBAI_PRICE_TABLE}"
        ).fetchall() == [
            (pd.Timestamp("2026-01-31").date(), 62.72818181818182),
            (pd.Timestamp("2026-02-28").date(), 68.5085),
        ]
    finally:
        read_con.close()

    workbook_path = tmp_path / "阿联酋.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        pd.DataFrame({"placeholder": [1]}).to_excel(
            writer, sheet_name="placeholder", index=False
        )
    merge_result = source_oil.merge(workbook_path)

    assert merge_result["status"] == "ok"
    data = load_oil_market_data(workbook_path, file_name=workbook_path.name)
    assert data.prices.columns.tolist() == ["布伦特现货"]
    assert data.production.iloc[-1] == 3_835_000.0
    assert data.metadata["布伦特现货"].source == "U.S. EIA"
    assert data.metadata["阿联酋原油产量"].source == "OPEC MOMR"

    merged_workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        assert source_oil.DUBAI_PRICE_SHEET in merged_workbook.sheetnames
        dubai_sheet = merged_workbook[source_oil.DUBAI_PRICE_SHEET]
        assert dubai_sheet.cell(2, 2).value == source_oil.DUBAI_PRICE_INDICATOR
        assert dubai_sheet.cell(7, 1).value == pd.Timestamp("2026-02-28").to_pydatetime()
        assert dubai_sheet.cell(7, 2).value == 68.5085
    finally:
        merged_workbook.close()
