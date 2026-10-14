"""Each UAE source owns its workbook merge and uses the shared writer."""

from __future__ import annotations

from datetime import date

import duckdb
import pytest
from openpyxl import Workbook, load_workbook

from htfa.jobs.uae_data.workbook_sheet_writer import write_indicator_sheets  # noqa: E402
from htfa.jobs.uae_data import (
    db, source_dot_t100, source_dubai_customs_air, source_employment,
    source_eurostat_air, source_salik, update_data, merge_workbook,
)


def test_all_update_sources_have_exactly_one_merge_owner():
    assert set(merge_workbook.SOURCES) == set(update_data.SOURCES)
    for source in update_data.SOURCES:
        assert update_data.workbook_merge_sources([source]) == [source]


@pytest.mark.parametrize("module,ddl,insert", [
    (source_dot_t100,
     "CREATE TABLE dot_t100_monthly(period DATE,indicator VARCHAR,value DOUBLE)",
     "INSERT INTO dot_t100_monthly VALUES ('2026-06-01','美国↔阿联酋_航空旅客_合计(人次)',123)"),
    (source_dubai_customs_air,
     "CREATE TABLE dubai_airway_bill_monthly(period DATE,indicator VARCHAR,value DOUBLE)",
     "INSERT INTO dubai_airway_bill_monthly VALUES ('2026-06-01','迪拜航空货运_进口总量(吨)',123)"),
    (source_salik,
     "CREATE TABLE salik_active_vehicles_quarterly(period DATE,value DOUBLE)",
     "INSERT INTO salik_active_vehicles_quarterly VALUES ('2026-06-30',123)"),
    (source_employment,
     "CREATE TABLE employment_search_index(month DATE,keyword VARCHAR,value DOUBLE)",
     "INSERT INTO employment_search_index VALUES ('2026-06-01','work in dubai',123)"),
    (source_eurostat_air,
     "CREATE TABLE detail.eurostat_air_monthly(period DATE,dataset VARCHAR,geo VARCHAR,partner VARCHAR,schedule VARCHAR,tra_meas VARCHAR,value DOUBLE)",
     "INSERT INTO detail.eurostat_air_monthly VALUES ('2026-06-30','avia_paexcc','EU27_2020','AE','TOTAL','PAS_BRD',123)"),
])
def test_native_merge_needs_only_its_own_observation_table(module, ddl, insert, tmp_path, monkeypatch):
    database = tmp_path / "source.duckdb"
    with duckdb.connect(str(database)) as con:
        con.execute("CREATE SCHEMA detail")
        con.execute(ddl)
        con.execute(insert)
        spec = module._sheet_spec(con)
        con.execute("CREATE TABLE meta_indicator_dictionary(indicator_name VARCHAR,frequency VARCHAR,unit VARCHAR,source VARCHAR)")
        con.executemany("INSERT INTO meta_indicator_dictionary VALUES (?,?,?,?)",
                        [(name, "月", "测试单位", "测试来源") for name in spec["indicators"]])
    monkeypatch.setattr(db, "connect", lambda **kwargs: duckdb.connect(str(database), **kwargs))
    workbook_path = tmp_path / "book.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "preserved"
    workbook.save(workbook_path)
    result = module.merge(workbook_path)
    assert result["status"] == "ok"
    updated = load_workbook(workbook_path, data_only=True)
    assert updated.sheetnames == ["Sheet", module.TARGET_SHEET]
    assert updated["Sheet"]["A1"].value == "preserved"
    assert updated[module.TARGET_SHEET].cell(7, 2).value == 123
    updated.close()


def test_write_indicator_sheets_uses_workbook_protocol(tmp_path) -> None:
    workbook_path = tmp_path / "workbook.xlsx"
    Workbook().save(workbook_path)
    metadata = {
        "指标A": {"frequency": "月", "unit": "点", "source": "测试源"},
        "指标B": {"frequency": "月", "unit": "人", "source": "测试源"},
    }

    counts = write_indicator_sheets(
        workbook_path,
        [
            {
                "name": "月度_测试来源",
                "title": "测试来源月度指标",
                "indicators": ["指标A", "指标B"],
                "metadata": metadata,
                "rows": [
                    (date(2026, 7, 31), {"指标A": 1.5, "指标B": 2}),
                ],
            }
        ],
    )

    assert counts == {"月度_测试来源": 1}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook["月度_测试来源"]
    assert [worksheet.cell(row, 1).value for row in range(2, 7)] == [
        "指标名称",
        "频率",
        "单位",
        "来源",
        "更新时间",
    ]
    assert [worksheet.cell(2, column).value for column in range(2, 4)] == [
        "指标A",
        "指标B",
    ]
    assert worksheet.cell(7, 2).value == 1.5
    assert worksheet.cell(7, 3).value == 2
