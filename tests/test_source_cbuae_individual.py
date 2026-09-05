"""针对当前 ``source_cbuae`` 的个人信贷（Individual / Private-Retail）提取与入库测试。"""

import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from htfa.jobs.uae_data import db  # noqa: E402
from htfa.jobs.uae_data import source_cbuae as source  # noqa: E402


def _bulletin_workbook(path: Path, *, style: str) -> None:
    """构造一份含 Domestic Credit 明细行的迷你公报。

    ``style="old"`` 使用老格式标签（Private - Corporate / Private - Retail）；
    ``style="new"`` 使用新格式标签（Corporate / Individual），且不单列
    Business and Industrial Sector 行（触发 Corporate - OtherFinancial 推导）。
    数据列落在 "January 2020"（期间 2020-01）。
    """

    workbook = Workbook()
    credit = workbook.active
    credit.title = "20 Dom Crd"
    if style == "old":
        credit.append(
            ["Table 20: Domestic Credit (All Banks)", "January 2020"]
        )
        credit.append([None, None])
        credit.append(["Government", 100])
        credit.append(["Public Sector ( GREs )", 50])
        credit.append(["Private Sector", 500])
        credit.append(["Private - Corporate", 300])
        credit.append(["Financial Institutions", 20])
        credit.append(["Business and Industrial Sector", 280])
        credit.append(["Private - Retail", 200])
    else:
        credit.append(
            ["Table 20: Domestic Credit (All Banks)", "January 2020"]
        )
        credit.append([None, None])
        credit.append(["Government", 100])
        credit.append(["Public Sector ( GREs )", 50])
        credit.append(["Private Sector", 500])
        credit.append(["Corporate", 300])
        credit.append(["Other Financial Corporations", 20])
        credit.append(["Individual", 200])

    deposit = workbook.create_sheet("25 Dep")
    deposit.append(
        [
            "Deposits distributed Residents / Non Residents ( All Banks )",
            "January 2020",
        ]
    )
    deposit.append([None, None])
    deposit.append(["Government", 200])
    deposit.append(["GREs", 100])

    currency = workbook.create_sheet("Curr")
    currency.append(
        ["Deposits by Type and Currency (All Banks)", "January 2020"]
    )
    currency.append([None, None])

    foreign = workbook.create_sheet("FA")
    foreign.append(
        ["Foreign Assets and Liabilities (All Banks)", "January 2020"]
    )
    foreign.append([None, None])

    workbook.save(path)
    workbook.close()


def test_extract_individual_from_old_and_new_format(tmp_path) -> None:
    old = tmp_path / "2024-01.xlsx"
    _bulletin_workbook(old, style="old")
    observations = source.extract_workbook(old)
    assert len(observations) == 1
    assert len(observations[0].values) == 15
    # 老格式：Private - Corporate=300，Business=280，Private - Retail=200
    assert observations[0].values[4] == Decimal("300.000")
    assert observations[0].values[5] == Decimal("280.000")
    assert observations[0].values[14] == Decimal("200.000")

    new = tmp_path / "2026-06.xlsx"
    _bulletin_workbook(new, style="new")
    observations_new = source.extract_workbook(new)
    assert len(observations_new) == 1
    # 新格式：Corporate=300，Business=300-20 推导，Individual=200
    assert observations_new[0].values[4] == Decimal("300.000")
    assert observations_new[0].values[5] == Decimal("280.000")
    assert observations_new[0].values[14] == Decimal("200.000")


def test_update_writes_individual_credit_row_and_dictionary(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(source, "RAW_DIR", tmp_path)
    _bulletin_workbook(tmp_path / "2020-03.xlsx", style="new")

    con = db.connect(":memory:")
    db.init_schema(con)
    outcome = source.update(con, skip_download=True)

    assert outcome["status"] == "ok"
    row = con.execute(
        "SELECT value, source_file FROM cbuae_monthly "
        "WHERE period = ? AND indicator = ?",
        [date(2020, 1, 31), "阿联酋:国内信贷:个人信贷(Individual Credit)"],
    ).fetchone()
    assert row == (Decimal("200.000"), "2020-03.xlsx")
    dictionary = con.execute(
        "SELECT indicator_name, type, frequency, unit, source "
        "FROM meta_indicator_dictionary WHERE indicator_name = ?",
        ["阿联酋:国内信贷:个人信贷(Individual Credit)"],
    ).fetchone()
    assert dictionary == (
        "阿联酋:国内信贷:个人信贷(Individual Credit)",
        "信贷",
        "月",
        "百万迪拉姆",
        "CBUAE",
    )
    con.close()
