"""Regression tests for the DuckDB and workbook metadata completion path."""

from __future__ import annotations

from datetime import date
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = PROJECT_ROOT / "data" / "UAE" / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import db  # noqa: E402
import uae_metadata  # noqa: E402


def test_complete_metadata_registers_extended_indicators_and_columns() -> None:
    con = db.connect(":memory:")
    try:
        con.execute(
            """
            CREATE TABLE portwatch_uae_daily (
                date DATE, portcalls BIGINT
            )
            """
        )
        con.execute("CREATE SCHEMA dld")
        con.execute(
            """
            CREATE TABLE dld.transactions (
                transaction_id VARCHAR
            )
            """
        )
        result = uae_metadata.complete_metadata(con)
        assert result["indicator_rows"] == len(uae_metadata.INDICATOR_DEFINITIONS)
        assert result["column_rows"] > 0

        names = {
            row[0]
            for row in con.execute(
                "SELECT indicator_name FROM meta_indicator_dictionary"
            ).fetchall()
        }
        assert "菲律宾_DMW部署_总计" in names
        assert "迪拜:当月新发执照数(执照号筛重)" in names
        assert "阿联酋:Google搜索热度(work in uae)" in names

        detail_columns = con.execute(
            """
            SELECT object_name, column_name
            FROM meta_column_dictionary
            WHERE object_name IN (
                'foreign_labour_monthly', 'ded_monthly',
                'detail.portwatch_uae_daily', 'detail.dld_transactions'
            )
            """
        ).fetchall()
        assert ("foreign_labour_monthly", "quality_flag") in detail_columns
        assert ("ded_monthly", "series") in detail_columns
        assert ("detail.portwatch_uae_daily", "portcalls") in detail_columns
        assert ("detail.dld_transactions", "transaction_id") in detail_columns
    finally:
        con.close()


def test_complete_metadata_is_idempotent() -> None:
    con = db.connect(":memory:")
    try:
        uae_metadata.complete_metadata(con)
        first = con.execute(
            "SELECT count(*) FROM meta_indicator_dictionary"
        ).fetchone()[0]
        uae_metadata.complete_metadata(con)
        second = con.execute(
            "SELECT count(*) FROM meta_indicator_dictionary"
        ).fetchone()[0]
        assert second == first
    finally:
        con.close()


def test_detail_schema_separates_fine_grained_physical_tables() -> None:
    con = db.connect(":memory:")
    try:
        db.init_schema(con)
        for table in (
            "comtrade_partner_detail",
            "comtrade_quantity_detail",
            "eurostat_air_monthly",
            "emirates_post_monthly",
        ):
            physical = con.execute(
                "SELECT table_type FROM information_schema.tables "
                "WHERE table_schema='detail' AND table_name=?",
                [table],
            ).fetchone()
            compatibility = con.execute(
                "SELECT table_type FROM information_schema.tables "
                "WHERE table_schema='main' AND table_name=?",
                [table],
            ).fetchone()
            assert physical == ("BASE TABLE",)
            assert compatibility == ("VIEW",)
    finally:
        con.close()


def test_legacy_dld_table_is_migrated_to_detail_with_read_only_alias() -> None:
    con = db.connect(":memory:")
    try:
        con.execute("CREATE SCHEMA dld")
        con.execute(
            "CREATE TABLE dld.transactions "
            "(transaction_id VARCHAR, instance_date DATE)"
        )
        con.execute("INSERT INTO dld.transactions VALUES ('t1', DATE '2026-01-01')")
        db.init_schema(con)
        assert con.execute("SELECT count(*) FROM detail.dld_transactions").fetchone() == (1,)
        assert con.execute("SELECT count(*) FROM dld.transactions").fetchone() == (1,)
        assert con.execute(
            "SELECT table_type FROM information_schema.tables "
            "WHERE table_schema='dld' AND table_name='transactions'"
        ).fetchone() == ("VIEW",)
    finally:
        con.close()


def test_legacy_workbook_indicators_have_source_metadata() -> None:
    names = {
        "阿联酋DFM综合股票指数",
        "全球:现货均价:原油(阿联酋穆尔班)",
        "期货结算价(连续):布伦特原油",
        "全球:现货价:原油(英国布伦特Dtd)",
        "全球:现货价:原油(阿联酋迪拜)",
        "阿联酋:现货价(CIF,低端价):石脑油",
        "阿联酋:现货价(CIF,低端价):汽油(优质无铅)",
        "阿联酋:现货价(CIF,低端价):柴油(Gasoil.1)",
        "阿联酋:现货价(CIF,低端价):煤油(航空)",
        "中国:广州市场:市场价(主流价):线型低密度聚乙烯(FB2230膜料):阿联酋博禄化工",
        "阿联酋:现货价(CIF,低端价):柴油(10ppm)",
        "印度:出口数量:咖啡:阿联酋:累计值",
        "阿联酋:银行间同业拆借利率(EIBOR):1年",
        "美元兑阿联酋迪拉姆",
        "人民币兑阿联酋迪拉姆",
        "阿联酋:银行间同业拆借利率(EIBOR):隔夜",
        "阿联酋:港口到港总次数:当月值",
        "阿联酋:港口集装箱到港次数:当月值",
        "阿联酋:港口油轮到港次数:当月值",
        "阿联酋:港口进口总量:当月值",
        "阿联酋:港口集装箱进口量:当月值",
        "阿联酋:港口油轮进口量:当月值",
        "阿联酋:港口出口总量:当月值",
        "阿联酋:港口集装箱出口量:当月值",
        "阿联酋:港口油轮出口量:当月值",
        "霍尔木兹:过境总次数:当月值",
        "霍尔木兹:载货容量:当月值",
        "霍尔木兹:油轮过境次数:当月值",
        "霍尔木兹:油轮载货容量:当月值",
    }
    definitions = {
        row["indicator_name"]: row for row in uae_metadata.INDICATOR_DEFINITIONS
    }

    assert names <= definitions.keys()
    assert all(definitions[name]["source"] for name in names)


def test_sync_workbook_dictionary_removes_unbacked_rows(tmp_path, monkeypatch) -> None:
    from openpyxl import Workbook, load_workbook

    workbook_path = tmp_path / "workbook.xlsx"
    workbook = Workbook()
    dictionary = workbook.active
    dictionary.title = "指标字典"
    dictionary.append(["指标名称", "类型", "行业", "数据来源", "预测变量"])
    dictionary.append(["指标A", "旧类型", "旧行业", "旧来源", None])
    dictionary.append(["孤立指标", "类型", "行业", "来源", None])
    data = workbook.create_sheet("月度_测试")
    data.append(["测试数据", None])
    data.append(["指标名称", "指标A"])
    data.append(["频率", "月"])
    data.append(["单位", "点"])
    data.append(["来源", "测试源"])
    data.append(["更新时间", "2026-08-20"])
    data.append(["2026-07-31", 1])
    data.append(["2026-08-31", 1])
    workbook.save(workbook_path)

    class _Result:
        def fetchall(self):
            return [("指标A", "新类型", "新行业", "新来源")]

    class _Connection:
        def execute(self, _query):
            return _Result()

        def close(self):
            pass

    monkeypatch.setattr(uae_metadata.db, "connect", lambda read_only=True: _Connection())
    result = uae_metadata.sync_workbook_dictionary(workbook_path)

    assert result["removed"] == 1
    loaded = load_workbook(workbook_path, read_only=True, data_only=True)
    rows = list(loaded.worksheets[0].iter_rows(min_row=2, values_only=True))
    assert [row[0] for row in rows if row[0]] == ["指标A"]
    assert rows[0][1:4] == ("新类型", "新行业", "新来源")


def test_sync_workbook_dictionary_splits_extended_date_columns(tmp_path, monkeypatch) -> None:
    from openpyxl import Workbook, load_workbook

    workbook_path = tmp_path / "extended-workbook.xlsx"
    workbook = Workbook()
    dictionary = workbook.active
    dictionary.title = "指标字典"
    dictionary.append([
        "指标名称", "类型", "行业", "频率", "起止时间", "缺失期数",
        "数据来源", "预测变量",
    ])
    dictionary.append([
        "指标A", "旧类型", "旧行业", "月度", "2020-01-01 至 2020-02-01", 1,
        "旧来源", "是",
    ])
    data = workbook.create_sheet("月度_测试")
    data.append(["测试数据", None])
    data.append(["指标名称", "指标A"])
    data.append(["频率", "月"])
    data.append(["单位", "点"])
    data.append(["来源", "测试源"])
    data.append(["更新时间", "2026-08-20"])
    data.append(["2026-07-31", 1])
    workbook.save(workbook_path)

    class _Result:
        def fetchall(self):
            return [("指标A", "新类型", "新行业", "新来源")]

    class _Connection:
        def execute(self, _query):
            return _Result()

        def close(self):
            pass

    monkeypatch.setattr(uae_metadata.db, "connect", lambda read_only=True: _Connection())
    uae_metadata.sync_workbook_dictionary(workbook_path, as_of=date(2026, 8, 22))

    loaded = load_workbook(workbook_path, read_only=True, data_only=True)
    assert [cell.value for cell in loaded.worksheets[0][1]] == [
        "指标名称", "类型", "行业", "频率", "开始日期", "最新日期",
        "缺失期数", "数据来源", "预测变量",
    ]
    row = next(loaded.worksheets[0].iter_rows(min_row=2, max_row=2, values_only=True))
    assert row[1:3] == ("新类型", "新行业")
    assert row[3:7] == ("月度", "2026-07-31", "2026-07-31", 1)
    assert row[7:9] == ("新来源", "是")
