"""Regression checks for live official downloads and WAM scan integrity."""

from datetime import date
import json
from urllib.error import HTTPError
from urllib.parse import parse_qs, urlparse
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from htfa.jobs.uae_data import (
    _official_download,
    db,
    fetch_wam,
    source_baker_hughes,
    source_cbuae,
    source_oil,
    source_wam,
    uae_metadata,
)


def test_baker_discovers_current_official_file_without_manual_input(monkeypatch, tmp_path):
    url = "https://rigcount.bakerhughes.com/static-files/c9738c82-95d8-45ed-8b53-a86c336598f1"
    page = f'<a href="{url}">Worldwide Rig Count Report - New Report</a>'.encode()
    monkeypatch.setattr(source_baker_hughes, "fetch_bytes", lambda *a, **k: page)
    monkeypatch.setattr(source_baker_hughes, "RAW_DIR", tmp_path)
    calls = []

    def download(request_url, destination, **kwargs):
        calls.append((request_url, destination, kwargs))
        return "downloaded"

    monkeypatch.setattr(source_baker_hughes, "download_file", download)
    note = source_baker_hughes._download_latest_worldwide_report(force=False)

    assert len(calls) == 1
    assert calls[0][0] == url
    assert calls[0][1].parent == tmp_path
    assert calls[0][2]["force"] is True
    assert url in note


def test_baker_missing_current_report_raises_without_using_old_report(monkeypatch):
    monkeypatch.setattr(source_baker_hughes, "fetch_bytes", lambda *a, **k: b"no report")
    with pytest.raises(ValueError, match="未发现"):
        source_baker_hughes._discover_latest_worldwide_report()


def _wam_paths(monkeypatch, tmp_path):
    root = tmp_path / "raw" / "wam"
    monkeypatch.setattr(fetch_wam, "RAW_DIR", root)
    monkeypatch.setattr(fetch_wam, "ARCHIVE_DIR", root / "sources")
    monkeypatch.setattr(fetch_wam, "MANIFEST_PATH", root / "source_manifest.json")
    monkeypatch.setattr(fetch_wam, "DAILY_PATH", root / "daily.csv")
    monkeypatch.setattr(fetch_wam, "AUTO_ARTICLES_PATH", root / "auto_articles.json")
    return root


def test_wam_records_third_party_archive_failure_and_continues(monkeypatch, tmp_path):
    root = _wam_paths(monkeypatch, tmp_path)
    external = "https://www.arabnews.com/node/example"
    official = "https://www.wam.ae/en/article/example"
    monkeypatch.setattr(fetch_wam, "_source_urls", lambda: [external, official])

    def fetch(url):
        if url == external:
            raise HTTPError(url, 403, "Forbidden", None, None)
        return 200, "text/html", b"official archive"

    monkeypatch.setattr(fetch_wam, "_fetch", fetch)
    outcome = fetch_wam.refresh_source_manifest()
    manifest = json.loads((root / "source_manifest.json").read_text(encoding="utf-8"))

    assert outcome == {"total": 2, "downloaded": 1, "reused": 0, "failed": 1}
    assert manifest["sources"][external]["status"] == 403
    assert manifest["sources"][official]["status"] == 200


def test_wam_official_archive_failure_still_propagates(monkeypatch, tmp_path):
    _wam_paths(monkeypatch, tmp_path)
    official = "https://www.wam.ae/en/article/example"
    monkeypatch.setattr(fetch_wam, "_source_urls", lambda: [official])

    def fetch(url):
        raise HTTPError(url, 403, "Forbidden", None, None)

    monkeypatch.setattr(fetch_wam, "_fetch", fetch)
    with pytest.raises(HTTPError):
        fetch_wam.refresh_source_manifest()


def test_wam_missing_month_cannot_be_accepted_as_zero_events(monkeypatch, tmp_path):
    _wam_paths(monkeypatch, tmp_path)
    index = (
        '<sitemapindex xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">'
        '<sitemap><loc>https://www.wam.ae/sitemap/news/en/2026/9/english.xml</loc>'
        '</sitemap></sitemapindex>'
    ).encode()
    monkeypatch.setattr(fetch_wam, "_fetch", lambda url: (200, "application/xml", index))
    with pytest.raises(RuntimeError, match="lacks requested months"):
        fetch_wam._sitemap_month_urls(date(2026, 9, 1), date(2026, 10, 12))


def test_wam_renders_javascript_shell_before_accepting_article(monkeypatch, tmp_path):
    _wam_paths(monkeypatch, tmp_path)
    url = "https://www.wam.ae/en/article/example"
    item = {
        "url": url,
        "sitemap_title": "UAE Air Force responds to UAV",
        "sitemap_published_at": "2026-10-12T10:00:00+04:00",
        "sitemap_last_modified": "2026-10-12T10:00:00+04:00",
    }
    monkeypatch.setattr(fetch_wam, "discover_auto_articles", lambda **kwargs: [item])
    monkeypatch.setattr(fetch_wam, "_fetch", lambda url: (200, "text/html", b"<app-root></app-root>"))
    article = {"@type": "NewsArticle", "headline": item["sitemap_title"], "articleBody": "Official body"}
    rendered_urls = []

    def render(request_url):
        rendered_urls.append(request_url)
        return '<script type="application/ld+json">' + json.dumps(article) + '</script>'

    monkeypatch.setattr(fetch_wam, "_render_article", render)
    counts = fetch_wam.refresh_auto_articles(end_date=date(2026, 10, 13))

    assert rendered_urls == [url]
    assert counts["downloaded"] == 1
    assert fetch_wam.load_auto_articles()[0]["article_body"] == "Official body"


def test_wam_update_writes_detail_table_behind_main_view(monkeypatch):
    row = {
        "date": date(2026, 2, 28),
        "ballistic_missiles": 0,
        "cruise_missiles": 0,
        "uavs": 0,
        "unclassified_missiles": 0,
        "strike_intensity_log": 0.0,
        "attack_any": False,
        "uae_asset_attack": False,
        "origin": None,
        "observation_status": "reviewed",
        "count_basis": None,
        "external_threat_alert": False,
        "source_url": None,
        "source_url_2": None,
        "source_url_3": None,
        "notes": None,
    }
    monkeypatch.setattr(source_wam, "load_daily_observations", lambda: [row])
    monkeypatch.setattr(source_wam, "load_effective_daily_observations", lambda: [row])
    monkeypatch.setattr(fetch_wam, "last_complete_uae_date", lambda: row["date"])
    monkeypatch.setattr(fetch_wam, "load_auto_articles", lambda: [])
    monkeypatch.setattr(fetch_wam, "refresh_source_manifest", lambda **kwargs: {
        "total": 0, "downloaded": 0, "reused": 0, "failed": 0,
    })
    monkeypatch.setattr(fetch_wam, "refresh_auto_articles", lambda **kwargs: {
        "discovered": 0, "downloaded": 0, "reused": 0,
    })
    con = db.connect(":memory:")
    try:
        db.init_schema(con)
        con.execute("CREATE VIEW main.wam_military_strike_daily AS SELECT * FROM detail.wam_military_strike_daily")
        assert source_wam.update(con)["status"] == "ok"
        assert con.execute("SELECT count(*) FROM detail.wam_military_strike_daily").fetchone()[0] == 1
        assert con.execute("SELECT count(*) FROM main.wam_military_strike_daily").fetchone()[0] == 1
    finally:
        con.close()


def test_wam_retains_daily_counts_before_later_cumulative_sentence():
    result = source_wam.extract_auto_article_observation({
        "url": "https://www.wam.ae/en/article/example",
        "headline": "UAE air defences intercept two drones",
        "article_body": (
            "ABU DHABI, 12th October, 2026 (WAM) -- UAE air defences "
            "intercepted two drones today. The forces have dealt with "
            "165 ballistic missiles since the start of hostilities."
        ),
    })
    assert result is not None
    assert result["uavs"] == 2
    assert result["ballistic_missiles"] == 0


def test_wam_does_not_publish_incomplete_initial_month():
    assert source_wam.build_monthly_observations([{
        "date": date(2026, 2, 28),
        "ballistic_missiles": 0,
        "cruise_missiles": 0,
        "uavs": 0,
        "unclassified_missiles": 0,
        "strike_intensity_log": 0.0,
        "attack_any": False,
        "uae_asset_attack": False,
        "external_threat_alert": False,
    }]) == []


def test_dictionary_accepts_empty_formatted_trailing_column():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(uae_metadata._EXTENDED_DICTIONARY_HEADERS)
    sheet.cell(2, 10).number_format = "0.00"
    assert sheet.max_column == 10
    uae_metadata._validate_dictionary_layout(sheet)
    assert sheet.max_column == 9
    workbook.close()


def test_dictionary_preserves_unexpected_trailing_data():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(uae_metadata._EXTENDED_DICTIONARY_HEADERS)
    sheet.cell(2, 10).value = "user data"
    with pytest.raises(ValueError, match="拒绝删除"):
        uae_metadata._validate_dictionary_layout(sheet)
    assert sheet.cell(2, 10).value == "user data"
    workbook.close()


def test_cbuae_automatic_download_requires_matching_observation_month(monkeypatch, tmp_path):
    url = "https://centralbank.ae/media/example/statistical-bulletin-august-2026.xlsx"
    monkeypatch.setattr(source_cbuae, "_discover_bulletins", lambda: [("2026-08", url)])
    monkeypatch.setattr(source_cbuae, "RAW_DIR", tmp_path)
    calls = []

    def download(request_url, destination, **kwargs):
        calls.append((request_url, destination, kwargs))
        destination.write_bytes(b"validated workbook")
        return "downloaded"

    monkeypatch.setattr(source_cbuae, "download_file", download)
    monkeypatch.setattr(source_cbuae, "extract_workbook", lambda path: [SimpleNamespace(period="2026-08")])
    assert "2026-08" in source_cbuae._download_latest_bulletins()
    assert calls[0][0] == url
    assert calls[0][1].name == "2026-08.xlsx"
    assert calls[0][2]["force"] is True
    cached = tmp_path / "2026-08.xlsx"
    assert cached.read_bytes() == b"validated workbook"
    cached.write_bytes(b"previous good workbook")
    monkeypatch.setattr(source_cbuae, "extract_workbook", lambda path: [SimpleNamespace(period="2026-07")])
    with pytest.raises(ValueError, match="do not match"):
        source_cbuae._download_latest_bulletins()
    assert cached.read_bytes() == b"previous good workbook"


def _eia_rows():
    return [
        {"period": "1987-12-21", "series": "RBRTE", "value": "15.4", "units": "$/BBL"},
        {"period": "1987-12-22", "series": "RBRTE", "value": "15.8", "units": "$/BBL"},
        {"period": "1987-12-23", "series": "RBRTE", "value": "15.5", "units": "$/BBL"},
    ]


def test_eia_api_validates_complete_daily_series():
    rows = _eia_rows()
    payload = {"response": {"frequency": "daily", "total": 3, "data": rows}}
    parsed = source_oil.parse_brent_api(payload)
    assert len(parsed) == 3
    assert parsed[0].period == date(1987, 12, 21)
    assert parsed[0].price_usd_bbl == 15.4
    payload["response"]["total"] = 4
    with pytest.raises(ValueError, match="incomplete"):
        source_oil.parse_brent_api(payload)
    payload["response"]["total"] = 3
    rows[1]["units"] = "cents/gal"
    with pytest.raises(ValueError, match="different series or units"):
        source_oil.parse_brent_api(payload)


def test_eia_api_fetches_every_page_without_storing_key(monkeypatch, tmp_path):
    rows = _eia_rows()
    monkeypatch.setattr(source_oil, "PRICE_RAW_PATH", tmp_path / "brent.xls")
    monkeypatch.setattr(source_oil, "_EIA_PAGE_SIZE", 2)
    monkeypatch.setenv("HTFA_EIA_API_KEY", "test-only-api-key")
    offsets = []

    def fetch(url):
        offset = int(parse_qs(urlparse(url).query)["offset"][0])
        offsets.append(offset)
        return json.dumps({"response": {
            "frequency": "daily", "total": "3", "data": rows[offset:offset + 2],
        }}).encode()

    monkeypatch.setattr(source_oil, "fetch_bytes", fetch)
    source_oil._download_brent_api()
    saved = source_oil._brent_api_raw_path().read_bytes()
    assert offsets == [0, 2]
    assert len(source_oil.parse_brent_api(saved)) == 3
    assert b"test-only-api-key" not in saved


def test_eia_failed_later_page_preserves_previous_raw_file(monkeypatch, tmp_path):
    rows = _eia_rows()
    monkeypatch.setattr(source_oil, "PRICE_RAW_PATH", tmp_path / "brent.xls")
    monkeypatch.setattr(source_oil, "_EIA_PAGE_SIZE", 2)
    target = source_oil._brent_api_raw_path()
    target.write_bytes(b"previous validated input")

    def fetch(url):
        offset = int(parse_qs(urlparse(url).query)["offset"][0])
        if offset:
            raise HTTPError(url, 503, "Unavailable", None, None)
        return json.dumps({"response": {
            "frequency": "daily", "total": 3, "data": rows[:2],
        }}).encode()

    monkeypatch.setattr(source_oil, "fetch_bytes", fetch)
    with pytest.raises(HTTPError):
        source_oil._download_brent_api()
    assert target.read_bytes() == b"previous validated input"


def test_eia_reconciliation_keeps_workbook_values_and_api_only_dates():
    workbook = [
        source_oil.PriceObservation(date(1987, 12, 21), 15.4),
        source_oil.PriceObservation(date(1987, 12, 22), 15.8),
    ]
    api = [
        source_oil.PriceObservation(date(1987, 12, 22), 15.7),
        source_oil.PriceObservation(date(1987, 12, 23), 15.5),
    ]
    rows, api_only, conflicts = source_oil.combine_brent_observations(workbook, api)
    assert [(row.period, row.price_usd_bbl) for row in rows] == [
        (date(1987, 12, 21), 15.4),
        (date(1987, 12, 22), 15.8),
        (date(1987, 12, 23), 15.5),
    ]
    assert api_only == {date(1987, 12, 23)}
    assert conflicts == 1


def test_eia_incomplete_live_coverage_cannot_delete_stored_dates():
    stored = {date(1987, 12, 21), date(1987, 12, 22)}
    with pytest.raises(ValueError, match="refusing historical row deletion"):
        source_oil._require_brent_coverage(stored, [
            source_oil.PriceObservation(date(1987, 12, 22), 15.8),
        ])


def _cbuae_listing():
    return {"Data": {"MediaListingPages": [
        {"Title": "Statistical Bulletin July 2026", "MediaList": [
            {"Url": "/media/old/statistical-bulletin-july-2026.xlsx"},
        ]},
        {"Title": "Statistical Bulletin August 2026", "MediaList": [
            {"Url": "/media/new/statistical-bulletin-august-2026.pdf"},
            {"Url": "/media/new/statistical-bulletin-august-2026.xlsx"},
        ]},
    ]}}


def test_cbuae_listing_sorts_periods_and_prefers_excel():
    rows = source_cbuae._parse_bulletin_listing(json.dumps(_cbuae_listing()).encode())
    assert [period for period, url in rows] == ["2026-08", "2026-07"]
    assert rows[0][1] == "https://centralbank.ae/media/new/statistical-bulletin-august-2026.xlsx"


def test_cbuae_listing_rejects_unofficial_media():
    payload = _cbuae_listing()
    payload["Data"]["MediaListingPages"][1]["MediaList"][1]["Url"] = (
        "https://unofficial.example/media/statistical-bulletin-august-2026.xlsx"
    )
    with pytest.raises(ValueError, match="non-official"):
        source_cbuae._parse_bulletin_listing(json.dumps(payload).encode())


def test_cbuae_discovery_reads_current_content_id(monkeypatch, tmp_path):
    identifier = "e98b57a7-c71a-494e-a6f7-8d0be9a12dbc"
    page = f'<section data-media-listing data-content-id="{identifier}">'.encode()
    calls = []

    def fetch(url):
        calls.append(url)
        return page if len(calls) == 1 else json.dumps(_cbuae_listing()).encode()

    monkeypatch.setattr(source_cbuae, "RAW_DIR", tmp_path)
    monkeypatch.setattr(source_cbuae, "_fetch_cbuae_listing", fetch)
    assert source_cbuae._discover_bulletins()[0][0] == "2026-08"
    assert parse_qs(urlparse(calls[1]).query)["ContentId"] == [identifier]
    assert (tmp_path / "listing.json").is_file()


@pytest.mark.parametrize("code,enabled,expected_calls", [
    (403, True, 1), (403, False, 0), (404, True, 1),
])
def test_browser_transport_is_selected_before_request_and_never_switches(
    monkeypatch, code, enabled, expected_calls
):
    from curl_cffi import requests
    calls = []
    ordinary_calls = []
    url = "https://centralbank.ae/en/research-and-statistics/latest-statistics/"

    def fail(*args, **kwargs):
        ordinary_calls.append(args)
        raise HTTPError(url, code, "failed", None, None)

    def compatibility_get(request_url, **kwargs):
        calls.append((request_url, kwargs))
        return SimpleNamespace(status_code=200, content=b"fresh official payload")

    monkeypatch.setattr(_official_download, "urlopen", fail)
    monkeypatch.setattr(requests, "get", compatibility_get)
    if expected_calls:
        assert _official_download.fetch_bytes(url, browser_transport=enabled) == (
            b"fresh official payload"
        )
    else:
        with pytest.raises(HTTPError):
            _official_download.fetch_bytes(url, browser_transport=enabled)
    assert len(calls) == expected_calls
    assert len(ordinary_calls) == (0 if enabled else 1)


def test_failed_browser_download_preserves_existing_raw_without_switching(monkeypatch, tmp_path):
    from curl_cffi import requests
    url = "https://publications.opec.org/momr"
    target = tmp_path / "previous.html"
    target.write_bytes(b"previous validated raw payload")

    def fail(*args, **kwargs):
        raise AssertionError("Browser transport must not switch to urllib")

    monkeypatch.setattr(_official_download, "urlopen", fail)
    monkeypatch.setattr(requests, "get", lambda *a, **k: SimpleNamespace(status_code=403))
    with pytest.raises(HTTPError):
        _official_download.download_file(url, target, force=True, browser_transport=True)
    assert target.read_bytes() == b"previous validated raw payload"
