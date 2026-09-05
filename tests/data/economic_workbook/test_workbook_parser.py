from io import BytesIO

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from htfa.data.economic_workbook.core.workbook_parser import parse_economic_workbook


def _build_workbook(
    *,
    dictionary_sheet_name: str = "指标字典",
    frequency_label: str = "频率",
    extended_dictionary: bool = False,
) -> BytesIO:
    workbook = Workbook()
    dictionary = workbook.active
    dictionary.title = dictionary_sheet_name
    if extended_dictionary:
        dictionary.append([
            "指标名称", "类型", "行业", "频率", "开始日期", "最新日期",
            "缺失期数", "数据来源", "预测变量",
        ])
        dictionary.append([
            "指标A", "指数", "金融", "日度", "2026-07-24", "2026-07-25",
            0, "Wind", "是",
        ])
        dictionary.append([
            "指标B", "产量", "能源", "月度", "2026-06-30", "2026-06-30",
            0, "Wind", None,
        ])
    else:
        dictionary.append(["指标名称", "类型", "行业", "数据来源", "预测变量"])
        dictionary.append(["指标A", "指数", "金融", "Wind", "是"])
        dictionary.append(["指标B", "产量", "能源", "Wind", None])

    daily = workbook.create_sheet("日度_Wind")
    daily.append(["Wind", None])
    daily.append(["指标名称", "指标A"])
    daily.append([frequency_label, "日"])
    daily.append(["单位", "点"])
    daily.append(["来源", "交易所"])
    daily.append(["更新时间", "2026-07-25"])
    daily.append([pd.Timestamp("2026-07-24"), 10.0])
    daily.append([pd.Timestamp("2026-07-25"), 11.0])

    monthly = workbook.create_sheet("月度_Wind")
    monthly.append(["Wind", None])
    monthly.append(["指标名称", "指标B"])
    monthly.append(["频率", "月"])
    monthly.append(["单位", "万吨"])
    monthly.append(["来源", "行业协会"])
    monthly.append(["更新时间", "2026-07-20"])
    monthly.append([pd.Timestamp("2026-06-30"), 20.0])

    weekly = workbook.create_sheet("周度_Wind")
    weekly.append([None])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    output.name = "测试数据库.xlsx"
    return output


def test_parse_workbook_merges_dictionary_and_sheet_metadata():
    result = parse_economic_workbook(_build_workbook(), module_name="test")

    assert result.get_dataframe("daily").shape == (2, 1)
    assert result.get_dataframe("monthly").shape == (1, 1)
    assert result.get_dataframe("weekly").empty
    assert result.get_dataframe("ten_day").empty

    metadata = result.indicator_metadata_map["指标A"]
    assert metadata.frequency == "daily"
    assert metadata.unit == "点"
    assert metadata.sheet_source == "交易所"
    assert metadata.dictionary_source == "Wind"
    assert metadata.updated_at == "2026-07-25"
    assert metadata.indicator_type == "指数"
    assert metadata.industry == "金融"
    assert metadata.forecast_variable == "是"
    assert metadata.file_name == "测试数据库.xlsx"
    assert metadata.sheet_name == "日度_Wind"

    assert result.source_map["指标A"] == "测试数据库|日度_Wind"
    assert result.indicator_unit_map["指标B"] == "万吨"
    assert result.indicator_freq_map["指标B"] == "monthly"


def test_parse_workbook_reads_source_and_forecast_after_date_columns():
    result = parse_economic_workbook(
        _build_workbook(extended_dictionary=True),
        module_name="test",
    )

    metadata = result.indicator_metadata_map["指标A"]
    assert metadata.dictionary_source == "Wind"
    assert metadata.forecast_variable == "是"


def test_parse_workbook_requires_named_dictionary_sheet():
    with pytest.raises(ValueError, match="指标字典"):
        parse_economic_workbook(
            _build_workbook(dictionary_sheet_name="Sheet1"),
            module_name="test",
        )


def test_parse_workbook_validates_fixed_metadata_rows():
    with pytest.raises(ValueError, match="第3行.*频率"):
        parse_economic_workbook(
            _build_workbook(frequency_label="周期"),
            module_name="test",
        )


def test_parse_workbook_accepts_reordered_unit_source_and_update_rows():
    workbook_file = _build_workbook()
    loaded_workbook = load_workbook(workbook_file)
    daily = loaded_workbook["日度_Wind"]
    daily_rows = [
        [daily.cell(row=row, column=column).value for column in range(1, 3)]
        for row in (4, 5, 6)
    ]
    for target_row, values in zip((4, 5, 6), daily_rows[1:] + daily_rows[:1]):
        for column, value in enumerate(values, start=1):
            daily.cell(row=target_row, column=column).value = value

    output = BytesIO()
    loaded_workbook.save(output)
    output.seek(0)
    output.name = "元数据换序.xlsx"

    result = parse_economic_workbook(output, module_name="test")
    metadata = result.indicator_metadata_map["指标A"]

    assert metadata.unit == "点"
    assert metadata.sheet_source == "交易所"
    assert metadata.updated_at == "2026-07-25"


def test_parse_workbook_accepts_blank_unit():
    workbook_file = _build_workbook()
    loaded_workbook = load_workbook(workbook_file)
    daily = loaded_workbook["日度_Wind"]
    daily.cell(row=4, column=2).value = None

    output = BytesIO()
    loaded_workbook.save(output)
    output.seek(0)
    output.name = "空单位.xlsx"

    result = parse_economic_workbook(output, module_name="test")

    assert result.indicator_metadata_map["指标A"].unit == ""
    assert result.indicator_unit_map["指标A"] == ""


def test_parse_workbook_rejects_value_with_blank_date():
    workbook_file = _build_workbook()
    loaded_workbook = load_workbook(workbook_file)
    daily = loaded_workbook["日度_Wind"]
    daily.cell(row=7, column=1).value = None

    output = BytesIO()
    loaded_workbook.save(output)
    output.seek(0)
    output.name = "空日期.xlsx"

    with pytest.raises(ValueError, match="第7行日期"):
        parse_economic_workbook(output, module_name="test")


def test_parse_workbook_discards_indicators_not_registered_in_dictionary():
    workbook_file = _build_workbook()
    loaded_workbook = load_workbook(workbook_file)
    dictionary = loaded_workbook["指标字典"]
    dictionary.delete_rows(3)

    output = BytesIO()
    loaded_workbook.save(output)
    output.seek(0)
    output.name = "白名单数据库.xlsx"

    result = parse_economic_workbook(output, module_name="test")

    assert result.get_dataframe("monthly").empty
    assert "指标B" not in result.indicator_metadata_map


def test_parse_workbook_skips_non_protocol_sheet_without_registered_indicators():
    workbook_file = _build_workbook()
    loaded_workbook = load_workbook(workbook_file)
    supplemental = loaded_workbook.create_sheet("月度_补充数据")
    supplemental.append(["补充来源", None])
    supplemental.append(["日期", "未登记指标"])
    supplemental.append([pd.Timestamp("2026-07-31"), 1.0])

    output = BytesIO()
    loaded_workbook.save(output)
    output.seek(0)
    output.name = "含未登记补充数据.xlsx"

    result = parse_economic_workbook(output, module_name="test")

    assert set(result.indicator_metadata_map) == {"指标A", "指标B"}


def test_parse_workbook_can_limit_parsing_to_an_indicator_allowlist():
    result = parse_economic_workbook(
        _build_workbook(),
        module_name="test",
        indicator_allowlist={"指标A"},
    )

    assert set(result.indicator_metadata_map) == {"指标A"}
    assert result.get_dataframe("monthly").empty


def test_parse_workbook_converts_zero_values_to_missing():
    workbook_file = _build_workbook()
    loaded_workbook = load_workbook(workbook_file)
    daily = loaded_workbook["日度_Wind"]
    daily.cell(row=7, column=2).value = 0

    output = BytesIO()
    loaded_workbook.save(output)
    output.seek(0)
    output.name = "零值占位数据库.xlsx"

    result = parse_economic_workbook(output, module_name="test")
    series = result.get_dataframe("daily")["指标A"]

    assert pd.isna(series.iloc[0])
    assert series.iloc[1] == 11.0
