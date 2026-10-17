"""Official quarterly GDP ingestion, units, rollback and workbook preservation."""

import csv
import io
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from htfa.jobs.uae_data import db, source_fcsc_gdp as source


def payload(system="CON", through=2):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "REF_AREA", "FREQ", "SOURCE_DETAIL", "QGDP_SYS", "QGDP_UNIT",
        "UNIT_MEASURE", "MEASURE", "TIME_PERIOD", "QUARTER", "OBS_VALUE",
    ])
    for quarter in range(1, through + 1):
        for measure in [*source.MEASURES, "MIN"]:
            level = {"TOT_GDP": 52000, "TOT_NO": 32000, "NFC": 30000,
                     "MIN": 20000}.get(measure, 2000)
            for kind, unit, value in [("VAL", "MILAED", level),
                                      ("GWTH_RATE", "PERCENT", 4.7)]:
                writer.writerow(["AE", "Q", "FCSC", system, kind, unit,
                                 measure, 2026, f"Q{quarter}", value])
    return output.getvalue().encode()


def test_parse_preserves_official_yoy_and_converts_millions_to_billions():
    rows = source.parse_gdp_csv(payload(), "CON")
    values = {(row["period"], row["indicator"]): row["value"] for row in rows}
    assert len(rows) == 44  # 19 levels + 3 official rates, for two quarters.
    assert values[date(2026, 6, 30), "阿联酋:GDP:不变价:制造业"] == 2
    assert values[date(2026, 6, 30), "阿联酋:GDP:不变价:当季同比"] == 4.7


def test_first_year_without_published_annual_rates_is_not_fabricated():
    lines = payload(through=4).decode().splitlines()
    text = "\n".join(line for line in lines if ",GWTH_RATE," not in line)
    text = text.replace(",2026,", ",2025,")
    rows = source.parse_gdp_csv(text.encode(), "CON")
    assert len(rows) == 19 * 4
    assert not any("当季同比" in row["indicator"] for row in rows)


@pytest.mark.parametrize("mutation", ["duplicate", "missing", "unit", "nonfinite"])
def test_invalid_source_is_rejected(mutation):
    text = payload().decode()
    lines = text.splitlines()
    if mutation == "duplicate":
        text += lines[1] + "\n"
    elif mutation == "missing":
        text = "\n".join(line for line in lines if ",MAN," not in line)
    elif mutation == "unit":
        text = text.replace("MILAED", "AED")
    else:
        text = text.replace(",2000", ",nan")
    with pytest.raises(ValueError):
        source.parse_gdp_csv(text.encode(), "CON")


def test_update_preserves_different_release_quarters_and_refuses_regression(tmp_path, monkeypatch):
    monkeypatch.setattr(source, "RAW_DIR", tmp_path)
    (tmp_path / "CON.csv").write_bytes(payload("CON", 2))
    (tmp_path / "CUR.csv").write_bytes(payload("CUR", 1))
    con = db.connect(":memory:")
    db.init_schema(con)
    source.update(con, skip_download=True)
    assert con.execute(
        "SELECT max(period) FROM fcsc_gdp_quarterly WHERE indicator=?",
        ["阿联酋:GDP:不变价"],
    ).fetchone()[0] == date(2026, 6, 30)
    assert con.execute(
        "SELECT max(period) FROM fcsc_gdp_quarterly WHERE indicator=?",
        ["阿联酋:GDP:现价"],
    ).fetchone()[0] == date(2026, 3, 31)
    before = con.execute("SELECT * FROM fcsc_gdp_quarterly ORDER BY 1,2").fetchall()
    (tmp_path / "CON.csv").write_bytes(payload("CON", 1))
    with pytest.raises(ValueError, match="回退"):
        source.update(con, skip_download=True)
    assert con.execute("SELECT * FROM fcsc_gdp_quarterly ORDER BY 1,2").fetchall() == before
    con.close()


def test_merge_reads_database_and_preserves_unrelated_columns(tmp_path, monkeypatch):
    database = tmp_path / "gdp.duckdb"
    monkeypatch.setattr(db, "DB_PATH", database)
    monkeypatch.setattr(source, "RAW_DIR", tmp_path)
    (tmp_path / "CON.csv").write_bytes(payload("CON", 2))
    (tmp_path / "CUR.csv").write_bytes(payload("CUR", 1))
    with db.connect() as con:
        db.init_schema(con)
        source.update(con, skip_download=True)
    path = tmp_path / "book.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = source.TARGET_SHEET
    for column, indicator in enumerate([*source.INDICATORS, "其他指标"], 2):
        sheet.cell(2, column, indicator)
    sheet.cell(7, 1, datetime(2026, 3, 31))
    sheet.cell(7, sheet.max_column, 123)
    workbook.save(path)
    workbook.close()
    source.merge(path)
    workbook = load_workbook(path)
    sheet = workbook[source.TARGET_SHEET]
    columns = {sheet.cell(2, c).value: c for c in range(2, sheet.max_column + 1)}
    assert sheet.cell(7, 1).value == datetime(2026, 6, 30)
    assert sheet.cell(7, columns["阿联酋:GDP:不变价"]).value == 52
    assert sheet.cell(7, columns["阿联酋:GDP:现价"]).value is None
    assert sheet.cell(8, columns["其他指标"]).value == 123
    workbook.close()
    from htfa.jobs.uae_data import workbook_sheet_writer

    before = path.read_bytes()

    def fail_replace(*_args):
        raise OSError("atomic replacement interrupted")

    with monkeypatch.context() as patch:
        patch.setattr(workbook_sheet_writer.os, "replace", fail_replace)
        with pytest.raises(OSError, match="interrupted"):
            source.merge(path)
    assert path.read_bytes() == before
