"""Public-source refresh regressions: periods, signs, grain and preservation."""

from datetime import date, datetime
from types import SimpleNamespace

import duckdb
import pytest
from openpyxl import Workbook, load_workbook

from htfa.jobs.uae_data import source_ded as ded
from htfa.jobs.uae_data import source_foreign_labour as labour
from htfa.jobs.uae_data import source_public_market as market
from htfa.jobs.uae_data import source_salik as salik
from htfa.jobs.uae_data import source_uaewps as wps


def test_wps_declining_wages_keep_the_negative_sign():
    assert wps._parse_wage_yoy("the average wage declined by 2.3% Y-o-Y") == -2.3
    assert wps._parse_wage_yoy("the average wage rose by 0.7% Y-o-Y") == 0.7


def test_wps_repeated_observation_month_uses_latest_publication():
    records = [("2024Q3", "2024-09", 4.0, 7.2, "first"),
               ("2024Q4", "2024-09", 4.0, 7.2, "later")]
    rows = wps._long_rows(records)
    assert len(rows) == 2
    assert {r["source_file"] for r in rows} == {"2024Q4.pdf"}


def test_wps_historical_validation_still_rejects_changed_values():
    with pytest.raises(ValueError, match="不一致"):
        wps.validate_observations({"2025Q1": ("2024-12", 0.0, 6.7, "")})


def test_wps_publication_vintage_comes_from_pdf_not_wrong_cache_name(monkeypatch, tmp_path):
    text = ("Quarterly Economic Review June 2024 Employment and Wages "
            "As of April 2024, employees covered by WPS increased by 7.5% "
            "and average wage increased by 9.4%. Government Revenue and Expenditure")
    page = SimpleNamespace(extract_text=lambda: text)
    monkeypatch.setattr(wps, "PdfReader", lambda path: SimpleNamespace(pages=[page]))
    tag, period, employees, wage, _ = wps.extract_qer_wps(tmp_path / "2024Q1.pdf")
    assert (tag, period, employees, wage) == ("2024Q2", "2024-04", 7.5, 9.4)
    rows = wps._long_rows([(tag, period, employees, wage, "")],
                          source_files={tag: "2024Q1.pdf"})
    assert rows[0]["source_file"] == "2024Q1.pdf"


@pytest.mark.parametrize("name, expected", [
    ("Salik-Earnings-Presentation-9M-2025.pdf", date(2025, 9, 30)),
    ("Salik-_-Earnings-Presentation_H1-_2024.pdf", date(2024, 6, 30)),
    ("Salik-Investor-Presentation-H1_2026.pdf", date(2026, 6, 30)),
])
def test_salik_financial_period_filename_variants(name, expected):
    assert salik._period_from_filename(name) == expected


def test_ded_future_source_row_is_preserved_but_not_aggregated():
    dirty = {"issue_date": "2027-02-23", "expiry_date": "2026-02-04",
             "commerce_number": "2536780", "main_license_number": "875617"}
    valid = {"issue_date": "2026-09-01", "commerce_number": "1", "main_license_number": "2"}
    result = ded.aggregate([dirty, valid], "issue_date", None, as_of=date(2026, 10, 14))
    assert result["total"] == 2
    assert result["monthly"] == {"2026-09": 1}
    assert result["future_rows"] == [dirty]
    assert result["enterprise_total"] == result["licence_total"] == 1


def test_nepal_new_progress_title_uses_the_cards_publication_year():
    page = '<h3 class="card__title"><a href="/new">भदौ महिनाको प्रगति विवरण</a></h3>'
    page += '<div class="post__date">२ असोज, २०८३</div>'
    def get(url, **kwargs):
        body = (page if url == labour.DOFE_MONTHLY_URL and kwargs.get("params", {}).get("page") == 1
                else '<a href="https://giwmscdnone.gov.np/media/pdf_upload/new.pdf">PDF</a>'
                if url.endswith('/new') else '')
        return SimpleNamespace(status_code=200, text=body, raise_for_status=lambda: None)
    result = labour.discover_nepal_reports(SimpleNamespace(get=get))
    assert [(r["bs_year"], r["bs_month"]) for r in result] == [(2083, 5)]


def test_public_fx_never_confuses_offshore_and_onshore_yuan():
    body = '<tr><td>Chinese Yuan - Offshore</td><td>0.8</td></tr>'
    body += '<tr><td>Chinese Yuan</td><td>0.55</td></tr>'
    assert market.parse_cny_daily(body, date(2026, 9, 1), "official")[0]["value"] == 0.55


def test_public_effr_rejects_a_different_benchmark():
    with pytest.raises(ValueError, match="different reference"):
        market.parse_effr({"refRates": [{"type": "SOFR"}]})


def test_fx_discovery_does_not_require_a_new_year_archive_before_publication():
    page = '<a href="/media/a/fx_dec26_en.xlsx">December</a>'
    page += '<a href="/media/b/fx_nov26_en.xlsx">November</a>'
    page += '<a href="/media/c/fx_dec25_en.xlsx">Old</a>'
    assert market._latest_fx_year_links(page) == [
        "/media/a/fx_dec26_en.xlsx", "/media/b/fx_nov26_en.xlsx"]


def test_public_merge_preserves_unmatched_columns_and_blanks_holidays(monkeypatch, tmp_path):
    database = tmp_path / "test.duckdb"
    with duckdb.connect(str(database)) as con:
        con.execute(f"CREATE TABLE {market.TABLE}(period DATE,indicator VARCHAR,value DOUBLE)")
        for name in market.INDICATORS:
            con.execute(f"INSERT INTO {market.TABLE} VALUES (?,?,?)", [date(2026, 9, 1), name, 3.0])
            con.execute(f"INSERT INTO {market.TABLE} VALUES (?,?,?)", [date(2026, 9, 3), name, 4.0])
    monkeypatch.setattr(market.db, "connect", lambda **k: duckdb.connect(str(database), **k))
    monkeypatch.setattr(market, "DATA_DIR", tmp_path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = market.SHEET
    names = list(market.INDICATORS) + ["unchanged"]
    for i, name in enumerate(names, 2):
        sheet.cell(2, i, name)
    sheet.cell(7, 1, datetime(2026, 9, 2))
    for i in range(2, len(names) + 2):
        sheet.cell(7, i, 99.0)
    path = tmp_path / "book.xlsx"
    workbook.save(path)
    market.merge(path)
    updated = load_workbook(path, data_only=True)
    row = next(r for r in updated[market.SHEET].iter_rows(min_row=7, values_only=True)
               if r[0] == datetime(2026, 9, 2))
    assert row[1:-1] == (None,) * len(market.INDICATORS)
    assert row[-1] == 99.0
    updated.close()


def test_public_credit_aliases_convert_million_aed_to_billion(monkeypatch, tmp_path):
    database = tmp_path / "credit.duckdb"
    with duckdb.connect(str(database)) as con:
        con.execute(f"CREATE TABLE {market.TABLE}(period DATE,indicator VARCHAR,value DOUBLE)")
        con.execute("CREATE TABLE cbuae_monthly(period DATE,indicator VARCHAR,value DOUBLE)")
        for name in market.INDICATORS:
            con.execute(f"INSERT INTO {market.TABLE} VALUES (?,?,?)", [date(2026, 9, 1), name, 3.0])
        for original in market.MONTHLY_CREDIT.values():
            con.execute("INSERT INTO cbuae_monthly VALUES (?,?,?)", [date(2026, 8, 31), original, 391964.551])
    monkeypatch.setattr(market.db, "connect", lambda **k: duckdb.connect(str(database), **k))
    monkeypatch.setattr(market, "DATA_DIR", tmp_path)
    workbook = Workbook()
    daily = workbook.active
    daily.title = market.SHEET
    monthly = workbook.create_sheet("月度_Wind")
    for sheet, names in ((daily, market.INDICATORS), (monthly, market.MONTHLY_CREDIT)):
        for i, name in enumerate(names, 2):
            sheet.cell(2, i, name)
    path = tmp_path / "book.xlsx"
    workbook.save(path)
    market.merge(path)
    updated = load_workbook(path, data_only=True)
    assert updated["月度_Wind"].cell(7, 2).value == pytest.approx(391.964551)
    assert updated["月度_Wind"].cell(4, 2).value == "十亿阿联酋迪拉姆"
    updated.close()
