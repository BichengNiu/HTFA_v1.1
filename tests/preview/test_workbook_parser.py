from io import BytesIO

import pandas as pd
import pytest
from openpyxl import Workbook

from dashboard.preview.core.workbook_parser import parse_preview_workbook


def _build_workbook(
    *,
    dictionary_sheet_name: str = "指标字典",
    frequency_label: str = "频率",
) -> BytesIO:
    workbook = Workbook()
    dictionary = workbook.active
    dictionary.title = dictionary_sheet_name
    dictionary.append(["指标名称", "类型", "行业", "数据来源", "预测变量"])
    dictionary.append(["指标A", "指数", "金融", "Wind", "是"])
    dictionary.append(["指标B", "产量", "能源", "Wind", None])

    daily = workbook.create_sheet("日度_Wind")
    daily.append(["Wind", None])
    daily.append(["指标名称", "指标A"])
    daily.append([frequency_label, "日"])
    daily.append(["单位", "点"])
    daily.append(["来源", "交易所"])
    daily.append(["更新时间", "2026-07-25"])
    daily.append([pd.Timestamp("2026-07-24"), 10.0])
    daily.append([pd.Timestamp("2026-07-25"), 11.0])

    monthly = workbook.create_sheet("月度_Wind")
    monthly.append(["Wind", None])
    monthly.append(["指标名称", "指标B"])
    monthly.append(["频率", "月"])
    monthly.append(["单位", "万吨"])
    monthly.append(["来源", "行业协会"])
    monthly.append(["更新时间", "2026-07-20"])
    monthly.append([pd.Timestamp("2026-06-30"), 20.0])

    weekly = workbook.create_sheet("周度_Wind")
    weekly.append([None])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    output.name = "测试数据库.xlsx"
    return output


def test_parse_workbook_merges_dictionary_and_sheet_metadata():
    result = parse_preview_workbook(_build_workbook(), module_name="test")

    assert result.get_dataframe("daily").shape == (2, 1)
    assert result.get_dataframe("monthly").shape == (1, 1)
    assert result.get_dataframe("weekly").empty
    assert result.get_dataframe("ten_day").empty

    metadata = result.indicator_metadata_map["指标A"]
    assert metadata.frequency == "daily"
    assert metadata.unit == "点"
    assert metadata.sheet_source == "交易所"
    assert metadata.dictionary_source == "Wind"
    assert metadata.updated_at == "2026-07-25"
    assert metadata.indicator_type == "指数"
    assert metadata.industry == "金融"
    assert metadata.forecast_variable == "是"
    assert metadata.file_name == "测试数据库.xlsx"
    assert metadata.sheet_name == "日度_Wind"

    assert result.source_map["指标A"] == "测试数据库|日度_Wind"
    assert result.indicator_unit_map["指标B"] == "万吨"
    assert result.indicator_freq_map["指标B"] == "monthly"


def test_parse_workbook_requires_named_dictionary_sheet():
    with pytest.raises(ValueError, match="指标字典"):
        parse_preview_workbook(
            _build_workbook(dictionary_sheet_name="Sheet1"),
            module_name="test",
        )


def test_parse_workbook_validates_fixed_metadata_rows():
    with pytest.raises(ValueError, match="第3行.*频率"):
        parse_preview_workbook(
            _build_workbook(frequency_label="周期"),
            module_name="test",
        )
