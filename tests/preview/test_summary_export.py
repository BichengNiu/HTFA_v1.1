from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from dashboard.preview.core.summary_export import build_summary_workbook


def test_summary_workbook_contains_calculation_note_sheet():
    summary = pd.DataFrame(
        [
            {
                '季度指标名称': '阿联酋GDP现价',
                '单位': '百万阿联酋迪拉姆',
                '类型': '金额',
                '最新值': 591946,
                '环比上季': 0.0442,
            },
            {
                '季度指标名称': '阿联酋GDP不变价：当季同比',
                '单位': '%',
                '类型': '同比增速',
                '最新值': 11.08,
                '环比上季': 6.8,
            },
        ]
    )

    workbook = load_workbook(
        BytesIO(build_summary_workbook(summary)),
        data_only=True,
    )

    assert workbook.sheetnames == ['数据摘要', '口径说明']
    note_sheet = workbook['口径说明']
    assert note_sheet['A2'].value == '按增速计算'
    assert note_sheet['B2'].value == '（本期值 - 对比期值）/ 对比期值'
    assert note_sheet['C2'].value == '阿联酋GDP现价'
    assert note_sheet['A3'].value == '按减差形式计算'
    assert note_sheet['B3'].value == '本期值 - 对比期值'
    assert note_sheet['C3'].value == '阿联酋GDP不变价：当季同比'
    assert note_sheet['D3'].value == '单位为“%”的指标，计算结果按百分点理解。'
