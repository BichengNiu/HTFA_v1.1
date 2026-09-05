"""Tests for the extended DuckDB-to-workbook sheet writer."""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]

from htfa.jobs.uae_data.workbook_sheet_writer import write_indicator_sheets  # noqa: E402
from htfa.jobs.uae_data import source_extended  # noqa: E402


def test_extended_writer_leaves_foreign_labour_to_dedicated_source() -> None:
    assert not hasattr(source_extended, "_foreign_labour")


def test_write_indicator_sheets_uses_workbook_protocol(tmp_path) -> None:
    workbook_path = tmp_path / "workbook.xlsx"
    Workbook().save(workbook_path)
    metadata = {
        "指标A": {"frequency": "月", "unit": "点", "source": "测试源"},
        "指标B": {"frequency": "月", "unit": "人", "source": "测试源"},
    }

    counts = write_indicator_sheets(
        workbook_path,
        [
            {
                "name": "月度_测试来源",
                "title": "测试来源月度指标",
                "indicators": ["指标A", "指标B"],
                "metadata": metadata,
                "rows": [
                    (date(2026, 7, 31), {"指标A": 1.5, "指标B": 2}),
                ],
            }
        ],
    )

    assert counts == {"月度_测试来源": 1}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook["月度_测试来源"]
    assert [worksheet.cell(row, 1).value for row in range(2, 7)] == [
        "指标名称",
        "频率",
        "单位",
        "来源",
        "更新时间",
    ]
    assert [worksheet.cell(2, column).value for column in range(2, 4)] == [
        "指标A",
        "指标B",
    ]
    assert worksheet.cell(7, 2).value == 1.5
    assert worksheet.cell(7, 3).value == 2
