"""Tests for the current CBUAE DuckDB source pipeline."""

import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "data" / "UAE" / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import db  # noqa: E402
import source_cbuae as source  # noqa: E402


def _bulletin_workbook(
    path: Path,
    *,
    government_deposits: tuple[int, int] = (200, 201),
) -> None:
    """Build a minimal four-table CBUAE bulletin with two observation months."""

    first_government, second_government = government_deposits
    workbook = Workbook()

    credit = workbook.active
    credit.title = "20 Dom Crd"
    credit.append(
        ["Table 20: Domestic Credit (All Banks)", "January 2020", "February 2020"]
    )
    credit.append([None, None, None])
    credit.append(["Government", 100, 101])
    credit.append(["Public Sector ( GREs )", 50, 51])
    credit.append(["Corporate", 300, 301])
    credit.append(["Other Financial Corporations", 40, 41])
    credit.append(["Individual", 200, 201])

    deposit = workbook.create_sheet("25 Dep")
    deposit.append(
        [
            "Deposits distributed Residents / Non Residents ( All Banks )",
            None,
            None,
            None,
            "January 2020",
            "February 2020",
        ]
    )
    deposit.append([None, None, None, None, None, None])
    deposit.append(["Government", None, None, None, first_government, second_government])
    deposit.append(["GREs", None, None, None, 100, 101])
    deposit.append([None, "(2)", None, None, None, None])
    deposit.append([None, None, None, "Corporate", 300, 301])
    deposit.append([None, None, None, "Individuals", 50, 51])
    deposit.append(
        [None, None, None, "Government and Non Commercial Entities", 60, 61]
    )
    deposit.append([None, None, None, "Other Financial Corporations", 40, 41])

    currency = workbook.create_sheet("Curr")
    currency.append(
        ["Deposits by Type and Currency (All Banks)", "January 2020", "February 2020"]
    )
    currency.append([None, None, None])
    currency.append(["Total Foreign Currencies", 400, 401])

    foreign = workbook.create_sheet("FA")
    foreign.append(
        ["Foreign Assets and Liabilities (All Banks)", "January 2020", "February 2020"]
    )
    foreign.append([None, None, None])
    foreign.append(["Foreign Assets", 500, 501])
    foreign.append(["Foreign Liabilities", 600, 601])

    workbook.save(path)
    workbook.close()


def test_extract_workbook_returns_both_periods(tmp_path) -> None:
    bulletin = tmp_path / "2020-03.xlsx"
    _bulletin_workbook(bulletin)

    observations = source.extract_workbook(bulletin)

    assert [item.period for item in observations] == ["2020-01", "2020-02"]
    assert observations[0].values == (
        Decimal("200.000"),
        Decimal("100.000"),
        Decimal("100.000"),
        Decimal("50.000"),
        Decimal("300.000"),
        Decimal("260.000"),
        Decimal("300.000"),
        Decimal("260.000"),
        Decimal("50.000"),
        Decimal("60.000"),
        Decimal("40.000"),
        Decimal("400.000"),
        Decimal("500.000"),
        Decimal("600.000"),
        Decimal("200.000"),
    )
    assert observations[0].source_period == "2020-03"
    assert observations[0].source_file == "2020-03.xlsx"


def test_update_roundtrip_long_table_without_wind_dependency(
    tmp_path, monkeypatch
) -> None:
    """Official bulletins populate the long table; no Excel fallback is read."""

    monkeypatch.setattr(source, "RAW_DIR", tmp_path)
    _bulletin_workbook(tmp_path / "2020-03.xlsx")

    con = db.connect(":memory:")
    db.init_schema(con)
    outcome = source.update(con, skip_download=True)

    assert outcome["status"] == "ok"
    assert outcome["rows"] == 30  # 2020-01 and 2020-02 × 15 indicators

    periods = con.execute(
        "SELECT DISTINCT period FROM cbuae_monthly ORDER BY period"
    ).fetchall()
    assert periods == [(date(2020, 1, 31),), (date(2020, 2, 29),)]

    value, source_file = con.execute(
        "SELECT value, source_file FROM cbuae_monthly "
        "WHERE period = ? AND indicator = ?",
        [date(2020, 1, 31), "阿联酋政府存款"],
    ).fetchone()
    assert value == Decimal("200.000")
    assert source_file == "2020-03.xlsx"
    con.close()


def test_update_writes_current_dictionary_rows(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(source, "RAW_DIR", tmp_path)
    _bulletin_workbook(tmp_path / "2020-03.xlsx")

    con = db.connect(":memory:")
    db.init_schema(con)
    source.update(con, skip_download=True)

    dictionary = con.execute(
        "SELECT indicator_name, type, industry, frequency, unit, source "
        "FROM meta_indicator_dictionary ORDER BY indicator_name"
    ).fetchall()
    con.close()
    assert len(dictionary) == len(source.CBUAE_INDICATORS)
    assert (
        "阿联酋政府存款",
        "存款",
        "金融",
        "月",
        "百万迪拉姆",
        "CBUAE",
    ) in dictionary
    assert (
        "阿联酋:国内信贷:个人信贷(Individual Credit)",
        "信贷",
        "金融",
        "月",
        "百万迪拉姆",
        "CBUAE",
    ) in dictionary


def test_select_latest_vintages_prefers_newer_bulletin() -> None:
    """The newer source vintage wins and the revision count is reported."""

    old = source.Observation(
        "2020-01", (Decimal("200"),), "2020-02", "2020-02.xlsx"
    )
    new = source.Observation(
        "2020-01", (Decimal("250"),), "2020-03", "2020-03.xlsx"
    )

    selected, missing, revised = source.select_latest_vintages(
        [old, new], start_period="2020-01", end_period="2020-01"
    )

    assert selected[0] is new
    assert missing == []
    assert revised == 1
