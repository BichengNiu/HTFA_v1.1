"""Atomic daily publication and portable workbook protocol regression tests."""

from datetime import datetime
from io import BytesIO
import hashlib
import json
from pathlib import Path
import shutil
from unittest.mock import MagicMock

import duckdb
from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest

from scripts import cloud_refresh as runner
from htfa.jobs.uae_data import db, _excel_helpers as helpers
from htfa.jobs.uae_data.portable_workbook import write_payload, materialize_report
from tests.analysis.uae._helpers import report_source


def report(value=0, include=True):
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame([
            ["测试", None], ["指标名称", "指标"], ["频率", "月"], ["单位", "点"],
            ["来源", "测试"], ["更新时间", "2026-01-31"],
            ["2026-01-31", value if include else None],
        ]).to_excel(writer, sheet_name="月度_测试", header=False, index=False)
    return report_source(buffer.getvalue())


@pytest.fixture
def live_data(tmp_path, monkeypatch):
    data = tmp_path / "data/UAE"
    data.mkdir(parents=True)
    source = report()
    shutil.copy2(source.path, data / "uae.duckdb")
    monkeypatch.setattr(runner, "DATA", data)
    monkeypatch.setattr(runner, "git", lambda *args: "baseline" if args[0] == "rev-parse" else "")
    return data


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def test_failed_refresh_never_replaces_live_database(live_data, monkeypatch):
    baseline = live_data / "uae.duckdb"
    before = digest(baseline)
    def fail(candidate, *args, **kwargs):
        with duckdb.connect(str(candidate / "uae.duckdb")) as con:
            con.execute("DELETE FROM report.observations")
        raise RuntimeError("source failed")
    monkeypatch.setattr(runner, "run_worker", fail)
    with pytest.raises(RuntimeError, match="source failed"):
        runner.refresh(("fcsc",))
    assert digest(baseline) == before


def test_successful_refresh_publishes_only_db_and_records_daily_success(live_data, monkeypatch):
    workbook = live_data / "阿联酋.xlsx"
    workbook.write_bytes(b"must not be overwritten")
    def revise(candidate, *args, **kwargs):
        with duckdb.connect(str(candidate / "uae.duckdb")) as con:
            con.execute("UPDATE report.observations SET value=2")
            con.execute("UPDATE report.release SET version='revised'")
    monkeypatch.setattr(runner, "run_worker", revise)
    result = runner.refresh(("fcsc",), day="2026-10-28")
    assert result["status"] == "refreshed"
    assert workbook.read_bytes() == b"must not be overwritten"
    with duckdb.connect(str(live_data / "uae.duckdb"), read_only=True) as con:
        assert con.execute("SELECT value FROM report.observations").fetchone()[0] == 2
        row = con.execute("SELECT refresh_date,sources_json FROM cloud_refresh_runs").fetchone()
        assert str(row[0]) == "2026-10-28"
        assert json.loads(row[1]) == ["fcsc"]


def test_remote_race_and_validation_only_leave_live_db_unchanged(live_data, monkeypatch):
    before = digest(live_data / "uae.duckdb")
    monkeypatch.setattr(runner, "run_worker", lambda *args, **kwargs: None)
    assert runner.refresh(("fcsc",), validate_only=True)["status"] == "validated_only"
    assert digest(live_data / "uae.duckdb") == before
    monkeypatch.setattr(runner, "git", lambda *args: "old" if args == ("rev-parse", "HEAD") else "changed")
    with pytest.raises(RuntimeError, match="Remote main changed"):
        runner.refresh(("fcsc",))
    assert digest(live_data / "uae.duckdb") == before


def test_history_guard_allows_revision_but_rejects_lost_values():
    old, revised, missing = report(0), report(3), report(include=False)
    runner.protect_history(old.path, revised.path)
    with pytest.raises(RuntimeError, match="remove existing report history"):
        runner.protect_history(old.path, missing.path)


def test_daily_guard_requires_today_db_commit_and_offline_cannot_publish(monkeypatch):
    monkeypatch.setattr(runner, "git", lambda *args: "chore: Streamlit keep-alive\n" + runner.COMMIT_PREFIX + "2026-10-27")
    assert runner.refresh_due("2026-10-28")
    assert not runner.refresh_due("2026-10-27")
    assert runner.refresh_due("2026-10-27", force=True)
    with pytest.raises(ValueError, match="Offline merge validation cannot publish"):
        runner.refresh(("fcsc",), merge_only=True)
    assert "comtrade" not in runner.MONTHLY_SOURCES
    assert "cloudflare_radar" not in runner.MONTHLY_SOURCES
    assert not set(runner.MONTHLY_SOURCES) & {"fcsc_gdp", "fcsc_annual", "wdi_gcc"}


def test_materialized_workbook_retains_zero_and_indicator_metadata(tmp_path):
    source = report(0)
    path = tmp_path / "scratch.xlsx"
    materialize_report(Path(source.path), path)
    book = load_workbook(path, data_only=True)
    try:
        assert book.sheetnames[0] == "指标字典"
        assert book["月度_测试"].cell(7, 2).value == 0
        assert book["月度_测试"].cell(2, 2).value == "指标"
    finally:
        book.close()


@pytest.fixture
def payload_book(tmp_path, monkeypatch):
    path = tmp_path / "book.xlsx"
    book = Workbook()
    book.active.title = "指标字典"
    book.active.append(["指标名称", "类型", "行业", "频率", "开始日期", "最新日期", "缺失期数", "数据来源", "预测变量"])
    book.save(path)
    book.close()
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "test.duckdb")
    with db.connect() as con:
        db.init_schema(con)
        db.upsert_dictionary_rows(con, [
            {"indicator_name": name, "frequency": "月", "unit": unit, "source": "官方", "type": "指标", "industry": "测试", "updated_at": datetime.now().date()}
            for name, unit in [("金额", "百万AED"), ("比例", "%")]
        ])
    payload = {"frequency": "月", "source": "官方", "unit": "点",
               "indicators": [{"name": "金额", "unit": "百万AED"}, {"name": "比例", "unit": "%"}],
               "records": [{"period": "2024-02", "values": [0, None]}]}
    json_path = tmp_path / "payload.json"
    json_path.write_text(json.dumps(payload))
    return path, json_path, payload


def test_portable_payload_preserves_units_leap_month_zero_and_null(payload_book):
    path, json_path, _ = payload_book
    write_payload(path, json_path, "月度_测试")
    book = load_workbook(path, data_only=True)
    try:
        sheet = book["月度_测试"]
        assert sheet.cell(4, 2).value == "百万AED"
        assert sheet.cell(4, 3).value == "%"
        assert sheet.cell(7, 1).value == datetime(2024, 2, 29)
        assert sheet.cell(7, 2).value == 0
        assert sheet.cell(7, 3).value is None
    finally:
        book.close()


def test_portable_mode_is_scoped_and_invalid_payload_does_not_write(payload_book, monkeypatch):
    path, json_path, payload = payload_book
    completed = MagicMock(returncode=0)
    command = MagicMock(return_value=completed)
    monkeypatch.setattr(helpers, "run_powershell_command", command)
    with helpers.portable_workbook_writes():
        helpers.run_powershell_sheet_writer(Path("writer.ps1"), path, json_path, "月度_测试")
    command.assert_not_called()
    helpers.run_powershell_sheet_writer(Path("writer.ps1"), path, json_path, "月度_测试")
    command.assert_called_once()
    before = digest(path)
    payload["records"][0]["values"] = [1]
    json_path.write_text(json.dumps(payload))
    with pytest.raises(ValueError, match="width mismatch"):
        write_payload(path, json_path, "月度_测试")
    assert digest(path) == before
